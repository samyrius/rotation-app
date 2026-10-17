import hashlib
from flask.testing import FlaskClient
import io
import sqlite3
from datetime import timedelta
import re
import zipfile
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import app as rotation_app


class CsrfAwareTestClient(FlaskClient):
    def post(self, *args, **kwargs):
        data = dict(kwargs.get("data") or {})
        if "_csrf_token" not in data:
            with self.session_transaction() as sess:
                token = sess.get("_csrf_token")
            if not token:
                self.get("/login")
                with self.session_transaction() as sess:
                    token = sess.get("_csrf_token")
            if token:
                data["_csrf_token"] = token
        kwargs["data"] = data
        return super().post(*args, **kwargs)


def setup_client(tmp_path, monkeypatch):
    rotation_app.app.test_client_class = CsrfAwareTestClient
    monkeypatch.setenv("ROTATION_INITIAL_PASSWORD", "test-strong-password-2026")
    monkeypatch.setattr(rotation_app, "DB", str(tmp_path / "test_rotation.db"))
    monkeypatch.setattr(rotation_app.app, "_ready", False, raising=False)
    rotation_app.app.config.update(TESTING=True, SECRET_KEY="test-only-secret")
    return rotation_app.app.test_client()


def login(client):
    return client.post("/login", data={"password": "test-strong-password-2026"}, follow_redirects=True)



def test_internal_navigation_is_grouped_responsive_and_marks_current_page(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.get("/")
    assert response.status_code == 200
    assert b'class="rt-app-shell"' in response.data
    assert b'aria-label="Main navigation"' in response.data or b'aria-label="\xd8\xa7\xd9\x84\xd8\xaa\xd9\x86\xd9\x82\xd9\x84 \xd8\xa7\xd9\x84\xd8\xb1\xd8\xa6\xd9\x8a\xd8\xb3\xd9\x8a"' in response.data
    assert b'href="/leave"' in response.data
    assert b'href="/external-rotations"' in response.data
    assert "التربصات والتخطيط".encode() in response.data or b'Rotations &amp; planning' in response.data or b'Rotations & planning' in response.data
    assert b'class="nav-home is-active"' in response.data
    assert b'grid-template-columns:260px minmax(0,1fr)' in response.data
    assert b'@media(max-width:680px)' in response.data
    residents = client.get("/residents")
    assert residents.status_code == 200
    assert b'href="/residents" class="is-active"' in residents.data


def test_login_page_and_language_selector(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    response = client.get("/login")
    assert response.status_code == 200
    assert "Français".encode() in response.data
    assert "English".encode() in response.data
    assert b'name="_csrf_token"' in response.data



def test_login_branding_uses_resimanage_and_navy_cyan_identity(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    response = client.get("/login")
    assert response.status_code == 200
    assert b"ResiManage" in response.data
    assert b"ResiTrack" not in response.data
    assert b"#00dff755" in response.data
    assert b"#19dfbd" in response.data
    assert b"prefers-reduced-motion:reduce" in response.data
    assert b"Tous droits r\xc3\xa9serv\xc3\xa9s" in response.data
    assert b"All rights reserved" in response.data
    assert "جميع الحقوق محفوظة".encode() in response.data
    assert "✚".encode() not in response.data


def test_login_and_protected_dashboard(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    response = login(client)
    assert response.status_code == 200
    assert "Resident".encode() in response.data or "المقيمون".encode() in response.data
    client.get("/logout")
    protected = client.get("/residents")
    assert protected.status_code == 302
    assert "/login" in protected.headers["Location"]


def test_can_add_resident_without_touching_real_database(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.post(
        "/residents",
        data={"surname": "TestSurname", "firstname": "TestFirst", "year": "2", "phone": "000"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"TestSurname" in response.data
    with rotation_app.db() as conn:
        row = conn.execute("SELECT surname, firstname, year FROM residents WHERE surname=?", ("TestSurname",)).fetchone()
    assert tuple(row) == ("TestSurname", "TestFirst", 2)


def test_seed_has_nineteen_terrains_and_four_periods(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM terrains").fetchone()[0] == 19
        assert conn.execute("SELECT COUNT(*) FROM periods").fetchone()[0] == 4
        optional = {r[0] for r in conn.execute("SELECT name FROM terrains WHERE kind='optional'")}
    assert optional == {"Sédation", "Maxillo-faciale", "CAC DBK", "Brûlés", "Chirurgie cardiaque"}
    assert "Gynécologie" not in optional
    with rotation_app.db() as conn:
        assert conn.execute("SELECT kind FROM terrains WHERE name='Gynécologie'").fetchone()[0] != "optional"
        assert conn.execute("SELECT kind FROM terrains WHERE name='Chirurgie générale'").fetchone()[0] == "mandatory"


def test_assignment_requires_capacity_and_blocks_same_period_duplicate(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        resident_id = conn.execute("SELECT id FROM residents WHERE surname='TestSurname'").fetchone()
        if resident_id is None:
            conn.execute(
                "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
                ("AssignTest", "Resident", 2),
            )
            resident_id = conn.execute("SELECT id FROM residents WHERE surname='AssignTest'").fetchone()
        rid = resident_id[0]
        terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        second_terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]

    response = client.post(
        "/assign",
        data={"resident": str(rid), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "حدّد السعة أولاً".encode() in response.data

    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO capacities(year_label, terrain_id, period_id, places) VALUES(?,?,?,1)",
            (year_label, terrain_id, period_id),
        )

    response = client.post(
        "/assign",
        data={"resident": str(rid), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assigned = conn.execute(
            "SELECT status FROM assignments WHERE resident_id=? AND period_id=?",
            (rid, period_id),
        ).fetchone()
    assert assigned is not None
    assert assigned["status"] == "confirmed"

    response = client.post(
        "/assign",
        data={"resident": str(rid), "terrain": str(second_terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert response.status_code == 200
    with rotation_app.db() as conn:
        count = conn.execute(
            "SELECT COUNT(*) FROM assignments WHERE resident_id=? AND period_id=?",
            (rid, period_id),
        ).fetchone()[0]
    assert count == 1


def test_capacity_limit_is_enforced_for_two_residents(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
            ("CapacityOne", "Resident", 3),
        )
        conn.execute(
            "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
            ("CapacityTwo", "Resident", 2),
        )
        ids = conn.execute(
            "SELECT id, surname FROM residents WHERE surname IN ('CapacityOne','CapacityTwo')"
        ).fetchall()
        resident_ids = {r["surname"]: r["id"] for r in ids}
        terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Réanimation polyvalente'").fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]
        conn.execute(
            "INSERT INTO capacities(year_label, terrain_id, period_id, places) VALUES(?,?,?,1)",
            (year_label, terrain_id, period_id),
        )

    first = client.post(
        "/assign",
        data={"resident": str(resident_ids["CapacityOne"]), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert first.status_code == 200
    second = client.post(
        "/assign",
        data={"resident": str(resident_ids["CapacityTwo"]), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert second.status_code == 200
    assert "السعة مكتملة".encode() in second.data
    with rotation_app.db() as conn:
        count = conn.execute(
            "SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status='confirmed'",
            (terrain_id, period_id),
        ).fetchone()[0]
    assert count == 1


def test_repeat_is_blocked_until_all_core_terrains_are_completed(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
            ("RepeatRule", "Resident", 4),
        )
        rid = conn.execute("SELECT id FROM residents WHERE surname='RepeatRule'").fetchone()[0]
        target = conn.execute(
            "SELECT id FROM terrains WHERE name='Réanimation chirurgicale'"
        ).fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]
        conn.execute(
            "INSERT INTO capacities(year_label, terrain_id, period_id, places) VALUES(?,?,?,2)",
            (year_label, target, period_id),
        )
        conn.execute(
            "INSERT INTO history(resident_id, terrain_id, year_label, period, date_done) VALUES(?,?,?,?,?)",
            (rid, target, "2025-2026", "01 Fév - 30 Avr", "2026-04-30"),
        )

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(conn, rid, target, period_id, year_label)
            assert allowed is False
            assert "التربص منجز سابقاً" in message

    with rotation_app.db() as conn:
        conn.execute(
            """
            INSERT INTO history(resident_id, terrain_id, year_label, period, date_done)
            SELECT ?, id, '2025-2026', 'historique', '2026-04-30'
            FROM terrains
            WHERE kind IN ('mandatory', 'basic')
              AND id NOT IN (SELECT terrain_id FROM history WHERE resident_id=?)
            """,
            (rid, rid),
        )

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(conn, rid, target, period_id, year_label)
    assert allowed is True
    assert status == "confirmed"
    assert "يمكن التكرار" in message

def test_resident_profile_and_objectives_are_saved_and_audited(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)", ("Profile", "Resident", 2))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Profile'").fetchone()[0]
        terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
    response = client.post("/profile/" + str(rid), data={
        "birth_place": "Tizi Ouzou", "birth_date": "1998-05-02",
        "phone": "0550000000", "photo_path": ""
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT birth_place,birth_date,phone FROM residents WHERE id=?", (rid,)).fetchone()
        assert tuple(row) == ("Tizi Ouzou", "1998-05-02", "0550000000")
    response = client.post("/objectives", data={
        "terrain": str(terrain_id), "study_year": "2", "objective": "Manage airway safely"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b"Manage airway safely" in response.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='resident.profile_updated'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM objectives WHERE objective='Manage airway safely'").fetchone()[0] == 1


def test_first_year_progression_requires_every_component_at_least_ten(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("Pass", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Pass'").fetchone()[0]
    for module, theory, practical in [("Module A", 12, 11), ("Module B", 10, 10)]:
        client.post("/exams", data={"action":"record","resident":str(rid),"study_year":"1","module":module,"theory":str(theory),"practical":str(practical)})
    client.post("/exams", data={"action":"evaluate_first_year","resident":str(rid),"study_year":"1","all_modules_confirmed":"yes"})
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 2
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("Fail", "Resident"))
        fail_id = conn.execute("SELECT id FROM residents WHERE surname='Fail'").fetchone()[0]
    for module, theory, practical in [("Module A", 15, 15), ("Module B", 9, 16)]:
        client.post("/exams", data={"action":"record","resident":str(fail_id),"study_year":"1","module":module,"theory":str(theory),"practical":str(practical)})
    client.post("/exams", data={"action":"evaluate_first_year","resident":str(fail_id),"study_year":"1","all_modules_confirmed":"yes"})
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (fail_id,)).fetchone()[0] == 1


def test_first_year_progression_does_not_advance_inactive_resident(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,1,'inactive')",
            ("InactivePass", "Resident"),
        )
        rid = conn.execute(
            "SELECT id FROM residents WHERE surname='InactivePass'"
        ).fetchone()[0]
        conn.execute(
            "INSERT INTO exam_records(resident_id,study_year,module,theory,practical) VALUES(?,1,?,?,?)",
            (rid, "Module A", 12, 12),
        )
        conn.commit()

    response = client.post(
        "/exams",
        data={
            "action": "evaluate_first_year",
            "resident": str(rid),
            "study_year": "1",
            "all_modules_confirmed": "yes",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "resident is inactive".encode() in response.data
    with rotation_app.db() as conn:
        resident = conn.execute(
            "SELECT year,status FROM residents WHERE id=?", (rid,)
        ).fetchone()
        assert (resident["year"], resident["status"]) == (1, "inactive")
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action='resident.year_advanced' AND details LIKE ?",
            (f"%resident_id={rid};%",),
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM decisions WHERE details LIKE ?",
            ("%InactivePass Resident — first-year exams%",),
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action='exam.first_year_evaluated' AND details LIKE ?",
            (f"%resident_id={rid};%",),
        ).fetchone()[0] == 0


def test_dems_archive_keeps_resident_and_history(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={"username": username, "password": "strong-test-password-2026", "role": "super_admin"}, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,4)", ("Graduate", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Graduate'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        conn.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                     (rid, tid, "2025-2026", "01 Mai - 30 Juillet", "2025-07-30"))
        conn.commit()

    with client.session_transaction() as sess:
        sess["user"] = "benhocine"
        sess["role"] = "super_admin"
    first_vote = client.post("/dems", data={"resident":str(rid),"success_date":"2026-10-01","outcome":"approve","reason":"DEMS passed"}, follow_redirects=True)
    assert first_vote.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM residents WHERE id=?", (rid,)).fetchone()[0] == "active"
        assert conn.execute("SELECT COUNT(*) FROM dems_archive WHERE resident_id=?", (rid,)).fetchone()[0] == 0

    with client.session_transaction() as sess:
        sess["user"] = "drdris"
        sess["role"] = "super_admin"
    second_vote = client.post("/dems", data={"resident":str(rid),"success_date":"2026-10-01","outcome":"approve","reason":"Reviewed and approved"}, follow_redirects=True)
    assert second_vote.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT status,dems_status FROM residents WHERE id=?", (rid,)).fetchone()
        assert tuple(row) == ("graduated", "success")
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=?", (rid,)).fetchone()[0] == 1
        archive = conn.execute("SELECT history_json FROM dems_archive WHERE resident_id=?", (rid,)).fetchone()
        assert "Urologie" in archive["history_json"]
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='dems.vote_recorded'").fetchone()[0] == 2


def test_standard_user_cannot_access_super_admin_pages(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.post("/users", data={
        "username":"ordinary-user","password":"ordinary-password-2026","role":"user"
    }, follow_redirects=True)
    assert response.status_code == 200
    client.get("/logout")
    response = client.post("/login", data={
        "username":"ordinary-user","password":"ordinary-password-2026"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert client.get("/settings").status_code == 403
    assert client.get("/dems").status_code == 403
    assert client.get("/audit").status_code == 403

def test_standard_user_cannot_create_residents(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    client.post(
        "/users",
        data={
            "username": "standard-operator",
            "password": "standard-operator-password-2026",
            "role": "user",
        },
        follow_redirects=True,
    )
    with client.session_transaction() as sess:
        sess["user"] = "standard-operator"
        sess["role"] = "user"

    response = client.post(
        "/residents",
        data={
            "surname": "Unauthorized",
            "firstname": "Resident",
            "year": "2",
            "phone": "",
        },
        follow_redirects=False,
    )
    assert response.status_code == 403
    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM residents WHERE surname='Unauthorized'"
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action='resident.created' AND details='Unauthorized Resident'"
        ).fetchone()[0] == 0


def test_standard_user_cannot_record_or_change_exam_marks(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    client.post(
        "/users",
        data={
            "username": "exam-operator",
            "password": "exam-operator-password-2026",
            "role": "user",
        },
        follow_redirects=True,
    )
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)",
            ("ProtectedExam", "Resident"),
        )
        rid = conn.execute(
            "SELECT id FROM residents WHERE surname='ProtectedExam'"
        ).fetchone()[0]
        conn.commit()
    with client.session_transaction() as sess:
        sess["user"] = "exam-operator"
        sess["role"] = "user"

    response = client.post(
        "/exams",
        data={
            "action": "record",
            "resident": str(rid),
            "study_year": "1",
            "module": "Unauthorized module",
            "theory": "20",
            "practical": "20",
        },
        follow_redirects=False,
    )
    assert response.status_code == 403
    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM exam_records WHERE resident_id=?", (rid,)
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action='exam.recorded'"
        ).fetchone()[0] == 0


def test_backup_and_restore_preserve_database_and_resident_photos(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    photo_dir = tmp_path / "resident_photos"
    photo_dir.mkdir(exist_ok=True)
    photo_path = photo_dir / "test-resident.webp"
    photo_path.write_bytes(b"original-photo-bytes")
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,?)", ("BeforeBackup", "Resident", 2))
    backup_response = client.get("/backup")
    assert backup_response.status_code == 200
    assert backup_response.mimetype == "application/zip"
    backup_bytes = backup_response.get_data()
    with zipfile.ZipFile(io.BytesIO(backup_bytes)) as bundle:
        assert "rotation.db" in bundle.namelist()
        assert "resident_photos/test-resident.webp" in bundle.namelist()
        assert bundle.read("resident_photos/test-resident.webp") == b"original-photo-bytes"
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,?)", ("AddedAfterBackup", "Resident", 1))
    photo_path.write_bytes(b"changed-photo-bytes")
    restored = client.post("/restore", data={
        "backup_file": (io.BytesIO(backup_bytes), "resimanage_full_backup.zip")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert restored.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM residents WHERE surname='BeforeBackup'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM residents WHERE surname='AddedAfterBackup'").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='database.restored'").fetchone()[0] == 1
    assert (tmp_path / "resident_photos" / "test-resident.webp").read_bytes() == b"original-photo-bytes"
    assert list(tmp_path.glob("rotation_backup_before_restore_*.db"))
    assert list(tmp_path.glob("rotation_backup_before_restore_*.zip"))

def test_csv_import_requires_preview_and_does_not_invent_history(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    source = b"surname,firstname,year,phone\nImportSurname,ImportFirst,3,0551234567\n"
    preview = client.post("/import", data={
        "file": (io.BytesIO(source), "residents.csv")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200, preview.data.decode("utf-8", errors="replace")
    assert b"ImportSurname" in preview.data, preview.data.decode("utf-8", errors="replace")
    match = re.search(rb'name="preview_id" value="(\d+)"', preview.data)
    assert match
    confirmed = client.post("/import", data={
        "action":"confirm", "preview_id":match.group(1).decode()
    }, follow_redirects=True)
    assert confirmed.status_code == 200
    with rotation_app.db() as conn:
        rid = conn.execute("SELECT id FROM residents WHERE surname='ImportSurname'").fetchone()[0]
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=?", (rid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='residents.imported'").fetchone()[0] == 1

def test_mandatory_year_override_requires_reason_and_records_audit(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("Override", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Override'").fetchone()[0]
        conn.execute("INSERT INTO terrains(name,kind,mandatory_year) VALUES(?,?,3)", ("Year Three Only", "mandatory"))
        tid = conn.execute("SELECT id FROM terrains WHERE name='Year Three Only'").fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]
        conn.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)", (year_label,tid,period_id))
    response = client.post("/assign", data={"resident":str(rid),"terrain":str(tid),"period":str(period_id)}, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        aid = conn.execute("SELECT id FROM assignments WHERE resident_id=? AND status='pending'", (rid,)).fetchone()[0]
    client.post("/decisions", data={"aid":str(aid),"outcome":"approve","reason":"Authorized educational exception"}, follow_redirects=True)
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM assignments WHERE id=?", (aid,)).fetchone()[0] == "confirmed"
        assert conn.execute("SELECT COUNT(*) FROM decisions WHERE reason='Authorized educational exception'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.override_approved'").fetchone()[0] == 1


def test_uploaded_resident_photo_is_stored_and_served_safely(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("Photo", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Photo'").fetchone()[0]
    from PIL import Image
    image = Image.new("RGB", (8, 8), (30, 120, 180))
    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")
    png = image_buffer.getvalue()
    response = client.post("/profile/" + str(rid), data={
        "birth_place":"", "birth_date":"", "phone":"",
        "photo":(io.BytesIO(png),"portrait.png")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        path = conn.execute("SELECT photo_path FROM residents WHERE id=?", (rid,)).fetchone()[0]
    assert path.startswith("/photos/resident_")
    photo_response = client.get(path)
    assert photo_response.status_code == 200
    assert photo_response.data.startswith(b"\x89PNG")

def test_assignment_resident_choices_are_sorted_by_year_priority(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("FirstYear", "Resident"))
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,3)", ("ThirdYear", "Resident"))
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,4)", ("FourthYear", "Resident"))
    response = client.get("/assign")
    html = response.data.decode("utf-8")
    assert html.index("FourthYear") < html.index("ThirdYear") < html.index("FirstYear")


def test_academic_year_change_creates_four_new_periods_and_preserves_old_periods(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        old_year = conn.execute("SELECT year_label FROM periods ORDER BY id LIMIT 1").fetchone()[0]
    response = client.post("/settings", data={
        "lang":"fr","year":"2027-2028","director":"Dr Dris","new_password":""
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM periods WHERE year_label='2027-2028'").fetchone()[0] == 4
        assert conn.execute("SELECT COUNT(*) FROM periods WHERE year_label=?", (old_year,)).fetchone()[0] == 4


def test_xlsx_import_preview(tmp_path, monkeypatch):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Nom","Prénom","Année","Téléphone"])
    sheet.append(["ExcelResident","ExcelFirst",2,"0551111111"])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    preview = client.post("/import", data={"file":(stream,"residents.xlsx")},
                          content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200
    assert b"ExcelResident" in preview.data, preview.data.decode("utf-8", errors="replace")

def test_docx_import_reads_table_headers_and_resident_rows(tmp_path, monkeypatch):
    from docx import Document
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Nom"
    table.rows[0].cells[1].text = "Prénom"
    table.rows[0].cells[2].text = "Année"
    row = table.add_row().cells
    row[0].text = "WordResident"
    row[1].text = "WordFirst"
    row[2].text = "4"
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    preview = client.post("/import", data={"file":(stream,"residents.docx")},
                          content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200
    assert b"WordResident" in preview.data


def test_csv_export_includes_exam_objective_archive_and_audit_tables_but_not_users(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.get("/export")
    assert response.status_code == 200
    content = response.get_data().decode("utf-8-sig")
    for table in ("exam_records", "objectives", "dems_archive", "leave_years", "leave_entries", "external_rotations", "promotion_votes", "dems_votes", "audit"):
        assert table in content
    assert "users" not in content
    assert "password_hash" not in content

def test_fresh_database_requires_admin_to_set_a_private_password(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    monkeypatch.delenv("ROTATION_INITIAL_PASSWORD", raising=False)
    page = client.get("/login")
    assert page.status_code == 200
    assert b"confirm_password" in page.data
    response = client.post("/login", data={
        "username":"admin",
        "password":"fresh-private-password-2026",
        "confirm_password":"fresh-private-password-2026"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        stored = conn.execute("SELECT value FROM settings WHERE key='password'").fetchone()[0]
        user = conn.execute("SELECT role,password_hash FROM users WHERE username='admin'").fetchone()
    assert stored.startswith("pbkdf2_sha256$")
    assert user["role"] == "super_admin"
    assert user["password_hash"].startswith("pbkdf2_sha256$")


def test_legacy_password_hash_is_upgraded_after_successful_login(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    client.get("/login")
    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            conn.execute("UPDATE settings SET value=? WHERE key='password'",
                         (hashlib.sha256(b"legacy-test-password-2026").hexdigest(),))
    response = client.post("/login", data={
        "username":"admin","password":"legacy-test-password-2026"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        stored = conn.execute("SELECT value FROM settings WHERE key='password'").fetchone()[0]
        user = conn.execute("SELECT password_hash,role FROM users WHERE username='admin'").fetchone()
    assert stored.startswith("pbkdf2_sha256$")
    assert user["password_hash"].startswith("pbkdf2_sha256$")
    assert user["role"] == "super_admin"

def test_terrain_deactivation_preserves_history_and_hides_it_from_new_assignments(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("TerrainHistory", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='TerrainHistory'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        conn.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                     (rid,tid,"2025-2026","01 Mai - 30 Juillet","2025-07-30"))
    response = client.post("/terrains", data={"action":"deactivate","tid":str(tid)}, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT active FROM terrains WHERE id=?", (tid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=? AND terrain_id=?", (rid,tid)).fetchone()[0] == 1
    assert b"Urologie" not in client.get("/assign").data


def test_assignment_cancellation_is_audited_and_does_not_delete_record(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,3)", ("Cancel", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Cancel'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        pid = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        conn.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,'confirmed')",(rid,tid,pid))
        aid = conn.execute("SELECT id FROM assignments WHERE resident_id=?", (rid,)).fetchone()[0]
    response = client.post("/assignments/" + str(aid) + "/cancel", follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM assignments WHERE id=?", (aid,)).fetchone()[0] == "cancelled"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.cancelled'").fetchone()[0] == 1


def test_user_deactivation_is_audited_and_blocks_login(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    client.post("/users", data={"username":"temporary-user","password":"temporary-user-password-2026","role":"user"})
    with rotation_app.db() as conn:
        uid = conn.execute("SELECT id FROM users WHERE username='temporary-user'").fetchone()[0]
    client.post("/users", data={"action":"deactivate","user_id":str(uid)})
    with rotation_app.db() as conn:
        assert conn.execute("SELECT active FROM users WHERE id=?", (uid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='user.deactivated'").fetchone()[0] == 1
    client.get("/logout")
    response = client.post("/login", data={"username":"temporary-user","password":"temporary-user-password-2026"}, follow_redirects=True)
    assert b"Temporary user" not in response.data
    assert client.get("/settings").status_code == 302


def test_email_registration_requires_verification_code(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    sent = {}
    monkeypatch.setattr(rotation_app, "send_verification_email", lambda recipient, code: sent.update(email=recipient, code=code))
    response = client.get("/register")
    assert response.status_code == 200
    assert b'type="email"' in response.data

    response = client.post("/register", data={
        "username": "emailuser",
        "email": "emailuser@example.com",
        "password": "strong-test-password-2026",
        "confirm_password": "strong-test-password-2026",
    })
    assert response.status_code == 302
    assert response.headers["Location"] == "/verify-email"
    assert sent["email"] == "emailuser@example.com"
    with rotation_app.db() as conn:
        assert conn.execute("SELECT 1 FROM users WHERE username='emailuser'").fetchone() is None
        pending = conn.execute("SELECT attempts FROM pending_registrations WHERE email=?", (sent["email"],)).fetchone()
        assert pending is not None
        assert pending["attempts"] == 0

    wrong = client.post("/verify-email", data={"code": "000000"})
    assert wrong.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT 1 FROM users WHERE username='emailuser'").fetchone() is None

    verified = client.post("/verify-email", data={"code": sent["code"]})
    assert verified.status_code == 302
    assert verified.headers["Location"] == "/login"
    with rotation_app.db() as conn:
        account = conn.execute("SELECT username,email,email_verified,role FROM users WHERE username='emailuser'").fetchone()
        assert tuple(account) == ("emailuser", "emailuser@example.com", 1, "user")
        assert conn.execute("SELECT 1 FROM pending_registrations WHERE email=?", (sent["email"],)).fetchone() is None


def test_annual_leave_records_days_and_carries_unused_balance(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Leave", "Test", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Leave'").fetchone()[0]
        conn.commit()
    from datetime import date, timedelta
    start = date.today().replace(month=1, day=1)
    end = start + timedelta(days=4)
    response = client.post("/leave", data={
        "resident_id": str(rid), "start_date": start.isoformat(),
        "end_date": end.isoformat(), "notes": "test leave"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        entry = conn.execute("SELECT days FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()
        assert entry[0] == 5
        earned = conn.execute("SELECT SUM(earned_days) FROM leave_years WHERE resident_id=?", (rid,)).fetchone()[0]
        used = conn.execute("SELECT SUM(days) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0]
        assert earned >= 30
        assert earned - used >= 25


def test_external_rotation_blocks_overlapping_local_period(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Away", "Resident", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Away'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods WHERE name LIKE '%Août%' OR name LIKE '%Aout%'").fetchone()
        if period is None:
            period = conn.execute("SELECT id,year_label,name FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Test", "Paris", "CHU Test", "France", "Observation", "2026-08-01", "2026-10-30", "active"))
        conn.commit()
        allowed, message, status = rotation_app.check(conn, rid, tid, period["id"], period["year_label"])
        assert allowed is False
        assert "External rotation" in message


def test_external_rotation_form_requires_all_destination_details(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Trip", "Resident", 3, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Trip'").fetchone()[0]
        conn.commit()
    response = client.post("/external-rotations", data={
        "resident_id": str(rid), "supervisor": "Dr Test", "destination": "Paris",
        "hospital": "CHU Test", "country": "France", "rotation_type": "Observation",
        "start_date": "2026-11-01", "end_date": "2026-12-01", "notes": "test"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT supervisor,hospital,country,rotation_type FROM external_rotations WHERE resident_id=?", (rid,)).fetchone()
        assert tuple(row) == ("Dr Test", "CHU Test", "France", "Observation")

def test_leave_carries_unused_days_forward_and_blocks_overdraw(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Carry", "Leave", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Carry'").fetchone()[0]
        conn.commit()

    first = client.post("/leave", data={"resident_id": rid, "start_date": "2024-01-01", "end_date": "2024-01-10"}, follow_redirects=True)
    assert first.status_code == 200
    second = client.post("/leave", data={"resident_id": rid, "start_date": "2025-01-01", "end_date": "2025-01-25"}, follow_redirects=True)
    assert second.status_code == 200
    with rotation_app.db() as conn:
        years = conn.execute("SELECT SUM(earned_days) FROM leave_years WHERE resident_id=? AND year IN (2024,2025)", (rid,)).fetchone()[0]
        used = conn.execute("SELECT SUM(days) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0]
    assert years == 60
    assert used == 35
    rejected = client.post("/leave", data={"resident_id": rid, "start_date": "2025-02-01", "end_date": "2025-02-26"}, follow_redirects=True)
    assert rejected.status_code == 200
    assert b"Insufficient accrued leave balance" in rejected.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT SUM(days) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0] == 35


def test_leave_rejects_overlapping_entries(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Overlap", "Leave", 1, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Overlap'").fetchone()[0]
        conn.commit()
    client.post("/leave", data={"resident_id": rid, "start_date": "2026-01-01", "end_date": "2026-01-10"}, follow_redirects=True)
    rejected = client.post("/leave", data={"resident_id": rid, "start_date": "2026-01-05", "end_date": "2026-01-12"}, follow_redirects=True)
    assert rejected.status_code == 200
    assert b"Leave dates overlap an existing entry" in rejected.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0] == 1


def test_external_rotation_returned_record_still_blocks_overlapping_period(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Returned", "Away", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Returned'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods WHERE name LIKE '%Août%' OR name LIKE '%Aout%'").fetchone()
        assert period is not None
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Test", "Paris", "CHU Test", "France", "Observation", "2026-08-01", "2026-10-30", "returned"))
        conn.commit()
        allowed, message, _ = rotation_app.check(conn, rid, tid, period["id"], period["year_label"])
    assert allowed is False
    assert "External rotation" in message


def test_assignment_resident_options_mark_external_period_conflicts(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Blocked", "Resident", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Blocked'").fetchone()[0]
        period = conn.execute("SELECT id FROM periods WHERE name LIKE '%Août%' OR name LIKE '%Aout%'").fetchone()
        assert period is not None
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Test", "Paris", "CHU Test", "France", "Observation", "2026-08-01", "2026-10-30", "returned"))
        conn.commit()
    response = client.get("/assign")
    assert response.status_code == 200
    assert f'data-blocked-periods="{period["id"]}"'.encode() in response.data
    assert b"o.disabled=blocked" in response.data


def test_external_rotation_can_be_edited_and_cancelled_with_audit(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Edit", "External", 3, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Edit'").fetchone()[0]
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Old", "Paris", "CHU Old", "France", "Observation", "2026-11-01", "2026-12-01", "active"))
        xid = conn.execute("SELECT id FROM external_rotations WHERE resident_id=?", (rid,)).fetchone()[0]
        conn.commit()
    edit = client.post(f"/external-rotations/{xid}/edit", data={
        "resident_id": str(rid), "supervisor": "Dr New", "destination": "Lyon",
        "hospital": "CHU New", "country": "France", "rotation_type": "Clinical",
        "start_date": "2026-11-10", "end_date": "2026-12-10", "notes": "updated"
    }, follow_redirects=True)
    assert edit.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT supervisor,destination,hospital,rotation_type,status FROM external_rotations WHERE id=?", (xid,)).fetchone()
        assert tuple(row) == ("Dr New", "Lyon", "CHU New", "Clinical", "active")
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='external_rotation.updated'").fetchone()[0] == 1
    cancelled = client.post(f"/external-rotations/{xid}/cancel", follow_redirects=True)
    assert cancelled.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM external_rotations WHERE id=?", (xid,)).fetchone()[0] == "cancelled"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='external_rotation.cancelled'").fetchone()[0] == 1

def test_post_without_csrf_token_is_rejected(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    client.get("/login")
    response = client.open("/login", method="POST", data={"username": "admin", "password": "not-a-real-password"})
    assert response.status_code == 400
    assert b"Invalid or missing CSRF token" in response.data


def test_official_admin_accounts_must_be_super_admin_and_cannot_be_disabled(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        invalid = client.post("/users", data={
            "username": username, "password": "strong-test-password-2026", "role": "user"
        }, follow_redirects=True)
        assert invalid.status_code == 200
        assert b"Official accounts must be Super Admins" in invalid.data
        with rotation_app.db() as conn:
            assert conn.execute("SELECT 1 FROM users WHERE lower(username)=?", (username,)).fetchone() is None

        created = client.post("/users", data={
            "username": username, "password": "strong-test-password-2026", "role": "super_admin"
        }, follow_redirects=True)
        assert created.status_code == 200
        with rotation_app.db() as conn:
            row = conn.execute("SELECT id,role,active FROM users WHERE lower(username)=?", (username,)).fetchone()
            assert tuple(row[1:]) == ("super_admin", 1)
            uid = row["id"]
        deactivated = client.post("/users", data={"action": "deactivate", "user_id": str(uid)}, follow_redirects=True)
        assert deactivated.status_code == 200
        assert b"Official accounts cannot be deactivated" in deactivated.data
        with rotation_app.db() as conn:
            assert conn.execute("SELECT active FROM users WHERE id=?", (uid,)).fetchone()[0] == 1

def test_year_progression_requires_approval_from_both_official_admins(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={"username": username, "password": "strong-test-password-2026", "role": "super_admin"}, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,2,'active')", ("Promotion", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Promotion'").fetchone()[0]
        conn.commit()

    unauthorized = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"2","outcome":"approve","reason":"ok"}, follow_redirects=False)
    assert unauthorized.status_code == 403

    with client.session_transaction() as sess:
        sess["user"] = "benhocine"
        sess["role"] = "super_admin"
    first = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"2","outcome":"approve","reason":"First official approves"}, follow_redirects=True)
    assert first.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM promotion_votes WHERE resident_id=?", (rid,)).fetchone()[0] == 1

    with client.session_transaction() as sess:
        sess["user"] = "drdris"
        sess["role"] = "super_admin"
    second = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"2","outcome":"approve","reason":"Second official approves"}, follow_redirects=True)
    assert second.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 3
        assert conn.execute("SELECT COUNT(*) FROM promotion_votes WHERE resident_id=?", (rid,)).fetchone()[0] == 0
        decision = conn.execute("SELECT decision,actor FROM decisions WHERE details LIKE '%Promotion Resident — year 2%'").fetchone()
        assert decision["decision"] == "Approved by both official administrators"
        assert decision["actor"] == "benhocine+drdris"


def test_year_progression_is_not_approved_if_either_official_rejects(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={"username": username, "password": "strong-test-password-2026", "role": "super_admin"}, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,3,'active')", ("Reject", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Reject'").fetchone()[0]
        conn.commit()
    for username,outcome in (("benhocine","approve"),("drdris","reject")):
        with client.session_transaction() as sess:
            sess["user"] = username
            sess["role"] = "super_admin"
        response = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"3","outcome":outcome,"reason":username+" reviewed"}, follow_redirects=True)
        assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 3
        decision = conn.execute("SELECT decision FROM decisions WHERE details LIKE '%Reject Resident — year 3%'").fetchone()
        assert decision[0] == "Rejected by official vote"



def test_dems_archive_is_blocked_if_either_official_rejects(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={
            "username": username,
            "password": "strong-test-password-2026",
            "role": "super_admin",
        }, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,4,'active')",
            ("DEMSReject", "Resident"),
        )
        rid = conn.execute(
            "SELECT id FROM residents WHERE surname='DEMSReject'"
        ).fetchone()[0]
        conn.commit()

    for username, outcome in (("benhocine", "approve"), ("drdris", "reject")):
        with client.session_transaction() as sess:
            sess["user"] = username
            sess["role"] = "super_admin"
        response = client.post("/dems", data={
            "resident": str(rid),
            "success_date": "2026-10-15",
            "outcome": outcome,
            "reason": username + " reviewed the DEMS result",
        }, follow_redirects=True)
        assert response.status_code == 200

    with rotation_app.db() as conn:
        resident = conn.execute(
            "SELECT status,dems_status FROM residents WHERE id=?", (rid,)
        ).fetchone()
        assert tuple(resident) == ("active", "")
        assert conn.execute(
            "SELECT COUNT(*) FROM dems_archive WHERE resident_id=?", (rid,)
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM dems_votes WHERE resident_id=?", (rid,)
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action='dems.vote_rejected'"
        ).fetchone()[0] == 1
        decision = conn.execute(
            "SELECT decision FROM decisions WHERE details LIKE '%DEMSReject Resident — DEMS%'"
        ).fetchone()
        assert decision is not None
        assert decision["decision"] == "Not approved by official vote"


def test_assignment_page_has_priority_filters_and_clear_conflict_feedback(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Priority", "Four", 4, "active"))
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Search", "One", 1, "active"))
        conn.commit()
    response = client.get("/assign")
    assert response.status_code == 200
    assert b'id="assignment-search"' in response.data
    assert b'id="assignment-year-filter"' in response.data
    assert b'value="4"' in response.data and ("السنة".encode() in response.data or b'Year 4' in response.data)
    assert "تخطيط التربصات وتعيين المقيمين".encode() in response.data or b'Rotation planning &amp; resident assignment' in response.data or b'Rotation planning & resident assignment' in response.data
    assert b'rt-assignment-stats' in response.data
    assert b'rt-assignment-notice' in response.data
    assert b'External rotation conflict' in response.data
    assert b'rt-assignment-table' in response.data

def test_inactive_core_terrain_does_not_block_repeat_eligibility(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,4)",
            ("InactiveCore", "Resident"),
        )
        rid = conn.execute(
            "SELECT id FROM residents WHERE surname='InactiveCore'"
        ).fetchone()[0]
        target = conn.execute(
            "SELECT id FROM terrains WHERE name='Réanimation chirurgicale'"
        ).fetchone()[0]
        inactive = conn.execute(
            "SELECT id FROM terrains WHERE name='Urologie'"
        ).fetchone()[0]
        period_id = conn.execute(
            "SELECT id FROM periods ORDER BY id LIMIT 1"
        ).fetchone()[0]
        year_label = conn.execute(
            "SELECT year_label FROM periods WHERE id=?", (period_id,)
        ).fetchone()[0]
        conn.execute("UPDATE terrains SET active=0 WHERE id=?", (inactive,))
        conn.execute(
            "INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,2)",
            (year_label, target, period_id),
        )
        conn.execute(
            "INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
            (rid, target, year_label, "01 Fév - 30 Avr", "2026-04-30"),
        )
        conn.execute(
            """
            INSERT INTO history(resident_id,terrain_id,year_label,period,date_done)
            SELECT ?, id, ?, 'historique', '2026-04-30'
            FROM terrains
            WHERE active=1 AND kind IN ('mandatory','basic') AND id<>?
            """,
            (rid, year_label, target),
        )

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(
                conn, rid, target, period_id, year_label
            )

    assert allowed is True
    assert status == "confirmed"
    assert "يمكن التكرار" in message

def test_pending_assignment_cannot_be_approved_after_terrain_deactivation(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("StaleTerrain", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='StaleTerrain'").fetchone()[0]
        conn.execute("INSERT INTO terrains(name,kind,mandatory_year) VALUES(?,?,3)", ("Pending Deactivated Terrain", "mandatory"))
        tid = conn.execute("SELECT id FROM terrains WHERE name='Pending Deactivated Terrain'").fetchone()[0]
        period = conn.execute("SELECT id,year_label FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)", (period["year_label"], tid, period["id"]))
    client.post("/assign", data={"resident":str(rid),"terrain":str(tid),"period":str(period["id"])})
    with rotation_app.db() as conn:
        aid = conn.execute("SELECT id FROM assignments WHERE resident_id=? AND status='pending'", (rid,)).fetchone()[0]
        conn.execute("UPDATE terrains SET active=0 WHERE id=?", (tid,))
    response = client.post("/decisions", data={"aid":str(aid),"outcome":"approve","reason":"Review stale request"}, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM assignments WHERE id=?", (aid,)).fetchone()[0] == "pending"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.override_approved'").fetchone()[0] == 0


def test_pending_assignment_cannot_be_approved_during_new_external_overlap(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("StaleExternal", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='StaleExternal'").fetchone()[0]
        conn.execute("INSERT INTO terrains(name,kind,mandatory_year) VALUES(?,?,3)", ("Pending External Terrain", "mandatory"))
        tid = conn.execute("SELECT id FROM terrains WHERE name='Pending External Terrain'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)", (period["year_label"], tid, period["id"]))
    client.post("/assign", data={"resident":str(rid),"terrain":str(tid),"period":str(period["id"])})
    start_day, end_day = rotation_app.rotation_period_bounds(period["year_label"], period["name"])
    with rotation_app.db() as conn:
        aid = conn.execute("SELECT id FROM assignments WHERE resident_id=? AND status='pending'", (rid,)).fetchone()[0]
        conn.execute(
            "INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
            (rid,"Dr Test","Paris","CHU Test","France","Observation",start_day.isoformat(),end_day.isoformat(),"active"),
        )
    response = client.post("/decisions", data={"aid":str(aid),"outcome":"approve","reason":"Review external conflict"}, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM assignments WHERE id=?", (aid,)).fetchone()[0] == "pending"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.override_approved'").fetchone()[0] == 0

def test_legacy_terrain_classification_migrates_once_without_rewriting_later_admin_edits(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    legacy = sqlite3.connect(rotation_app.DB)
    legacy.execute(
        "CREATE TABLE terrains(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL,kind TEXT DEFAULT 'basic',mandatory_year INTEGER DEFAULT 0,active INTEGER NOT NULL DEFAULT 1)"
    )
    legacy.executemany(
        "INSERT INTO terrains(name,kind) VALUES(?,?)",
        [("Sédation", "basic"), ("Maxillo-faciale", "basic"),
         ("Chirurgie cardiaque", "basic"), ("Brûlés", "basic"),
         ("CAC DBK", "basic"), ("Gynécologie", "optional")],
    )
    legacy.commit()
    legacy.close()

    assert client.get("/login").status_code == 200
    with rotation_app.db() as conn:
        for name in ("Sédation", "Maxillo-faciale", "Chirurgie cardiaque", "Brûlés", "CAC DBK"):
            assert conn.execute("SELECT kind FROM terrains WHERE name=?", (name,)).fetchone()[0] == "optional"
        assert conn.execute("SELECT kind FROM terrains WHERE name='Gynécologie'").fetchone()[0] == "basic"
        assert conn.execute(
            "SELECT value FROM settings WHERE key='terrain_classification_migration_20261017'"
        ).fetchone()[0] == "1"
        conn.execute("UPDATE terrains SET kind='mandatory' WHERE name='Sédation'")
        conn.commit()

    monkeypatch.setattr(rotation_app.app, "_ready", False, raising=False)
    assert client.get("/login").status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT kind FROM terrains WHERE name='Sédation'").fetchone()[0] == "mandatory"

def test_settings_rejects_password_shorter_than_twelve_characters(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        before = conn.execute(
            "SELECT password_hash FROM users WHERE username='admin'"
        ).fetchone()[0]
        year = conn.execute("SELECT value FROM settings WHERE key='year'").fetchone()[0]

    response = client.post(
        "/settings",
        data={"lang":"ar","year":year,"director":"","new_password":"weak"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "at least 12 characters".encode() in response.data
    with rotation_app.db() as conn:
        after = conn.execute(
            "SELECT password_hash FROM users WHERE username='admin'"
        ).fetchone()[0]
        assert after == before

def test_external_rotation_creation_rejects_overlap_with_existing_local_assignment(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("LocalFirst", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='LocalFirst'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,'confirmed')", (rid,tid,period["id"]))
        conn.commit()
    start_day, end_day = rotation_app.rotation_period_bounds(period["year_label"], period["name"])
    response = client.post("/external-rotations", data={
        "resident_id":str(rid),"supervisor":"Dr Test","destination":"Paris",
        "hospital":"CHU Test","country":"France","rotation_type":"Observation",
        "start_date":start_day.isoformat(),"end_date":end_day.isoformat(),"notes":"overlap"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b"overlaps an existing local assignment" in response.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM external_rotations WHERE resident_id=?", (rid,)).fetchone()[0] == 0


def test_external_rotation_edit_rejects_overlap_with_existing_local_assignment(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("LocalEdit", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='LocalEdit'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,'confirmed')", (rid,tid,period["id"]))
        conn.commit()
    start_day, end_day = rotation_app.rotation_period_bounds(period["year_label"], period["name"])
    nonoverlap_start = end_day + timedelta(days=1)
    nonoverlap_end = nonoverlap_start + timedelta(days=7)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
            (rid,"Dr Test","Paris","CHU Test","France","Observation",nonoverlap_start.isoformat(),nonoverlap_end.isoformat(),"returned"),
        )
        xid = conn.execute("SELECT id FROM external_rotations WHERE resident_id=?", (rid,)).fetchone()[0]
    response = client.post(f"/external-rotations/{xid}/edit", data={
        "resident_id":str(rid),"supervisor":"Dr Test","destination":"Paris",
        "hospital":"CHU Test","country":"France","rotation_type":"Observation",
        "start_date":start_day.isoformat(),"end_date":end_day.isoformat(),"notes":"overlap"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b"overlaps an existing local assignment" in response.data
    with rotation_app.db() as conn:
        row = conn.execute("SELECT start_date,end_date FROM external_rotations WHERE id=?", (xid,)).fetchone()
        assert tuple(row) == (nonoverlap_start.isoformat(), nonoverlap_end.isoformat())

def test_resident_edit_persists_and_cannot_bypass_or_reverse_dems_graduation(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,2,'active')", ("Editable", "Resident"))
        editable = conn.execute("SELECT id FROM residents WHERE surname='Editable'").fetchone()[0]
        conn.execute("INSERT INTO residents(surname,firstname,year,status,dems_status) VALUES(?,?,4,'graduated','success')", ("Graduated", "Resident"))
        graduated = conn.execute("SELECT id FROM residents WHERE surname='Graduated'").fetchone()[0]
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,4,'active')", ("Attempt", "Bypass"))
        bypass = conn.execute("SELECT id FROM residents WHERE surname='Attempt'").fetchone()[0]

    saved = client.post("/residents", data={
        "action":"edit","rid":str(editable),"year":"2","repeat_year":"0","status":"inactive"
    }, follow_redirects=True)
    assert saved.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT year,status FROM residents WHERE id=?", (editable,)).fetchone()
        assert tuple(row) == (2,"inactive")

    year_change = client.post("/residents", data={
        "action":"edit","rid":str(editable),"year":"3","repeat_year":"0","status":"active"
    }, follow_redirects=True)
    assert year_change.status_code == 200
    assert b"Study-year changes must use the official exam/promotion workflow" in year_change.data
    with rotation_app.db() as conn:
        row = conn.execute("SELECT year,status FROM residents WHERE id=?", (editable,)).fetchone()
        assert tuple(row) == (2,"inactive")

    bypass_response = client.post("/residents", data={
        "action":"edit","rid":str(bypass),"year":"4","repeat_year":"0","status":"graduated"
    }, follow_redirects=True)
    assert bypass_response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM residents WHERE id=?", (bypass,)).fetchone()[0] == "active"

    restore_response = client.post("/residents", data={
        "action":"edit","rid":str(graduated),"year":"4","repeat_year":"0","status":"active"
    }, follow_redirects=True)
    assert restore_response.status_code == 200
    assert b"Graduated residents are archived" in restore_response.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM residents WHERE id=?", (graduated,)).fetchone()[0] == "graduated"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='resident.updated'").fetchone()[0] == 1

def test_first_year_evaluation_does_not_record_decision_for_resident_in_later_year(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,2,'active')", ("LaterYear", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='LaterYear'").fetchone()[0]
        conn.execute(
            "INSERT INTO exam_records(resident_id,study_year,module,theory,practical) VALUES(?,1,?,?,?)",
            (rid,"Module A",12,12),
        )
        conn.commit()

    response = client.post("/exams", data={
        "action":"evaluate_first_year","resident":str(rid),"study_year":"1",
        "all_modules_confirmed":"yes"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b"Only current first-year residents" in response.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM decisions WHERE details LIKE ?", (f"%LaterYear Resident — first-year exams%",)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='exam.first_year_evaluated'").fetchone()[0] == 0

def test_startup_fills_missing_periods_in_partial_legacy_academic_year(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    legacy = sqlite3.connect(rotation_app.DB)
    legacy.execute(
        "CREATE TABLE periods(id INTEGER PRIMARY KEY,year_label TEXT NOT NULL,name TEXT NOT NULL,UNIQUE(year_label,name))"
    )
    legacy.execute("INSERT INTO periods(year_label,name) VALUES(?,?)", ("2026-2027","01 Fév - 30 Avr"))
    legacy.execute("CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT NOT NULL)")
    legacy.execute("INSERT INTO settings(key,value) VALUES('year','2026-2027')")
    legacy.commit()
    legacy.close()

    assert client.get("/login").status_code == 200
    expected = {"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
    with rotation_app.db() as conn:
        names = {row[0] for row in conn.execute("SELECT name FROM periods WHERE year_label='2026-2027'")}
        assert names == expected
        assert conn.execute("SELECT COUNT(*) FROM periods WHERE year_label='2026-2027'").fetchone()[0] == 4

    monkeypatch.setattr(rotation_app.app, "_ready", False, raising=False)
    assert client.get("/login").status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM periods WHERE year_label='2026-2027'").fetchone()[0] == 4

def test_import_rejects_fractional_study_year_instead_of_truncating(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    source = b"surname,firstname,year,phone\nFractional,Resident,2.5,0550000000\n"
    response = client.post(
        "/import",
        data={"file":(io.BytesIO(source),"fractional-year.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Year must be 1, 2, 3 or 4." in response.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM residents WHERE surname='Fractional'").fetchone()[0] == 0

def test_deactivated_account_loses_existing_session_immediately(tmp_path, monkeypatch):
    admin = setup_client(tmp_path, monkeypatch)
    login(admin)
    username = "revoked-session"
    password = "strong-revoked-password-2026"
    admin.post("/users", data={"username":username,"password":password,"role":"user"}, follow_redirects=True)

    user_client = rotation_app.app.test_client()
    signed_in = user_client.post(
        "/login", data={"username":username,"password":password}, follow_redirects=True
    )
    assert signed_in.status_code == 200
    assert user_client.get("/").status_code == 200

    with rotation_app.db() as conn:
        uid = conn.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()[0]
    admin.post("/users", data={"action":"deactivate","user_id":str(uid)}, follow_redirects=True)

    revoked = user_client.get("/")
    assert revoked.status_code == 302
    assert "/login" in revoked.headers["Location"]
    with user_client.session_transaction() as session:
        assert "user" not in session

def test_standard_user_cannot_create_rotation_assignments(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    client.post("/users", data={
        "username":"assignment-operator",
        "password":"assignment-operator-password-2026",
        "role":"user",
    }, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("NoAssign", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='NoAssign'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute(
            "INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)",
            (period["year_label"],tid,period["id"]),
        )
        conn.commit()
    with client.session_transaction() as session:
        session["user"]="assignment-operator"
        session["role"]="user"
        session["lang"]="en"

    read_only = client.get("/assign")
    assert read_only.status_code == 200
    assert b"Only Super Admins can create assignments." in read_only.data
    assert b'id="assignment-form"' not in read_only.data

    response = client.post("/assign", data={
        "resident":str(rid),"terrain":str(tid),"period":str(period["id"])
    })
    assert response.status_code == 403
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM assignments WHERE resident_id=?", (rid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.created'").fetchone()[0] == 0

def test_history_rejects_invalid_year_period_and_duplicate_records(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("History", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='History'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        conn.commit()

    invalid_year = client.post("/history", data={
        "resident":str(rid),"terrain":str(tid),"year_label":"2025-2027",
        "period":"01 Fév - 30 Avr","date_done":"2026-03-01"
    }, follow_redirects=True)
    assert invalid_year.status_code == 200
    assert b"Academic year must be YYYY-YYYY" in invalid_year.data

    invalid_period = client.post("/history", data={
        "resident":str(rid),"terrain":str(tid),"year_label":"2025-2026",
        "period":"invented period","date_done":"2026-03-01"
    }, follow_redirects=True)
    assert invalid_period.status_code == 200
    assert b"Choose one of the four official rotation periods" in invalid_period.data

    future_date = client.post("/history", data={
        "resident":str(rid),"terrain":str(tid),"year_label":"2098-2099",
        "period":"01 Fév - 30 Avr","date_done":"2098-03-01"
    }, follow_redirects=True)
    assert future_date.status_code == 200
    assert b"A historical rotation date cannot be in the future" in future_date.data

    payload = {
        "resident":str(rid),"terrain":str(tid),"year_label":"2025-2026",
        "period":"01 Fév - 30 Avr","date_done":"2026-03-01"
    }
    first = client.post("/history", data=payload, follow_redirects=True)
    assert first.status_code == 200
    second = client.post("/history", data=payload, follow_redirects=True)
    assert second.status_code == 200
    assert b"This historical rotation already exists" in second.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=? AND terrain_id=?", (rid,tid)).fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='history.created'").fetchone()[0] == 1

def test_inactive_resident_cannot_receive_promotion_votes(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    client.post("/users", data={
        "username":"benhocine",
        "password":"official-admin-password-2026",
        "role":"super_admin",
    }, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,2,'inactive')", ("Inactive", "Promotion"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Inactive'").fetchone()[0]
        conn.commit()

    with client.session_transaction() as session:
        session["user"]="benhocine"
        session["role"]="super_admin"
    first = client.post("/exams", data={
        "action":"approve_advance","resident":str(rid),"study_year":"2",
        "outcome":"approve","reason":"Should not count"
    }, follow_redirects=True)
    assert first.status_code == 200
    assert b"Inactive residents cannot be advanced" in first.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM promotion_votes WHERE resident_id=?", (rid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='promotion.vote_recorded'").fetchone()[0] == 0

def test_future_confirmed_rotations_do_not_count_as_completed_for_repeat_eligibility(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        rid = conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,4) RETURNING id",
            ("Future", "Rotations"),
        ).fetchone()[0]
        current_period = conn.execute("SELECT * FROM periods ORDER BY id LIMIT 1").fetchone()
        target = conn.execute(
            "SELECT id FROM terrains WHERE name='Réanimation chirurgicale'"
        ).fetchone()[0]
        conn.execute(
            "INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,2)",
            (current_period["year_label"],target,current_period["id"]),
        )
        conn.execute(
            "INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
            (rid,target,"2025-2026","01 Fév - 30 Avr","2026-04-30"),
        )
        future_terrains = conn.execute(
            "SELECT id FROM terrains WHERE active=1 AND kind IN ('mandatory','basic') AND id<>?",
            (target,),
        ).fetchall()
        for index, terrain_row in enumerate(future_terrains):
            start_year = 2099 + index
            future_period_id = conn.execute(
                "INSERT INTO periods(year_label,name) VALUES(?,?) RETURNING id",
                (f"{start_year}-{start_year+1}","01 Fév - 30 Avr"),
            ).fetchone()[0]
            conn.execute(
                "INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,'confirmed')",
                (rid,terrain_row["id"],future_period_id),
            )
        conn.commit()

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(
                conn,rid,target,current_period["id"],current_period["year_label"]
            )
    assert allowed is False
    assert "التربص منجز سابقاً" in message

def test_capacity_rejects_negative_and_cross_academic_year_periods(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        terrain_id = conn.execute(
            "SELECT id FROM terrains WHERE active=1 ORDER BY id LIMIT 1"
        ).fetchone()[0]
        current = conn.execute("SELECT id,year_label FROM periods ORDER BY id LIMIT 1").fetchone()
        other_period_id = conn.execute(
            "INSERT INTO periods(year_label,name) VALUES(?,?) RETURNING id",
            ("2099-2100","01 Fév - 30 Avr"),
        ).fetchone()[0]
        conn.commit()

    negative = client.post("/capacities", data={
        "terrain":str(terrain_id),"period":str(current["id"]),"places":"-1"
    }, follow_redirects=True)
    assert negative.status_code == 200
    assert b"Invalid capacity input" in negative.data

    wrong_year = client.post("/capacities", data={
        "terrain":str(terrain_id),"period":str(other_period_id),"places":"3"
    }, follow_redirects=True)
    assert wrong_year.status_code == 200
    assert b"Invalid capacity input" in wrong_year.data

    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM capacities WHERE terrain_id=? AND period_id=?",
            (terrain_id,other_period_id),
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM capacities WHERE terrain_id=? AND period_id=? AND places<0",
            (terrain_id,current["id"]),
        ).fetchone()[0] == 0

def test_import_confirmation_rechecks_duplicates_and_imports_nothing_on_stale_preview(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    source = (
        b"surname,firstname,year,phone\n"
        b"Fresh,Resident,2,0550000001\n"
        b"Race,Condition,3,0550000002\n"
    )
    preview = client.post("/import", data={
        "file": (io.BytesIO(source), "residents.csv")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200
    match = re.search(rb'name="preview_id" value="(\d+)"', preview.data)
    assert match

    # Simulate another administrator adding a resident after the preview was made.
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,3)",
            ("Race", "Condition"),
        )
        conn.commit()

    confirmed = client.post("/import", data={
        "action":"confirm", "preview_id":match.group(1).decode()
    }, follow_redirects=True)
    assert confirmed.status_code == 200
    assert b"Nothing was imported" in confirmed.data
    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM residents WHERE surname='Fresh' AND firstname='Resident'"
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM residents WHERE surname='Race' AND firstname='Condition' AND year=3"
        ).fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='residents.imported'").fetchone()[0] == 0

def test_repeat_of_core_terrain_is_blocked_across_academic_years(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        rid = conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,4) RETURNING id",
            ("CrossYear", "Repeat"),
        ).fetchone()[0]
        target = conn.execute(
            "SELECT id FROM terrains WHERE active=1 AND kind IN ('mandatory','basic') ORDER BY id LIMIT 1"
        ).fetchone()[0]
        current_year = conn.execute("SELECT value FROM settings WHERE key='year'").fetchone()[0]
        current_period = conn.execute(
            "SELECT id,year_label FROM periods WHERE year_label=? ORDER BY id LIMIT 1",
            (current_year,),
        ).fetchone()
        conn.execute(
            "INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)",
            (current_year,target,current_period["id"]),
        )
        previous_period = conn.execute(
            "INSERT INTO periods(year_label,name) VALUES(?,?) RETURNING id",
            ("2024-2025","01 Fév - 30 Avr"),
        ).fetchone()[0]
        conn.execute(
            "INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,'confirmed')",
            (rid,target,previous_period),
        )
        conn.commit()

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(
                conn,rid,target,current_period["id"],current_year
            )
    assert allowed is False
    assert status is None

def test_future_dated_history_does_not_unlock_core_terrain_repeats(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        rid = conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,4) RETURNING id",
            ("FutureHistory", "Resident"),
        ).fetchone()[0]
        current_period = conn.execute("SELECT * FROM periods ORDER BY id LIMIT 1").fetchone()
        terrains = conn.execute(
            "SELECT id FROM terrains WHERE active=1 AND kind IN ('mandatory','basic')"
        ).fetchall()
        target = terrains[0]["id"]
        conn.execute(
            "INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)",
            (current_period["year_label"],target,current_period["id"]),
        )
        for terrain in terrains:
            conn.execute(
                "INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                (rid,terrain["id"],"2098-2099","01 Fév - 30 Avr","2098-03-01"),
            )
        conn.commit()

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(
                conn,rid,target,current_period["id"],current_period["year_label"]
            )
    assert allowed is False
    assert status is None

def test_undated_future_history_does_not_unlock_core_terrain_repeats(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        rid = conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,4) RETURNING id",
            ("UndatedFuture", "History"),
        ).fetchone()[0]
        current_period = conn.execute("SELECT * FROM periods ORDER BY id LIMIT 1").fetchone()
        terrains = conn.execute(
            "SELECT id FROM terrains WHERE active=1 AND kind IN ('mandatory','basic')"
        ).fetchall()
        target = terrains[0]["id"]
        conn.execute(
            "INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)",
            (current_period["year_label"], target, current_period["id"]),
        )
        for terrain in terrains:
            conn.execute(
                "INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?, '')",
                (rid, terrain["id"], "2098-2099", "01 Fév - 30 Avr"),
            )
        conn.commit()

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(
                conn, rid, target, current_period["id"], current_period["year_label"]
            )
    assert allowed is False
    assert status is None
    assert "repetition" not in message.lower() or "completed" not in message.lower()


def test_exam_recording_rejects_future_year_and_graduated_resident(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        future_id = conn.execute(
            "INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,2,'active') RETURNING id",
            ("FutureYear", "Exam"),
        ).fetchone()[0]
        graduated_id = conn.execute(
            "INSERT INTO residents(surname,firstname,year,status,dems_status) VALUES(?, ?,4,'graduated','success') RETURNING id",
            ("Graduated", "Exam"),
        ).fetchone()[0]
        conn.commit()

    future = client.post("/exams", data={
        "action":"record","resident":str(future_id),"study_year":"3",
        "module":"Future Module","theory":"15","practical":"15"
    }, follow_redirects=True)
    assert future.status_code == 200
    assert b"Cannot record marks for a future study year" in future.data

    graduated = client.post("/exams", data={
        "action":"record","resident":str(graduated_id),"study_year":"4",
        "module":"Late Module","theory":"15","practical":"15"
    }, follow_redirects=True)
    assert graduated.status_code == 200
    assert b"Cannot record marks for a future study year" in graduated.data

    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM exam_records WHERE resident_id IN (?,?)", (future_id,graduated_id)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action IN ('exam.recorded','exam.updated')").fetchone()[0] == 0

def test_terrain_classification_cannot_change_after_rotation_records_exist(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        resident_id = conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,2) RETURNING id",
            ("Terrain", "Audit"),
        ).fetchone()[0]
        terrain = conn.execute(
            "SELECT id,kind,mandatory_year,name FROM terrains WHERE active=1 ORDER BY id LIMIT 1"
        ).fetchone()
        conn.execute(
            "INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
            (resident_id,terrain["id"],"2025-2026","01 Fév - 30 Avr","2026-04-30"),
        )
        conn.commit()

    response = client.post("/terrains", data={
        "action":"edit","tid":str(terrain["id"]),"name":terrain["name"],
        "kind":"optional" if terrain["kind"]!="optional" else "mandatory",
        "mandatory_year":str(terrain["mandatory_year"]),
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b"Terrain classification/year cannot be changed after rotation history or assignments exist" in response.data
    with rotation_app.db() as conn:
        unchanged = conn.execute("SELECT kind,mandatory_year FROM terrains WHERE id=?", (terrain["id"],)).fetchone()
        assert unchanged["kind"] == terrain["kind"]
        assert unchanged["mandatory_year"] == terrain["mandatory_year"]
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=? AND terrain_id=?", (resident_id,terrain["id"])).fetchone()[0] == 1

def test_profile_rejects_invalid_identity_and_preserves_photo_when_omitted(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        rid = conn.execute(
            "INSERT INTO residents(surname,firstname,year,birth_place,birth_date,photo_path,phone) VALUES(?,?,?,?,?,?,?) RETURNING id",
            ("Original","Resident",2,"Tizi Ouzou","1990-01-01","/photos/resident_1.png","0550000000"),
        ).fetchone()[0]
        conn.commit()
    client.get(f"/profile/{rid}")

    invalid = client.post(f"/profile/{rid}", data={
        "surname":"   ","firstname":"Resident","birth_date":"2099-01-01"
    }, follow_redirects=True)
    assert invalid.status_code == 200
    assert b"Surname and first name are required" in invalid.data
    with rotation_app.db() as conn:
        row = conn.execute("SELECT surname,photo_path FROM residents WHERE id=?",(rid,)).fetchone()
        assert row["surname"] == "Original"
        assert row["photo_path"] == "/photos/resident_1.png"

    valid = client.post(f"/profile/{rid}", data={
        "surname":"Updated","firstname":"Resident","birth_place":"Tizi Ouzou",
        "birth_date":"1990-01-01","phone":"0550000000"
    }, follow_redirects=True)
    assert valid.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT surname,photo_path FROM residents WHERE id=?",(rid,)).fetchone()
        assert row["surname"] == "Updated"
        assert row["photo_path"] == "/photos/resident_1.png"

def test_resident_creation_validates_repeat_flag_lengths_and_duplicate_names(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    invalid_flag = client.post("/residents", data={
        "surname":"Valid","firstname":"Flag","year":"1","repeat_year":"yes"
    }, follow_redirects=True)
    assert invalid_flag.status_code == 200
    assert b"repeat flag must be 0 or 1" in invalid_flag.data

    long_phone = client.post("/residents", data={
        "surname":"Valid","firstname":"Phone","year":"1","phone":"x"*41
    }, follow_redirects=True)
    assert long_phone.status_code == 200
    assert b"Phone must be at most 40 characters" in long_phone.data

    payload = {"surname":"Test","firstname":"Resident","year":"2","phone":"0550000000","repeat_year":"0"}
    first = client.post("/residents", data=payload, follow_redirects=True)
    assert first.status_code == 200
    second = client.post("/residents", data={**payload,"surname":" test "}, follow_redirects=True)
    assert second.status_code == 200
    assert b"A resident with the same name and study year already exists" in second.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM residents WHERE lower(surname)='test' AND lower(firstname)='resident' AND year=2").fetchone()[0] == 1

def test_objectives_reject_invalid_year_inactive_terrain_and_duplicates(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        terrain = conn.execute("SELECT id FROM terrains WHERE active=1 ORDER BY id LIMIT 1").fetchone()[0]
        conn.execute("INSERT INTO terrains(name,kind,active) VALUES('Inactive objective terrain','basic',0)")
        inactive = conn.execute("SELECT id FROM terrains WHERE name='Inactive objective terrain'").fetchone()[0]
        conn.commit()

    invalid_year = client.post("/objectives", data={
        "terrain":str(terrain),"study_year":"5","objective":"Objective"
    }, follow_redirects=True)
    assert invalid_year.status_code == 200
    assert b"Objective study year must be 1, 2, 3 or 4" in invalid_year.data

    inactive_terrain = client.post("/objectives", data={
        "terrain":str(inactive),"study_year":"1","objective":"Objective"
    }, follow_redirects=True)
    assert inactive_terrain.status_code == 200
    assert b"Choose an active terrain" in inactive_terrain.data

    payload = {"terrain":str(terrain),"study_year":"1","objective":"Clinical skills"}
    first = client.post("/objectives", data=payload, follow_redirects=True)
    assert first.status_code == 200
    second = client.post("/objectives", data=payload, follow_redirects=True)
    assert second.status_code == 200
    assert b"This objective already exists" in second.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM objectives WHERE terrain_id=? AND study_year=1 AND objective='Clinical skills'",(terrain,)).fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM objectives WHERE terrain_id=? AND study_year=5",(terrain,)).fetchone()[0] == 0

def test_documented_history_import_requires_preview_and_rejects_future_dates(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,3)", ("HistoryImport", "Resident"))
        conn.commit()

    valid_csv = (
        "resident,terrain,year_label,period,date_done\n"
        "HistoryImport Resident,Réanimation chirurgicale,2025-2026,01 Fév - 30 Avr,2026-04-30\n"
    ).encode("utf-8")
    preview = client.post(
        "/import-history",
        data={"file":(io.BytesIO(valid_csv),"history.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert preview.status_code == 200
    assert b"HistoryImport Resident" in preview.data
    assert b"Confirm import" in preview.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM history").fetchone()[0] == 0
    match = re.search(rb'name="preview_id" value="(\d+)"', preview.data)
    assert match
    confirmed = client.post("/import-history", data={
        "action":"confirm","preview_id":match.group(1).decode()
    }, follow_redirects=True)
    assert confirmed.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM history h JOIN residents r ON r.id=h.resident_id WHERE r.surname='HistoryImport'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='history.imported'").fetchone()[0] == 1

    future_csv = (
        "resident,terrain,year_label,period,date_done\n"
        "HistoryImport Resident,Réanimation polyvalente,2026-2027,01 Fév - 30 Avr,2099-03-01\n"
    ).encode("utf-8")
    future_preview = client.post(
        "/import-history",
        data={"file":(io.BytesIO(future_csv),"future-history.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert future_preview.status_code == 200
    assert b"A historical rotation date cannot be in the future" in future_preview.data
    assert b"disabled" in future_preview.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM history h JOIN residents r ON r.id=h.resident_id WHERE r.surname='HistoryImport'").fetchone()[0] == 1

def test_exam_route_rejects_malformed_ids_and_scores_without_server_error(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        rid = conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,1) RETURNING id",
            ("Malformed", "ExamInput"),
        ).fetchone()[0]
        conn.commit()

    bad_year = client.post("/exams", data={
        "action":"record","resident":str(rid),"study_year":"not-a-year",
        "module":"Module A","theory":"12","practical":"13",
    }, follow_redirects=True)
    assert bad_year.status_code == 200
    assert b"Choose a valid resident and study year" in bad_year.data

    bad_score = client.post("/exams", data={
        "action":"record","resident":str(rid),"study_year":"1",
        "module":"Module A","theory":"not-a-number","practical":"13",
    }, follow_redirects=True)
    assert bad_score.status_code == 200
    assert b"numeric marks from 0 to 20 are required" in bad_score.data

    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM exam_records WHERE resident_id=?", (rid,)
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action IN ('exam.recorded','exam.updated')"
        ).fetchone()[0] == 0



def test_dems_rejects_future_exam_date_without_vote_or_audit(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={
            "username": username,
            "password": "strong-test-password-2026",
            "role": "super_admin",
        }, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname,firstname,year) VALUES(?,?,4)",
            ("FutureDEMS", "Resident"),
        )
        rid = conn.execute(
            "SELECT id FROM residents WHERE surname='FutureDEMS'"
        ).fetchone()[0]
        conn.commit()

    future_date = (rotation_app.date.today() + timedelta(days=1)).isoformat()
    with client.session_transaction() as sess:
        sess["user"] = "benhocine"
        sess["role"] = "super_admin"
    response = client.post("/dems", data={
        "resident": str(rid),
        "success_date": future_date,
        "outcome": "approve",
        "reason": "Should not be accepted",
    }, follow_redirects=True)

    assert response.status_code == 200
    assert b"DEMS exam date cannot be in the future" in response.data
    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM dems_votes WHERE resident_id=?", (rid,)
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM dems_archive WHERE resident_id=?", (rid,)
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action='dems.vote_recorded'"
        ).fetchone()[0] == 0


def test_admin_routes_handle_malformed_record_ids_without_server_error(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)

    decision_response = client.post("/decisions", data={
        "aid": "not-an-id",
        "outcome": "reject",
        "reason": "Malformed request",
    }, follow_redirects=True)
    assert decision_response.status_code == 200

    user_response = client.post("/users", data={
        "action": "deactivate",
        "user_id": "not-a-user-id",
    }, follow_redirects=True)
    assert user_response.status_code == 200

    with rotation_app.db() as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM decisions WHERE reason='Malformed request'"
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM audit WHERE action IN ('assignment.override_approved','assignment.override_rejected','user.deactivated','user.activated')"
        ).fetchone()[0] == 0
