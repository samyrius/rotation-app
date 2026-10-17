# -*- coding: utf-8 -*-
from flask import Flask, request, redirect, session, render_template_string, flash, send_file, send_from_directory
import sqlite3, os, hashlib, hmac, io, csv, shutil, json, tempfile, re, subprocess, secrets, smtplib, time, zipfile
import warnings
from PIL import Image, ImageOps
from datetime import datetime, date
from werkzeug.utils import secure_filename

BASE=os.path.dirname(os.path.abspath(__file__))
DB=os.path.join(BASE,"rotation.db")
def local_secret_key():
 path=os.path.join(BASE,".rotation_secret")
 try:
  with open(path,"r",encoding="ascii") as handle: return handle.read().strip()
 except FileNotFoundError:
  value=secrets.token_hex(32)
  try:
   fd=os.open(path,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
   with os.fdopen(fd,"w",encoding="ascii") as handle: handle.write(value)
   return value
  except FileExistsError:
   with open(path,"r",encoding="ascii") as handle: return handle.read().strip()
app=Flask(__name__)
app.config["MAX_CONTENT_LENGTH"]=25*1024*1024
app.secret_key=os.environ.get("ROTATION_SECRET") or local_secret_key()
LANGS=("ar","fr","en")
D={
"ar":{"app":"ResiManage","home":"الرئيسية","residents":"المقيمون","manage_residents":"إضافة وإدارة المقيمين","import_residents":"استيراد المقيمين","promotion":"قرارات الترقية والامتحانات","dems_success":"قائمة الناجحين في DEMS","settings_page":"إعدادات التطبيق","terrains":"المصالح","capacities":"المقاعد","assign":"التوزيع","history":"التاريخ","summary":"ملخص المقيمين","decisions":"القرارات","settings":"الإعدادات","logout":"خروج","login":"دخول المسؤول","password":"كلمة المرور","signin":"دخول","badpass":"كلمة المرور غير صحيحة","year":"السنة الجامعية","resident":"المقيم","surname":"اللقب","firstname":"الاسم","study_year":"السنة","phone":"الهاتف","repeat":"يعيد السنة","status":"الحالة","active":"نشط","inactive":"غير نشط","add":"إضافة","save":"حفظ","terrain":"المصلحة","kind":"النوع","mandatory":"إجباري","basic":"أساسي","optional":"اختياري","mandatory_year":"السنة الإلزامية","any":"حسب البرنامج","period":"الفترة","places":"عدد المقاعد","used":"المستعمل","remaining":"المتبقي","date":"تاريخ الإنجاز","academic_year":"السنة الجامعية","choose":"اختيار","assign_now":"تعيين مقيم","check_save":"التحقق والحفظ","confirmed":"مؤكد","pending":"بانتظار القرار","rejected":"مرفوض","history_add":"تسجيل تربص سابق","reason":"سبب القرار","approve":"موافقة","reject":"رفض","language":"لغة الواجهة","director":"المسؤول","new_password":"كلمة مرور جديدة (اختياري)","change":"حفظ الإعدادات","note":"تغيير اللغة لا يغيّر البيانات المحفوظة.","priority":"الأولوية للسنة الرابعة ثم الثالثة والثانية والأولى.","none":"لا توجد بيانات","capacity_first":"حدّد السعة أولاً لهذه المصلحة والفترة","full":"السعة مكتملة","duplicate":"المقيم مسجّل في نفس الفترة بمصلحة: ","same_terrain_year":"سبق للمقيم إنجاز هذه المصلحة خلال السنة الجامعية في الفترة: ","history_block":"التربص منجز سابقاً؛ يمنع التكرار حتى إكمال جميع المصالح.","history_repeat":"يمكن التكرار لأن المقيم أكمل جميع المصالح.","mismatch":"السنة لا تطابق السنة الإلزامية؛ التعيين بانتظار قرار المسؤول","done":"تم الحفظ","overview":"ملخص التوزيع السنوي","count":"عدد التعيينات","type":"التصنيف" ,"export":"تصدير CSV","backup":"تنزيل نسخة احتياطية","all":"كل المصالح","four":"التعيينات المؤكدة / 4 فترات","waiting":"معلّق","users":"المستخدمون","import":"استيراد","objectives":"الأهداف","exams":"الامتحانات","dems":"أرشيف DEMS","audit":"سجل التدقيق","restore":"استرجاع","username":"اسم المستخدم","role":"الصلاحية","standard_user":"مستخدم عادي","super_admin":"مسؤول عام","created":"تاريخ الإنشاء","active_col":"نشط","actor":"الفاعل","action":"الإجراء","details":"التفاصيل","module":"المادة","theory":"نظري","practical":"تطبيقي","success_date":"تاريخ النجاح","file":"الملف","objective":"الهدف","archive":"أرشفة","users_help":"أنشئ حسابًا مستقلًا لكل مسؤول، وامنح الأستاذ Benhocine والدكتور Dris صلاحية المسؤول العام.","import_help":"ارفع ملفًا يحتوي على اللقب والاسم والسنة الدراسية. راجع المعاينة قبل التأكيد؛ لن تُختلق أي تربصات سابقة.","restore_help":"سيتم التحقق من سلامة قاعدة البيانات وإنشاء نسخة احتياطية من القاعدة الحالية قبل الاسترجاع.","confirm_password":"تأكيد كلمة المرور","confirm_import":"تأكيد الاستيراد","cancel":"إلغاء","upload_preview":"رفع ومعاينة","archive_note":"تُحفظ بيانات المقيم وتاريخه ولا تُحذف."},
"fr":{"app":"ResiManage","home":"Accueil","residents":"Résidents","manage_residents":"Gérer les résidents","import_residents":"Importer les résidents","promotion":"Décisions de promotion et examens","dems_success":"Lauréats du DEMS","settings_page":"Paramètres de l’application","terrains":"Services","capacities":"Capacités","assign":"Affectations","history":"Historique","summary":"Bilan résidents","decisions":"Décisions","settings":"Paramètres","logout":"Déconnexion","login":"Connexion administrateur","password":"Mot de passe","signin":"Connexion","badpass":"Mot de passe incorrect","year":"Année universitaire","resident":"Résident","surname":"Nom","firstname":"Prénom","study_year":"Année","phone":"Téléphone","repeat":"Redouble","status":"Statut","active":"Actif","inactive":"Inactif","add":"Ajouter","save":"Enregistrer","terrain":"Service","kind":"Type","mandatory":"Obligatoire","basic":"De base","optional":"Optionnel","mandatory_year":"Année requise","any":"Selon le programme","period":"Période","places":"Nombre de places","used":"Utilisées","remaining":"Restantes","date":"Date de réalisation","academic_year":"Année universitaire","choose":"Choisir","assign_now":"Affecter un résident","check_save":"Vérifier et enregistrer","confirmed":"Confirmé","pending":"En attente","rejected":"Refusé","history_add":"Enregistrer un stage antérieur","reason":"Motif de décision","approve":"Approuver","reject":"Refuser","language":"Langue de l’interface","director":"Responsable","new_password":"Nouveau mot de passe (facultatif)","change":"Enregistrer les paramètres","note":"Changer de langue ne modifie pas les données.","priority":"Priorité aux 4e, puis 3e, 2e et 1re années.","none":"Aucune donnée","capacity_first":"Définissez d’abord la capacité du service et de la période","full":"Capacité complète","duplicate":"Résident déjà affecté pendant cette période au service : ","same_terrain_year":"Ce résident a déjà effectué ce service cette année, pendant la période : ","history_block":"Stage déjà effectué ; répétition bloquée avant d’avoir terminé tous les services.","history_repeat":"Répétition autorisée : tous les services ont été effectués.","mismatch":"Année différente de l’année requise ; décision du responsable nécessaire","done":"Enregistré","overview":"Vue annuelle des affectations","count":"Nombre d’affectations","type":"Catégorie" ,"export":"Exporter CSV","backup":"Télécharger la sauvegarde","all":"Tous les services","four":"Affectations confirmées / 4 périodes","waiting":"En attente","users":"Utilisateurs","import":"Importer","objectives":"Objectifs","exams":"Examens","dems":"Archive DEMS","audit":"Journal d’audit","restore":"Restaurer","username":"Identifiant","role":"Rôle","standard_user":"Utilisateur standard","super_admin":"Super Admin","created":"Créé le","active_col":"Actif","actor":"Acteur","action":"Action","details":"Détails","module":"Matière","theory":"Théorie","practical":"Pratique","success_date":"Date de réussite","file":"Fichier","objective":"Objectif","archive":"Archiver","users_help":"Créez un compte distinct pour chaque responsable et attribuez le rôle Super Admin à Pr Benhocine et Dr Dris.","import_help":"Importez une liste contenant le nom, le prénom et l’année d’étude. Vérifiez l’aperçu avant confirmation ; aucun stage passé ne sera inventé.","restore_help":"L’intégrité de la base sera vérifiée et une copie de la base actuelle sera créée avant la restauration.","confirm_password":"Confirmer le mot de passe","confirm_import":"Confirmer l’importation","cancel":"Annuler","upload_preview":"Charger et prévisualiser","archive_note":"Les données et l’historique du résident sont conservés, pas supprimés."},
"en":{"app":"ResiManage","home":"Home","residents":"Residents","manage_residents":"Manage residents","import_residents":"Import residents","promotion":"Promotion decisions & exams","dems_success":"DEMS successful graduates","settings_page":"App settings","terrains":"Departments","capacities":"Capacity","assign":"Assignments","history":"History","summary":"Resident summary","decisions":"Decisions","settings":"Settings","logout":"Log out","login":"Administrator login","password":"Password","signin":"Sign in","badpass":"Incorrect password","year":"Academic year","resident":"Resident","surname":"Surname","firstname":"First name","study_year":"Year","phone":"Phone","repeat":"Repeating year","status":"Status","active":"Active","inactive":"Inactive","add":"Add","save":"Save","terrain":"Department","kind":"Type","mandatory":"Mandatory","basic":"Core","optional":"Optional","mandatory_year":"Required year","any":"As per curriculum","period":"Period","places":"Places","used":"Used","remaining":"Remaining","date":"Completion date","academic_year":"Academic year","choose":"Choose","assign_now":"Assign resident","check_save":"Validate and save","confirmed":"Confirmed","pending":"Pending decision","rejected":"Rejected","history_add":"Record previous rotation","reason":"Decision reason","approve":"Approve","reject":"Reject","language":"Interface language","director":"Administrator","new_password":"New password (optional)","change":"Save settings","note":"Changing language does not change saved data.","priority":"Priority: year 4, then 3, 2 and 1.","none":"No data","capacity_first":"Set capacity for this department and period first","full":"Capacity is full","duplicate":"Resident already assigned during this period to: ","same_terrain_year":"This resident already completed this department this academic year, during: ","history_block":"Rotation completed previously; repeats blocked until all departments are completed.","history_repeat":"Repeat allowed because all departments have been completed.","mismatch":"Year does not match required year; administrator decision needed","done":"Saved","overview":"Annual assignment overview","count":"Assignment count","type":"Category" ,"export":"Export CSV","backup":"Download backup","all":"All departments","four":"Confirmed assignments / 4 periods","waiting":"Pending","users":"Users","import":"Import","objectives":"Objectives","exams":"Exams","dems":"DEMS archive","audit":"Audit log","restore":"Restore","username":"Username","role":"Role","standard_user":"Standard user","super_admin":"Super Admin","created":"Created","active_col":"Active","actor":"Actor","action":"Action","details":"Details","module":"Module","theory":"Theory","practical":"Practical","success_date":"Success date","file":"File","objective":"Objective","archive":"Archive","users_help":"Create a separate account for each administrator and assign Super Admin to Pr Benhocine and Dr Dris.","import_help":"Upload a list containing surname, first name and study year. Review the preview before confirming; no past rotations are invented.","restore_help":"The database integrity will be checked and a copy of the current database will be created before restoring.","confirm_password":"Confirm password","confirm_import":"Confirm import","cancel":"Cancel","upload_preview":"Upload and preview","archive_note":"Resident data and history are retained, not deleted."}
}
def tr(k): return D.get(session.get("lang","ar"),D["ar"]).get(k,D["ar"].get(k,k))
def db():
 c=sqlite3.connect(DB,timeout=20); c.row_factory=sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); return c
def setting(c,k,default=""):
 r=c.execute("SELECT value FROM settings WHERE key=?",(k,)).fetchone(); return r["value"] if r else default
def hash_password(password):
 salt=os.urandom(16).hex()
 digest=hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),bytes.fromhex(salt),310000).hex()
 return "pbkdf2_sha256$"+salt+"$"+digest
def verify_password(password, stored):
 if stored.startswith("pbkdf2_sha256$"):
  try:
   _,salt,digest=stored.split("$",2)
   candidate=hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),bytes.fromhex(salt),310000).hex()
   return hmac.compare_digest(candidate,digest)
  except (ValueError,TypeError):
   return False
 return hmac.compare_digest(hashlib.sha256(password.encode("utf-8")).hexdigest(),stored)
def ensure_column(c, table, column, spec):
 existing={row[1] for row in c.execute("PRAGMA table_info("+table+")")}
 if column not in existing:
  c.execute("ALTER TABLE "+table+" ADD COLUMN "+column+" "+spec)
def audit_event(c, action, details):
 c.execute("INSERT INTO audit(action,details,actor) VALUES(?,?,?)",
           (action,str(details)[:1000],str(session.get("user","system"))))
def init():
 c=db()
 c.executescript("""
 CREATE TABLE IF NOT EXISTS residents(id INTEGER PRIMARY KEY,surname TEXT NOT NULL,firstname TEXT NOT NULL,year INTEGER NOT NULL CHECK(year BETWEEN 1 AND 4),phone TEXT DEFAULT '',repeat_year INTEGER DEFAULT 0,status TEXT DEFAULT 'active');
 CREATE TABLE IF NOT EXISTS terrains(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL,kind TEXT DEFAULT 'basic',mandatory_year INTEGER DEFAULT 0,active INTEGER NOT NULL DEFAULT 1);
 CREATE TABLE IF NOT EXISTS periods(id INTEGER PRIMARY KEY,year_label TEXT NOT NULL,name TEXT NOT NULL,UNIQUE(year_label,name));
 CREATE TABLE IF NOT EXISTS capacities(id INTEGER PRIMARY KEY,year_label TEXT NOT NULL,terrain_id INTEGER,period_id INTEGER,places INTEGER NOT NULL,UNIQUE(year_label,terrain_id,period_id));
 CREATE TABLE IF NOT EXISTS assignments(id INTEGER PRIMARY KEY,resident_id INTEGER,terrain_id INTEGER,period_id INTEGER,status TEXT DEFAULT 'confirmed',UNIQUE(resident_id,period_id));
 CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY,resident_id INTEGER,terrain_id INTEGER,year_label TEXT,period TEXT,date_done TEXT);
 CREATE TABLE IF NOT EXISTS decisions(id INTEGER PRIMARY KEY,details TEXT,decision TEXT,reason TEXT,actor TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,action TEXT,details TEXT,actor TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
 """)
 # Additive migrations preserve existing resident and rotation-history records.
 for col, spec in {"birth_place":"TEXT DEFAULT ''","birth_date":"TEXT DEFAULT ''","photo_path":"TEXT DEFAULT ''","dems_status":"TEXT DEFAULT ''"}.items():
  ensure_column(c,"residents",col,spec)
 ensure_column(c,"terrains","active","INTEGER NOT NULL DEFAULT 1")
 ensure_column(c,"audit","actor","TEXT DEFAULT ''")
 c.executescript("""
 CREATE TABLE IF NOT EXISTS exam_records(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,study_year INTEGER NOT NULL,module TEXT NOT NULL,theory REAL NOT NULL,practical REAL NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS objectives(id INTEGER PRIMARY KEY,terrain_id INTEGER NOT NULL,study_year INTEGER NOT NULL,objective TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS dems_archive(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL UNIQUE,success_date TEXT NOT NULL,profile_json TEXT NOT NULL,history_json TEXT NOT NULL,archived_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS promotion_votes(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,from_year INTEGER NOT NULL,decision TEXT NOT NULL CHECK(decision IN ('approve','reject')),reason TEXT NOT NULL,actor TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(resident_id,from_year,actor));
 CREATE TABLE IF NOT EXISTS dems_votes(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,success_date TEXT NOT NULL,decision TEXT NOT NULL CHECK(decision IN ('approve','reject')),reason TEXT NOT NULL,actor TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(resident_id,success_date,actor));
 CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,username TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,role TEXT NOT NULL CHECK(role IN ('super_admin','user')),active INTEGER NOT NULL DEFAULT 1,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS pending_registrations(id INTEGER PRIMARY KEY,username TEXT NOT NULL,email TEXT NOT NULL UNIQUE,password_hash TEXT NOT NULL,code_hash TEXT NOT NULL,expires_at INTEGER NOT NULL,attempts INTEGER NOT NULL DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS import_previews(id INTEGER PRIMARY KEY,owner TEXT NOT NULL,rows_json TEXT NOT NULL,errors_json TEXT NOT NULL,filename TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS leave_years(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,year INTEGER NOT NULL,earned_days INTEGER NOT NULL DEFAULT 30,created_by TEXT DEFAULT '',UNIQUE(resident_id,year));
 CREATE TABLE IF NOT EXISTS leave_entries(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,start_date TEXT NOT NULL,end_date TEXT NOT NULL,days INTEGER NOT NULL,notes TEXT DEFAULT '',created_by TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS external_rotations(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,supervisor TEXT NOT NULL,destination TEXT NOT NULL,hospital TEXT NOT NULL,country TEXT NOT NULL,rotation_type TEXT NOT NULL,start_date TEXT NOT NULL,end_date TEXT NOT NULL,notes TEXT DEFAULT '',status TEXT NOT NULL DEFAULT 'active',created_by TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
 """)
 # Migrate user email fields only after the users table exists (fresh DB and legacy DB).
 ensure_column(c,"users","email","TEXT DEFAULT ''")
 ensure_column(c,"users","email_verified","INTEGER NOT NULL DEFAULT 0")
 ensure_column(c,"import_previews","import_type","TEXT NOT NULL DEFAULT 'residents'")
 y=str(datetime.now().year)+"-"+str(datetime.now().year+1)
 c.execute("INSERT OR IGNORE INTO settings VALUES('year',?)",(y,))
 initial_password=os.environ.get("ROTATION_INITIAL_PASSWORD")
 if initial_password and not setting(c,"password"):
  c.execute("INSERT OR IGNORE INTO settings VALUES('password',?)",(hash_password(initial_password),))
 c.execute("INSERT OR IGNORE INTO settings VALUES('director',?)",("Professor Benhocine",))
 for n,k in [("Réanimation chirurgicale","basic"),("Réanimation polyvalente","basic"),("Chirurgie générale","mandatory"),("Urologie","basic"),("Orthopédie","basic"),("Neurochirurgie","basic"),("Sédation","optional"),("Maxillo-faciale","optional"),("CCI","basic"),("Obstétrique","basic"),("Chirurgie cardiaque","optional"),("PU Médecine","basic"),("PU Chirurgie","basic"),("ORL","basic"),("Ophtalmologie","basic"),("Gynécologie","basic"),("Chirurgie thoracique","basic"),("Brûlés","optional"),("CAC DBK","optional")]:
  c.execute("INSERT OR IGNORE INTO terrains(name,kind) VALUES(?,?)",(n,k))
 # One-time correction for databases created before the agreed terrain classification.
 # Do not reapply on every startup, so later authorized admin edits remain intact.
 if not setting(c,"terrain_classification_migration_20261017"):
  for name in ("Sédation","Maxillo-faciale","Chirurgie cardiaque","Brûlés","CAC DBK"):
   c.execute("UPDATE terrains SET kind='optional' WHERE name=?",(name,))
  c.execute("UPDATE terrains SET kind='basic' WHERE name='Gynécologie'")
  c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",("terrain_classification_migration_20261017","1"))
 y=setting(c,"year",y)
 # Fill any missing canonical periods in legacy/partially initialized databases.
 for n in ["01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"]:
  c.execute("INSERT OR IGNORE INTO periods(year_label,name) VALUES(?,?)",(y,n))
 c.commit(); c.close()
@app.before_request
def boot():
 if not getattr(app,"_ready",False): init(); app._ready=True

@app.before_request
def verify_csrf():
 if request.method=="POST":
  expected=session.get("_csrf_token","")
  supplied=request.form.get("_csrf_token","")
  if not expected or not supplied or not hmac.compare_digest(expected,supplied):
   return "Invalid or missing CSRF token",400

def page(body,**kw):
 lang=session.get("lang","ar")
 csrf_token=session.setdefault("_csrf_token",secrets.token_urlsafe(32))
 body=re.sub(r'(<form\b(?=[^>]*\bmethod\s*=\s*["\']post["\'])[^>]*>)',lambda match:match.group(1)+'<input type="hidden" name="_csrf_token" value="{{ csrf_token }}">',body,flags=re.IGNORECASE)
 shell="""<!doctype html><html lang="{{lang}}" dir="{{'rtl' if lang=='ar' else 'ltr'}}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#103c66"><meta name="application-name" content="ResiManage"><link rel="icon" type="image/svg+xml" href="/static/resimanage-3d-icon.jpg"><title>{{t('app')}} — {{t('year')}}</title></head><style>*{box-sizing:border-box}html{width:100%;max-width:100%;overflow-x:hidden}body{font-family:Arial,sans-serif;background:#f1f5f9;margin:0;color:#182b3d;width:100%;max-width:100%;overflow-x:hidden;line-height:1.5}header{background:linear-gradient(135deg,#103c66 0%,#17689a 58%,#23a6a5 100%);color:white;padding:clamp(10px,3vw,18px);overflow-wrap:anywhere}.brand-lockup{display:flex;align-items:center;gap:14px;max-width:1050px;margin:0 auto}.brand-icon{width:clamp(48px,12vw,70px);height:auto;flex:0 0 auto;filter:drop-shadow(0 5px 5px #001b3380)}.brand-lockup h2{margin:0;font-size:clamp(21px,4vw,30px);letter-spacing:.2px}.brand-tagline{margin:2px 0 5px;opacity:.93;font-size:clamp(12px,2.5vw,15px)}.brand-year{font-size:12px;opacity:.9}main{width:100%;max-width:1050px;margin:0 auto;padding:clamp(7px,2.5vw,16px);min-width:0}nav{display:flex;align-items:flex-start;flex-wrap:wrap;gap:6px;overflow:visible;max-width:100%;background:white;padding:10px}.nav-group{position:relative;min-width:0}.nav-group summary{list-style:none;background:#164e83;color:white;padding:9px 12px;border-radius:7px;cursor:pointer;white-space:nowrap}.nav-group summary::-webkit-details-marker{display:none}.nav-group summary:after{content:" ▾";font-size:.85em}.nav-menu{display:flex;flex-direction:column;gap:5px;min-width:180px;padding:7px;margin-top:5px;background:#fff;border:1px solid #d7e1eb;border-radius:9px;box-shadow:0 6px 18px #182b3d22}.nav-menu a{display:block;width:100%}.main-nav>a{flex:0 0 auto}.nav-group[open]{z-index:10}.nav-group[open] .nav-menu{position:absolute;top:100%;inset-inline-start:0;max-width:min(85vw,280px);max-height:65vh;overflow:auto}.nav-group[open] .nav-menu a{white-space:normal}a,button{background:#164e83;color:white;padding:9px 12px;border:0;border-radius:7px;text-decoration:none;display:inline-block;white-space:normal;overflow-wrap:anywhere;max-width:100%;cursor:pointer}section{background:white;padding:clamp(9px,2.5vw,15px);border:1px solid #d7e1eb;border-radius:10px;margin:10px 0;min-width:0;max-width:100%;overflow-wrap:anywhere}input,select,textarea{font:inherit;padding:9px;border:1px solid #bccada;border-radius:7px;width:100%;max-width:100%;margin:4px 0 9px}label{display:block;min-width:0}table{border-collapse:collapse;width:100%;min-width:540px}td,th{text-align:start;border-bottom:1px solid #ddd;padding:8px;overflow-wrap:anywhere}.wrap{max-width:100%;overflow-x:auto;-webkit-overflow-scrolling:touch}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(170px,100%),1fr));gap:8px;min-width:0}.msg{background:#fff0c2;padding:10px;border-radius:7px;overflow-wrap:anywhere}.good{color:green}.bad{color:#b21d1d}.muted{color:#526274}img,video,iframe{max-width:100%;height:auto}@media(max-width:600px){main{padding:7px}header{font-size:clamp(13px,3.5vw,16px)}nav{padding:7px;gap:5px}.nav-group summary,nav>a{font-size:13px;padding:8px 9px}.nav-menu{min-width:min(220px,82vw)}section{margin:7px 0;padding:10px}.grid{grid-template-columns:minmax(0,1fr)}table{min-width:540px}} [dir=rtl]{text-align:right}[dir=rtl] nav{direction:rtl}
/* ResiManage unified interior visual identity */
:root{--rt-navy:#142b49;--rt-blue:#2768a8;--rt-mint:#36cfae;--rt-cyan:#45c7e8;--rt-coral:#ff8b72;--rt-amber:#ffc36c;--rt-violet:#9b8cf5;--rt-ink:#21334b;--rt-line:#dce6f0;--rt-glass:#ffffffd9}
body{background:radial-gradient(ellipse at 8% 0%,#ffd2ad55,transparent 28%),radial-gradient(ellipse at 95% 22%,#7ce8d533,transparent 26%),linear-gradient(145deg,#f3f7fc,#edf3fa 58%,#f8f1f8);color:var(--rt-ink)}
header.brand-header{background:radial-gradient(ellipse at 85% 0%,#ffb66c45,transparent 34%),linear-gradient(115deg,#101d36,#1c4265 58%,#28566d);box-shadow:0 12px 32px #172d4825;border-bottom:3px solid transparent;border-image:linear-gradient(90deg,var(--rt-mint),var(--rt-cyan),var(--rt-coral),var(--rt-violet)) 1}
.brand-icon{filter:drop-shadow(0 7px 12px #0005)}main{padding-top:18px;padding-bottom:26px}
.main-nav{margin:0 auto;position:sticky;top:0;z-index:20;justify-content:center;align-items:center;gap:8px;background:#ffffffdf;backdrop-filter:blur(16px);border:1px solid #ffffff;border-bottom-color:#dce6f0;border-radius:0 0 18px 18px;box-shadow:0 9px 25px #17365b12}
.main-nav>a,.nav-group summary{border:1px solid #ffffff25;background:linear-gradient(135deg,#214e79,#286da2);box-shadow:0 4px 10px #1c42651a;transition:transform .18s ease,box-shadow .18s ease,filter .18s ease}
.main-nav>a:hover,.nav-group summary:hover{transform:translateY(-2px);box-shadow:0 7px 15px #1c42652a;filter:saturate(1.12)}
.main-nav>a:nth-of-type(1){background:linear-gradient(135deg,#2768a8,#3986c9)}
.main-nav>a:nth-of-type(2){background:linear-gradient(135deg,#168f88,#35c7a9)}
.main-nav>a:nth-of-type(3){background:linear-gradient(135deg,#cf7547,#ffab68)}
.main-nav>a:nth-of-type(4){background:linear-gradient(135deg,#6d62bf,#9c8df5)}
.nav-group summary{background:linear-gradient(135deg,#254a70,#477e9e)}
.nav-menu{background:#fffffff2;backdrop-filter:blur(15px);border:1px solid #d9e7f0;border-radius:14px;box-shadow:0 18px 45px #17365b25}
.nav-menu a{background:#f1f7fb;color:#244461;border:1px solid transparent;border-inline-start:3px solid var(--rt-cyan);border-radius:9px;transition:background .18s ease,transform .18s ease}
.nav-menu a:hover{background:#e3f7f1;transform:translateX(-2px)}
section{background:var(--rt-glass);backdrop-filter:blur(10px);border:1px solid #ffffff;box-shadow:0 9px 28px #2444610d;border-radius:17px;transition:box-shadow .2s ease}
section h2,section h3{color:#204c70}.nav-icon{display:inline-flex;align-items:center;justify-content:center;min-width:20px;margin-inline-end:5px;font-size:16px;line-height:1;color:#b9ffec}.nav-group summary .nav-icon{color:#ffcf91}.rt-quick-actions{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(210px,100%),1fr));gap:12px;margin:14px 0}.rt-action-card{display:flex;align-items:center;gap:13px;min-width:0;padding:15px;border-radius:16px;background:linear-gradient(135deg,#ffffffef,#f1f8fb);border:1px solid #fff;box-shadow:0 8px 22px #24446112;color:#244461;text-decoration:none}.rt-action-card:hover{transform:translateY(-3px)}.rt-action-icon{display:flex;align-items:center;justify-content:center;flex:0 0 46px;width:46px;height:46px;border-radius:14px;font-size:24px;color:#15364d;background:linear-gradient(135deg,#b9f6e8,#a8e8f7)}.rt-action-card:nth-child(2) .rt-action-icon{background:linear-gradient(135deg,#ffd6b2,#ffad8e)}.rt-action-card:nth-child(3) .rt-action-icon{background:linear-gradient(135deg,#d9d1ff,#b7b0fa)}.rt-action-card:nth-child(4) .rt-action-icon{background:linear-gradient(135deg,#fff0b9,#ffd18a)}.rt-action-card strong,.rt-action-card small{display:block}.rt-action-card strong{color:#214c6b}.rt-action-card small{font-size:12px;color:#657b8e}.rt-stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(145px,100%),1fr));gap:10px;margin:14px 0}.rt-stat{padding:15px;border-radius:15px;background:linear-gradient(135deg,#fff,#f0f7fc);border:1px solid #fff;box-shadow:0 6px 18px #2444610d}.rt-stat span{display:block;font-size:12px;color:#617b8f}.rt-stat strong{display:block;font-size:26px;color:#214f73;line-height:1.3}.rt-stat:nth-child(2) strong{color:#15977f}.rt-stat:nth-child(3) strong{color:#c8754f}.rt-stat:nth-child(4) strong{color:#7668c5}.rt-dashboard-title{display:flex;align-items:center;gap:10px}.rt-dashboard-title .rt-dot{width:11px;height:11px;border-radius:50%;background:linear-gradient(135deg,var(--rt-mint),var(--rt-cyan));box-shadow:0 0 0 5px #42d9bd1b}.rt-year-chip{display:inline-flex;min-width:28px;justify-content:center;padding:3px 8px;border-radius:9px;background:#e4f0ff;color:#285a92;font-weight:700}.rt-progress{display:inline-block;vertical-align:middle;width:66px;height:7px;overflow:hidden;border-radius:10px;background:#e5edf4}.rt-progress span{display:block;height:100%;border-radius:10px;background:linear-gradient(90deg,#36cfae,#45c7e8)}.rt-resident-link{background:transparent;color:#275b7f;padding:2px 0;box-shadow:none}
section h2:after,section h3:after{content:"";display:block;width:54px;height:3px;margin-top:7px;border-radius:5px;background:linear-gradient(90deg,var(--rt-mint),var(--rt-cyan),var(--rt-coral))}
a,button{background:linear-gradient(135deg,#24659c,#2b80b1);box-shadow:0 3px 9px #214e7918;transition:filter .18s ease,transform .18s ease,box-shadow .18s ease}
a:hover,button:hover{filter:brightness(1.07) saturate(1.1);box-shadow:0 6px 14px #214e7925}
button:active,a:active{transform:translateY(1px)}
input,select,textarea{background:#ffffffdc;border-color:#c8d8e7;border-radius:11px;transition:border-color .15s ease,box-shadow .15s ease}
input:focus,select:focus,textarea:focus{outline:none;border-color:#37bfae;box-shadow:0 0 0 3px #37cfae26}
table{border-collapse:separate;border-spacing:0;border:1px solid var(--rt-line);border-radius:12px;overflow:hidden;background:#fff}
th{background:linear-gradient(100deg,#e7f4f7,#eef0fc);color:#234b68;font-weight:700}
td,th{border-bottom:1px solid #e7edf4;padding:10px}
tr:last-child td{border-bottom:0}
tbody tr:nth-child(even),table tr:nth-child(even){background:#f8fbfe}
.msg{background:linear-gradient(100deg,#fff2d8,#fff8eb);border:1px solid #f4d49a;color:#72501c;border-radius:12px}
.good{color:#13856d;font-weight:700}.bad{color:#c34e50;font-weight:700}.muted{color:#60758a}
.rt-language{display:flex;align-items:center;gap:5px;margin-inline-start:auto}
.rt-language a{padding:7px 9px;border-radius:20px;background:#e9f5f7;color:#244c66;border:1px solid #d5e8ef;box-shadow:none;font-size:12px}
.rt-language a[aria-current=page]{background:linear-gradient(100deg,#2ac7a7,#55c9e8);color:#102f47;font-weight:700}
.rt-footer{text-align:center;padding:17px 10px 10px;color:#637b91;font-size:12px}
.rt-footer strong{color:#2d5877}.rt-footer .rt-sep{color:#ff986c;padding:0 4px}
@media(max-width:600px){.main-nav{position:relative;border-radius:0 0 14px 14px;justify-content:flex-start}.rt-language{width:100%;justify-content:center;margin:3px 0 0}.main-nav>a,.nav-group summary{font-size:12px}.rt-footer{line-height:1.9}}
@media(prefers-reduced-motion:reduce){*,*:before,*:after{scroll-behavior:auto!important;transition:none!important;animation:none!important}}
/* ResiManage dark navy mockup theme */
:root{color-scheme:dark;--rt-cyan:#19d9e8;--rt-mint:#19dfbd;--rt-text:#e8f7ff}
html,body{background:#041426!important;color:#e8f7ff!important}
body{background:radial-gradient(ellipse at 10% 0%,#0a467055,transparent 34%),linear-gradient(150deg,#020d1c,#061a30 58%,#041426)!important;color:#e8f7ff}
header.brand-header{background:linear-gradient(115deg,#020d1c,#07375b 58%,#075a70)!important;border-bottom:2px solid var(--rt-cyan)!important}
main{max-width:1120px}
.main-nav{background:#061a2be8!important;border:1px solid #0c4564!important;border-radius:0 0 18px 18px!important}
.main-nav>a,.nav-group summary{background:linear-gradient(135deg,#0b3557,#0b5572)!important;border:1px solid #17668a!important;color:#e9fbff!important;border-radius:11px!important}
.nav-menu{background:#071d32f5!important;border-color:#155a80!important}
.nav-menu a{background:#0a2b47!important;color:#dff8ff!important;border-color:#104765!important}
section,.rt-stat,.rt-action-card,.login-brand,.login-panel{background:linear-gradient(145deg,#08223aee,#061a30f5)!important;border:1px solid #125173!important;box-shadow:0 10px 28px #0004!important;color:#e8f7ff!important;border-radius:17px!important}
h1,h2,h3,h4,strong,th{color:#eafaff}
p,label,td,small,.muted{color:#a9ccdf}
a,button{background:linear-gradient(110deg,#087f9f,#09b9bd)!important;color:#f3feff!important;border:1px solid #24cce0!important;border-radius:11px!important}
input,select,textarea{background:#04182b!important;color:#e8f7ff!important;border:1px solid #176083!important;border-radius:10px!important}
th{background:#0b3552!important;color:#70eaf2!important}td{background:#061d32;border-bottom:1px solid #12415e!important}
.rt-dot{background:var(--rt-mint)!important}.rt-year-chip{background:#0b4c69!important;color:#8ff5f4!important}
.rt-progress{background:#123750!important}.rt-progress span{background:linear-gradient(90deg,#16d8bc,#16c8ed)!important}
.login-emblem{color:#20e2e7!important;text-shadow:0 0 25px #16d8ed!important}
@media(max-width:600px){main{padding:8px}.main-nav{position:relative;top:auto;justify-content:flex-start}.rt-stats,.rt-quick-actions{grid-template-columns:repeat(2,minmax(0,1fr))}}

/* Functional navigation redesign: task-based sidebar, clear active states, responsive mobile groups */
.rt-app-shell{display:grid;grid-template-columns:260px minmax(0,1fr);align-items:start;gap:18px;width:100%;max-width:1500px;margin:0 auto;padding:16px;direction:inherit}
.rt-app-shell>main{width:100%;max-width:none;min-width:0;margin:0;padding:0}
.main-nav{position:sticky;top:12px;z-index:30;display:flex;flex-direction:column;align-items:stretch;justify-content:flex-start;gap:7px;width:100%;min-width:0;padding:12px;background:#071a2beF!important;border:1px solid #15506b!important;border-radius:18px!important;box-shadow:0 12px 30px #00101f30!important;backdrop-filter:blur(18px)}
.main-nav>a,.nav-group summary{display:flex;align-items:center;gap:9px;width:100%;min-height:44px;padding:10px 12px!important;border:1px solid #174c68!important;border-radius:11px!important;background:#0a2b47!important;color:#dff8ff!important;box-shadow:none!important;text-align:start;font-size:14px;font-weight:600;transition:background .16s ease,border-color .16s ease,transform .16s ease}
.nav-group{width:100%;min-width:0}
.nav-group summary{list-style:none;cursor:pointer;justify-content:space-between}
.nav-group summary::-webkit-details-marker{display:none}
.nav-group summary:after{content:'⌄';font-size:14px;color:#74e7e8;transition:transform .16s ease}
.nav-group[open] summary:after{transform:rotate(180deg)}
.nav-group summary .nav-icon,.main-nav .nav-icon{display:inline-flex;align-items:center;justify-content:center;flex:0 0 24px;min-width:24px;height:24px;margin:0;color:#71e5e4;font-size:17px}
.nav-menu{position:static!important;display:flex;flex-direction:column;gap:4px;min-width:0;width:100%;max-width:none!important;max-height:none!important;overflow:visible!important;margin:5px 0 3px;padding:5px 7px;background:#041728!important;border:1px solid #123e57!important;border-radius:12px!important;box-shadow:none!important}
.nav-menu a{display:flex;align-items:center;width:100%;min-height:38px;padding:8px 10px!important;background:transparent!important;color:#b8d5e4!important;border:1px solid transparent!important;border-inline-start:3px solid transparent!important;border-radius:8px!important;box-shadow:none!important;font-size:13px;font-weight:500}
.main-nav a:hover,.nav-group summary:hover,.nav-menu a:hover{background:#0e3a56!important;color:#fff!important;border-color:#1a6a83!important;transform:none!important;box-shadow:none!important}
.main-nav a.is-active,.nav-menu a.is-active{background:linear-gradient(100deg,#0a6b7a,#0d536b)!important;color:#fff!important;border-color:#22c9c7!important;border-inline-start-color:#39e0c1!important;font-weight:700}
.rt-language{display:flex;align-items:center;justify-content:center;gap:6px;margin:8px 0 0;padding-top:12px;border-top:1px solid #17435b}
.rt-language a{min-width:40px;text-align:center}
.rt-language a[aria-current=page]{background:#19b9b5!important;color:#031b2c!important;border-color:#51e4d9!important}
.rt-app-shell section{margin:0 0 14px}
.rt-app-shell .rt-quick-actions{margin-top:16px}
:focus-visible{outline:3px solid #39dfcf!important;outline-offset:2px}
 .rt-app-shell:not(:has(.main-nav)){display:block;max-width:1120px} @media(max-width:900px){.rt-app-shell{grid-template-columns:220px minmax(0,1fr);gap:12px;padding:10px}.main-nav{padding:9px}}
@media(max-width:680px){.rt-app-shell{display:flex;flex-direction:column;gap:12px;padding:8px}.main-nav{position:relative;top:auto;display:grid;grid-template-columns:repeat(2,minmax(0,1fr));align-items:start;gap:7px;padding:9px;border-radius:14px!important}.main-nav>.nav-home{grid-column:1/-1}.main-nav>.nav-group{align-self:start}.nav-group summary{min-height:44px;font-size:12px!important;padding:8px!important}.nav-group[open] .nav-menu{position:static!important;grid-column:1/-1;max-height:42vh!important;overflow:auto!important}.nav-menu a{font-size:12px}.rt-language{grid-column:1/-1;margin-top:2px;padding-top:8px}.rt-app-shell>main{padding:0}.rt-app-shell section{margin-bottom:10px}}
@media(prefers-reduced-motion:reduce){.main-nav *{transition:none!important}}
.rt-profile-panel{padding:clamp(16px,3vw,28px);background:linear-gradient(145deg,#ffffffef,#f1f8fcf2);border:1px solid #fff;box-shadow:0 16px 38px #17365b14}.rt-profile-head{display:flex;align-items:center;gap:20px;padding:4px 0 22px;border-bottom:1px solid #dce7f0;margin-bottom:20px}.rt-profile-avatar[hidden]{display:none!important}.rt-profile-avatar{display:block;width:clamp(88px,18vw,132px);height:clamp(88px,18vw,132px);object-fit:cover;border-radius:28px;border:4px solid white;box-shadow:0 10px 24px #17365b24;background:#e8f2f7}.rt-avatar-placeholder{display:grid;place-items:center;font-size:48px}.rt-profile-heading{min-width:0}.rt-profile-heading h2{margin:4px 0 8px;font-size:clamp(22px,4vw,32px);overflow-wrap:anywhere}.rt-eyebrow{font-size:11px;letter-spacing:1.3px;font-weight:800;color:#168c91}.rt-chip{display:inline-flex;align-items:center;padding:4px 10px;margin-inline-end:6px;border-radius:999px;background:#e5efff;color:#285b8c;font-size:12px;font-weight:700}.rt-chip-active{background:#dcf8eb;color:#167b58}.rt-chip-muted{background:#eef1f5;color:#596a7c}.rt-profile-form{display:grid;gap:16px}.rt-field-grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:6px 16px}.rt-field-grid label{font-weight:700;color:#34536d}.rt-field-grid input{margin-top:7px;background:#fff}.rt-photo-picker{padding:15px;border:1px dashed #8cb7ca;border-radius:15px;background:linear-gradient(135deg,#f0fbfc,#f7f6ff)}.rt-photo-picker label{font-weight:800;color:#285c79}.rt-photo-picker input[type=file]{background:#fff;border:1px solid #d6e4ee;padding:10px}.rt-photo-picker small{display:block;color:#61778a}.rt-savebar{display:flex;flex-wrap:wrap;align-items:center;gap:10px}.rt-savebar button{min-width:190px;font-weight:800;padding:12px 18px}.rt-secondary-action{background:#edf4f9;color:#31536f;box-shadow:none;border:1px solid #dce7f0}.rt-profile-links{display:flex;flex-wrap:wrap;gap:10px;margin-top:18px}.rt-profile-links a{background:transparent;color:#28668e;box-shadow:none;padding:4px 0}.rt-profile-form input:focus-visible,button:focus-visible,a:focus-visible{outline:3px solid #39bda9;outline-offset:3px}@media(max-width:620px){.rt-profile-head{align-items:flex-start;gap:12px}.rt-profile-avatar{border-radius:20px}.rt-field-grid{grid-template-columns:minmax(0,1fr)}.rt-savebar{align-items:stretch;flex-direction:column}.rt-savebar button,.rt-savebar a{width:100%;text-align:center}.rt-profile-panel{padding:14px}}@media(prefers-reduced-motion:reduce){*,*:before,*:after{scroll-behavior:auto!important;transition:none!important;animation:none!important}}</style><header class="brand-header"><div class="brand-lockup"><img class="brand-icon" src="/static/resimanage-3d-icon.jpg" width="64" height="64" alt="ResiManage"><div><h2>{{t('app')}}</h2><p class="brand-tagline">{% if lang=='ar' %}تسيير تربصات الأطباء المقيمين{% elif lang=='fr' %}Gestion des stages des médecins résidents{% else %}Resident rotation management{% endif %}</p><span class="brand-year">{{t('year')}}: {{year}}</span></div></div></header><div class="rt-app-shell">{% if user %}<nav class="main-nav" aria-label="{% if lang=='ar' %}التنقل الرئيسي{% elif lang=='fr' %}Navigation principale{% else %}Main navigation{% endif %}">
<a class="nav-home {% if request.path=='/' %}is-active{% endif %}" href="/"><span class="nav-icon">⌂</span><span>{{t('home')}}</span></a>
<details class="nav-group" {% if request.path in ['/residents','/objectives','/import','/exams','/dems'] %}open{% endif %}><summary><span class="nav-icon">♙</span><span>{{t('residents')}}</span></summary><div class="nav-menu"><a href="/residents" class="{% if request.path=='/residents' %}is-active{% endif %}">{{t('manage_residents')}}</a><a href="/objectives" class="{% if request.path=='/objectives' %}is-active{% endif %}">{{t('objectives')}}</a><a href="/import" class="{% if request.path=='/import' %}is-active{% endif %}">{{t('import_residents')}}</a><a href="/exams" class="{% if request.path=='/exams' %}is-active{% endif %}">{{t('promotion')}}</a><a href="/dems" class="{% if request.path=='/dems' %}is-active{% endif %}">{{t('dems_success')}}</a></div></details>
<details class="nav-group" {% if request.path in ['/terrains','/capacities','/assign','/history','/leave','/external-rotations'] %}open{% endif %}><summary><span class="nav-icon">⇄</span><span>{% if lang=='ar' %}التربصات والتخطيط{% elif lang=='fr' %}Stages et planification{% else %}Rotations & planning{% endif %}</span></summary><div class="nav-menu"><a href="/terrains" class="{% if request.path=='/terrains' %}is-active{% endif %}">{{t('terrains')}}</a><a href="/capacities" class="{% if request.path=='/capacities' %}is-active{% endif %}">{{t('capacities')}}</a><a href="/assign" class="{% if request.path=='/assign' %}is-active{% endif %}">{{t('assign')}}</a><a href="/history" class="{% if request.path=='/history' %}is-active{% endif %}">{{t('history')}}</a><a href="/leave" class="{% if request.path=='/leave' %}is-active{% endif %}">{% if lang=='ar' %}الإجازات{% elif lang=='fr' %}Congés{% else %}Annual leave{% endif %}</a><a href="/external-rotations" class="{% if request.path.startswith('/external-rotations') %}is-active{% endif %}">{% if lang=='ar' %}التربصات الخارجية{% elif lang=='fr' %}Stages externes{% else %}External rotations{% endif %}</a></div></details>
<details class="nav-group" {% if request.path in ['/decisions','/summary'] %}open{% endif %}><summary><span class="nav-icon">▤</span><span>{% if lang=='ar' %}المتابعة والتقارير{% elif lang=='fr' %}Suivi et rapports{% else %}Monitoring & reports{% endif %}</span></summary><div class="nav-menu"><a href="/decisions" class="{% if request.path=='/decisions' %}is-active{% endif %}">{{t('decisions')}}</a><a href="/summary" class="{% if request.path=='/summary' %}is-active{% endif %}">{{t('summary')}}</a></div></details>
<details class="nav-group" {% if request.path in ['/settings','/users','/audit','/backup','/export','/restore'] %}open{% endif %}><summary><span class="nav-icon">⚙</span><span>{{t('settings')}}</span></summary><div class="nav-menu"><a href="/settings" class="{% if request.path=='/settings' %}is-active{% endif %}">{{t('settings_page')}}</a><a href="/users" class="{% if request.path=='/users' %}is-active{% endif %}">{{t('users')}}</a><a href="/audit" class="{% if request.path=='/audit' %}is-active{% endif %}">{{t('audit')}}</a><a href="/backup" class="{% if request.path=='/backup' %}is-active{% endif %}">{{t('backup')}}</a><a href="/export" class="{% if request.path=='/export' %}is-active{% endif %}">{{t('export')}}</a><a href="/restore" class="{% if request.path=='/restore' %}is-active{% endif %}">{{t('restore')}}</a><a href="/logout">{{t('logout')}}</a></div></details>
<div class="rt-language" aria-label="Language"><a href="/lang/ar" {% if lang=='ar' %}aria-current="page"{% endif %} title="العربية">ع</a><a href="/lang/fr" {% if lang=='fr' %}aria-current="page"{% endif %} title="Français">FR</a><a href="/lang/en" {% if lang=='en' %}aria-current="page"{% endif %} title="English">EN</a></div></nav>{% endif %}<main>{% for m in get_flashed_messages() %}<p class="msg">{{m}}</p>{% endfor %}{{body|safe}}</main></div><footer class="rt-footer">© <strong>Tous droits réservés · ResiManage · Dr Dris</strong><span class="rt-sep"> | </span><strong>All rights reserved · ResiManage · Dr Dris</strong><span class="rt-sep"> | </span><strong>جميع الحقوق محفوظة · ResiManage · Dr Dris</strong></footer></html>"""
 context={"t":tr,"lang":lang,"year":setting(db(),"year")}
 context.update(kw)
 context.setdefault("role",session.get("role",""))
 context.setdefault("csrf_token",csrf_token)
 inner=render_template_string(body,**context)
 return render_template_string(shell,body=inner,lang=lang,t=tr,year=setting(db(),"year"),user=bool(session.get("user")),role=session.get("role",""))
def guard():
 username=session.get("user")
 if not username: return redirect("/login")
 c=db()
 account=c.execute("SELECT role,active FROM users WHERE username=?",(username,)).fetchone()
 c.close()
 if not account or not account["active"]:
  session.clear()
  return redirect("/login")
 # Reflect the current role from the database instead of trusting stale session data.
 session["role"]=account["role"]
def super_guard():
 denied=guard()
 if denied: return denied
 if session.get("role")!="super_admin": return ("Forbidden — Super Admin only",403)
@app.get("/lang/<lang>")
def lang(lang):
 if lang in LANGS: session["lang"]=lang
 return redirect(request.referrer or "/login")
@app.route("/login",methods=["GET","POST"])
def login():
 c=db(); p=setting(c,"password")
 if request.method=="POST":
  username=request.form.get("username","admin").strip()
  password=request.form.get("password","")
  account=c.execute("SELECT * FROM users WHERE username=? AND active=1",(username,)).fetchone()
  if not p:
   confirm=request.form.get("confirm_password","")
   if username!="admin" or len(password)<12 or password!=confirm:
    c.close(); flash("First setup requires username admin and matching password of at least 12 characters.")
    return redirect("/login")
   p=hash_password(password)
   c.execute("INSERT INTO settings(key,value) VALUES('password',?)",(p,))
   c.execute("INSERT OR IGNORE INTO users(username,password_hash,role) VALUES('admin',?,'super_admin')",(p,))
   c.commit(); c.close(); session.clear(); session["user"]="admin"; session["role"]="super_admin"; return redirect("/")
  if account and verify_password(password,account["password_hash"]):
   c.close(); session.clear(); session["user"]=account["username"]; session["role"]=account["role"]; return redirect("/")
  if username=="admin" and not account and verify_password(password,p):
   secure_hash=hash_password(password)
   c.execute("UPDATE settings SET value=? WHERE key='password'",(secure_hash,))
   c.execute("INSERT OR IGNORE INTO users(username,password_hash,role) VALUES('admin',?,'super_admin')",(secure_hash,))
   c.commit(); c.close(); session.clear(); session["user"]="admin"; session["role"]="super_admin"; return redirect("/")
  c.close(); flash(tr("badpass"))
 else:
  c.close()
 setup_required=not p
 return page("""<section class="login-hero"><div class="login-language"><span>🌐</span> <a href="/lang/ar">العربية</a><a href="/lang/fr">🇫🇷 Français</a><a href="/lang/en">🇬🇧 English</a></div><div class="login-brand"><div class="login-emblem" aria-hidden="true">💓</div><p class="login-kicker">CHU Nedir Mohamed · Tizi Ouzou</p><h1>Resi<span>Manage</span></h1><p class="login-subtitle">Gestion des rotations des résidents</p><p class="login-motto">Aujourd’hui, un pas de plus vers votre avenir</p></div><div class="login-panel"><h2>{{t('login')}}</h2><p class="login-hint">Accédez à votre espace sécurisé</p><form method="post"><label>Username / Identifiant / اسم المستخدم</label><input name="username" value="admin" autocomplete="username" required><label>{{t('password')}}</label><div class="password-wrap"><input id="login-password" name="password" type="password" autocomplete="current-password" minlength="{{12 if setup_required else 1}}" required><button class="password-toggle" type="button" onclick="const p=document.getElementById('login-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'Show password':'Hide password')">◉</button></div>{% if setup_required %}<label>Confirm password / Confirmer le mot de passe / تأكيد كلمة المرور</label><input name="confirm_password" type="password" autocomplete="new-password" minlength="12" required><p>Create the first administrator password (12+ characters).</p>{% endif %}<button class="login-submit">{{'Create administrator / Créer l’administrateur / إنشاء المسؤول' if setup_required else t('signin')}} <span>➜</span></button></form><p class="login-security">♧ Connexion sécurisée · Secure sign-in · دخول آمن</p><p class="login-security"><a href="/register" style="color:#8df2df;text-decoration:underline">Créer un compte / Create account / إنشاء حساب</a></p></div><div class="login-feature-row"><span>🩺 Médecine</span><span>💓 Réanimation</span><span>✨ Anesthésie</span><span>📋 Formation</span></div><footer class="login-rights">Tous droits réservés · Dr Dris <span>|</span> All rights reserved · Dr Dris <span>|</span> جميع الحقوق محفوظة · Dr Dris</footer></section><style>.login-hero{position:relative;isolation:isolate;overflow:hidden;max-width:980px;margin:12px auto;padding:clamp(16px,4vw,38px);border-radius:28px;color:#f8fbff;background:radial-gradient(ellipse at 12% 15%,#00dff755,transparent 32%),radial-gradient(ellipse at 88% 24%,#3278ff55,transparent 30%),radial-gradient(ellipse at 80% 85%,#19dfbd55,transparent 30%),linear-gradient(135deg,#020b22,#073d85 48%,#041426);box-shadow:0 24px 60px #152d4c35}.login-hero:before,.login-hero:after{content:'';position:absolute;z-index:-1;border-radius:50%;filter:blur(2px);animation:orbit 12s ease-in-out infinite alternate}.login-hero:before{width:230px;height:230px;right:-55px;top:-75px;background:linear-gradient(135deg,#47e4d8aa,#4c8bff44)}.login-hero:after{width:190px;height:190px;left:-65px;bottom:90px;background:linear-gradient(135deg,#00eaffaa,#087dff55);animation-delay:-4s}.login-language{display:flex;justify-content:center;align-items:center;gap:clamp(8px,2vw,20px);flex-wrap:wrap}.login-language a{color:#fff;text-decoration:none;border:1px solid #ffffff45;background:#ffffff12;border-radius:22px;padding:7px 12px}.login-brand{text-align:center;padding:20px 0 14px}.login-emblem{font-size:56px;color:#53e2c4;text-shadow:0 0 28px #53e2c477;animation:pulse 3s ease-in-out infinite}.login-brand h1{font-size:clamp(38px,8vw,68px);line-height:1;margin:4px 0;font-weight:800;letter-spacing:-2px}.login-brand h1 span{background:linear-gradient(90deg,#19d9e8,#58a7ff,#19dfbd);-webkit-background-clip:text;background-clip:text;color:transparent}.login-kicker,.login-subtitle,.login-motto{margin:7px 0}.login-kicker{color:#b8d8f0}.login-subtitle{font-size:clamp(15px,3vw,20px)}.login-motto{color:#70e8f2;font-style:italic}.login-panel{max-width:480px;margin:16px auto 10px;padding:clamp(18px,4vw,30px);border:1px solid #ffffff55;border-radius:22px;background:#ffffff13;backdrop-filter:blur(14px);box-shadow:0 15px 35px #050d2040}.login-panel h2{text-align:center;margin:0;font-size:26px}.login-hint,.login-security{text-align:center;color:#d0e3f4}.login-panel label{display:block;margin-top:12px}.login-panel input{width:100%;min-width:0;background:#ffffff17;color:white;border:1px solid #ffffff48;border-radius:14px;padding:13px}.login-panel input::placeholder{color:#d4e2f1}.password-wrap{display:flex;align-items:center;gap:8px}.password-wrap input{flex:1}.password-toggle{border:1px solid #ffffff55;background:#ffffff18;color:white;padding:10px 14px;border-radius:12px;cursor:pointer}.login-submit{width:100%;margin-top:18px;border:0;border-radius:15px;padding:14px;background:linear-gradient(100deg,#5ce7b0,#22c9e5 55%,#6697ff);color:#08223d;font-weight:800;font-size:17px;box-shadow:0 7px 22px #31d5cf44;cursor:pointer}.login-submit span{margin-inline-start:10px}.login-feature-row{display:flex;justify-content:center;gap:10px;flex-wrap:wrap;margin:22px 0 12px}.login-feature-row span{border:1px solid #ffffff35;background:#ffffff10;border-radius:12px;padding:9px 13px}.login-rights{text-align:center;color:#d6e5f3;border-top:1px solid #ffffff24;padding-top:13px;font-size:12px}.login-rights span{color:#45c7e8;padding:0 5px}@keyframes orbit{to{transform:translate3d(-18px,22px,0) scale(1.08)}}@keyframes pulse{50%{transform:scale(1.08);filter:drop-shadow(0 0 12px #5ce7b0)}}@media(prefers-reduced-motion:reduce){.login-hero:before,.login-hero:after,.login-emblem{animation:none}}@media(max-width:520px){.login-hero{margin:4px;padding:14px;border-radius:18px}.login-feature-row span{font-size:12px;padding:7px}.login-rights span{padding:0 2px}}</style>""",setup_required=setup_required)

def send_verification_email(recipient, code):
 host=os.environ.get("RESITRACK_SMTP_HOST","").strip()
 port=int(os.environ.get("RESITRACK_SMTP_PORT","587"))
 sender=os.environ.get("RESITRACK_SMTP_FROM","").strip()
 username=os.environ.get("RESITRACK_SMTP_USER","").strip()
 password=os.environ.get("RESITRACK_SMTP_PASSWORD","")
 if not host or not sender or not username or not password:
  raise RuntimeError("Email verification is not configured. Set RESITRACK_SMTP_HOST, RESITRACK_SMTP_PORT, RESITRACK_SMTP_FROM, RESITRACK_SMTP_USER and RESITRACK_SMTP_PASSWORD.")
 from email.message import EmailMessage
 message=EmailMessage()
 message["Subject"]="ResiManage — email verification code"
 message["From"]=sender
 message["To"]=recipient
 message.set_content("Your ResiManage verification code is: "+code+"\n\nIt expires in 10 minutes. If you did not request this, ignore this message.")
 with smtplib.SMTP(host,port,timeout=15) as server:
  server.starttls()
  server.login(username,password)
  server.send_message(message)

@app.route("/register",methods=["GET","POST"])
def register():
 if session.get("user"): return redirect("/")
 c=db()
 if request.method=="POST":
  username=request.form.get("username","").strip()
  email=request.form.get("email","").strip().lower()
  password=request.form.get("password","")
  confirm=request.form.get("confirm_password","")
  if not re.fullmatch(r"[A-Za-z0-9_.-]{3,40}",username):
   flash("Username must be 3–40 letters, numbers, dots, underscores or hyphens.")
  elif not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+",email) or len(email)>254:
   flash("Enter a valid email address / أدخل بريدًا إلكترونيًا صالحًا.")
  elif len(password)<12 or password!=confirm:
   flash("Password must be 12+ characters and match confirmation / كلمة المرور يجب أن تكون 12 حرفًا على الأقل.")
  elif c.execute("SELECT 1 FROM users WHERE username=?",(username,)).fetchone() or c.execute("SELECT 1 FROM users WHERE lower(email)=?",(email,)).fetchone() or c.execute("SELECT 1 FROM pending_registrations WHERE username=? OR email=?",(username,email)).fetchone():
   flash("Username or email already registered / اسم المستخدم أو البريد مسجل مسبقًا.")
  else:
   code=("%06d" % secrets.randbelow(1000000))
   c.execute("DELETE FROM pending_registrations WHERE expires_at<?",(int(time.time()),))
   try:
    c.execute("INSERT INTO pending_registrations(username,email,password_hash,code_hash,expires_at) VALUES(?,?,?,?,?)",(username,email,hash_password(password),hashlib.sha256(code.encode()).hexdigest(),int(time.time())+600))
    c.commit()
    try:
     send_verification_email(email,code)
    except Exception as exc:
     c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
     app.logger.error("ResiManage verification email failed: %s",exc)
     flash("Email could not be sent; SMTP setup is required / تعذر إرسال البريد، يلزم إعداد SMTP.")
    else:
     session["pending_email"]=email
     flash("Verification code sent; valid for 10 minutes / أرسلنا رمز التأكيد، صالح لمدة 10 دقائق.")
     c.close()
     return redirect("/verify-email")
   except sqlite3.IntegrityError:
    flash("Username or email already registered / اسم المستخدم أو البريد مسجل مسبقًا.")
 c.close()
 return page("""<section class="login-hero"><div class="login-brand"><div class="login-emblem">💓</div><h1>Resi<span>Manage</span></h1><h2>Create account / إنشاء حساب</h2><p class="login-subtitle">A valid email is required to confirm the account.</p></div><div class="login-panel"><form method="post"><label>Username / اسم المستخدم</label><input name="username" minlength="3" maxlength="40" pattern="[A-Za-z0-9_.-]{3,40}" autocomplete="username" required><label>Email / البريد الإلكتروني</label><input name="email" type="email" maxlength="254" autocomplete="email" required><label>Password (12+) / كلمة المرور</label><input name="password" type="password" minlength="12" autocomplete="new-password" required><label>Confirm password / تأكيد كلمة المرور</label><input name="confirm_password" type="password" minlength="12" autocomplete="new-password" required><button class="login-submit">Create account / إنشاء الحساب</button></form><p class="login-security"><a href="/login">Back to sign in / الرجوع إلى الدخول</a></p></div></section>""")

@app.route("/verify-email",methods=["GET","POST"])
def verify_email():
 email=session.get("pending_email","").strip().lower()
 if not email: return redirect("/register")
 c=db()
 if request.method=="POST":
  code=request.form.get("code","").strip()
  pending=c.execute("SELECT * FROM pending_registrations WHERE email=?",(email,)).fetchone()
  if not pending:
   session.pop("pending_email",None); flash("Registration expired; please register again / أعد التسجيل."); c.close(); return redirect("/register")
  if pending["expires_at"]<int(time.time()):
   c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
   session.pop("pending_email",None); flash("Code expired; register again / انتهت صلاحية الرمز."); c.close(); return redirect("/register")
  if pending["attempts"]>=5:
   c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
   session.pop("pending_email",None); flash("Too many attempts; register again / تجاوزت عدد المحاولات."); c.close(); return redirect("/register")
  if not re.fullmatch(r"\d{6}",code) or not hmac.compare_digest(hashlib.sha256(code.encode()).hexdigest(),pending["code_hash"]):
   c.execute("UPDATE pending_registrations SET attempts=attempts+1 WHERE email=?",(email,)); c.commit()
   flash("Incorrect verification code / رمز غير صحيح.")
  else:
   try:
    c.execute("INSERT INTO users(username,password_hash,role,email,email_verified) VALUES(?,?,'user',?,1)",(pending["username"],pending["password_hash"],pending["email"]))
    c.execute("DELETE FROM pending_registrations WHERE email=?",(email,))
    audit_event(c,"user.self_registered",f"username={pending['username']}; email_verified=1")
    c.commit()
    session.pop("pending_email",None)
    flash("Email verified; account created. Please sign in / تم تأكيد البريد وإنشاء الحساب.")
    c.close()
    return redirect("/login")
   except sqlite3.IntegrityError:
    c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
    flash("Account already exists / الحساب موجود مسبقًا.")
 c.close()
 return page("""<section class="login-hero"><div class="login-brand"><div class="login-emblem">✉</div><h1>Resi<span>Manage</span></h1><h2>Email verification / تأكيد البريد</h2><p class="login-subtitle">Code sent to {{email}}</p></div><div class="login-panel"><form method="post"><label>6-digit code / رمز التأكيد</label><input name="code" inputmode="numeric" pattern="[0-9]{6}" minlength="6" maxlength="6" autocomplete="one-time-code" required><button class="login-submit">Verify email / تأكيد البريد</button></form><p class="login-security">Code expires after 10 minutes / تنتهي صلاحية الرمز بعد 10 دقائق.</p></div></section>""",email=email)

@app.get("/logout")
def logout(): session.clear(); return redirect("/login")
@app.get("/")
def home():
 if guard(): return redirect("/login")
 c=db(); y=setting(c,"year"); rs=c.execute("SELECT * FROM residents WHERE status='active' ORDER BY year DESC,surname").fetchall()
 count=c.execute("SELECT COUNT(*) FROM terrains").fetchone()[0]; confirmed=c.execute("SELECT COUNT(*) FROM assignments a JOIN periods p ON p.id=a.period_id WHERE p.year_label=? AND a.status='confirmed'",(y,)).fetchone()[0]
 rows=[]
 for r in rs:
  n=c.execute("SELECT COUNT(*) FROM assignments a JOIN periods p ON p.id=a.period_id WHERE a.resident_id=? AND p.year_label=? AND a.status='confirmed'",(r["id"],y)).fetchone()[0]; rows.append((r,n))
 c.close()
 return page("""<section><h3 class="rt-dashboard-title"><span class="rt-dot"></span>{{t('overview')}}</h3><p>{{t('priority')}}</p><div class="rt-stats"><div class="rt-stat"><span>{{t('residents')}}</span><strong>{{rows|length}}</strong></div><div class="rt-stat"><span>{{t('terrains')}}</span><strong>{{count}}</strong></div><div class="rt-stat"><span>{{t('confirmed')}}</span><strong>{{confirmed}}</strong></div><div class="rt-stat"><span>{{t('year')}}</span><strong style="font-size:17px">{{year}}</strong></div></div></section><div class="rt-quick-actions"><a class="rt-action-card" href="/leave"><span class="rt-action-icon">◷</span><span><strong>الإجازات / Congés / Leave</strong><small>30 days yearly · carry forward</small></span></a><a class="rt-action-card" href="/external-rotations"><span class="rt-action-icon">✈</span><span><strong>التربصات الخارجية / Stages à l’étranger / External rotations</strong><small>Supervisor · hospital · country · dates</small></span></a><a class="rt-action-card" href="/residents"><span class="rt-action-icon">♙</span><span><strong>{{t('manage_residents')}}</strong><small>Profils · années · coordonnées</small></span></a><a class="rt-action-card" href="/assign"><span class="rt-action-icon">⇄</span><span><strong>{{t('assign')}}</strong><small>Répartition des stages</small></span></a><a class="rt-action-card" href="/terrains"><span class="rt-action-icon">▥</span><span><strong>{{t('terrains')}}</strong><small>Terrains obligatoires et optionnels</small></span></a><a class="rt-action-card" href="/import"><span class="rt-action-icon">⇧</span><span><strong>{{t('import_residents')}}</strong><small>Import avec aperçu et validation</small></span></a></div><section><h3 class="rt-dashboard-title"><span class="rt-dot"></span>{{t('four')}}</h3><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('count')}}</th></tr>{% for r,n in rows %}<tr><td><a class="rt-resident-link" href="/profile/{{r.id}}">{{r.surname}} {{r.firstname}}</a></td><td><span class="rt-year-chip">{{r.year}}</span></td><td><span class="rt-progress"><span style="width:{{[n*25,100]|min}}%"></span></span> {{n}} / 4</td></tr>{% else %}<tr><td colspan="3">{{t('none')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows,count=count,confirmed=confirmed,year=y)
@app.route("/residents",methods=["GET","POST"])
def residents():
 if guard(): return redirect("/login")
 c=db()
 if request.method=="POST":
  if request.form.get("action")=="edit":
   if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
   try:
    rid=int(request.form["rid"])
    year=int(request.form.get("year","0"))
    repeat_year=request.form.get("repeat_year","0")
    status=request.form.get("status","")
   except (TypeError,ValueError):
    flash("Invalid resident update.")
   else:
    target=c.execute("SELECT status,year FROM residents WHERE id=?",(rid,)).fetchone()
    if not target:
     flash("Resident not found.")
    elif target["status"]=="graduated":
     flash("Graduated residents are archived; status cannot be changed here.")
    elif year not in (1,2,3,4) or repeat_year not in ("0","1") or status not in ("active","inactive"):
     flash("Year, repeat flag or status is invalid. DEMS graduation must use the official archive workflow.")
    elif year!=target["year"]:
     flash("Study-year changes must use the official exam/promotion workflow.")
    else:
     c.execute("UPDATE residents SET year=?,repeat_year=?,status=? WHERE id=?",(year,int(repeat_year),status,rid))
     audit_event(c,"resident.updated",f"resident_id={rid}; year={year}; status={status}")
     c.commit(); flash(tr("done"))
  else:
   if session.get("role")!="super_admin":
    c.close(); return ("Forbidden — Super Admin only",403)
   surname=request.form.get("surname","").strip()
   firstname=request.form.get("firstname","").strip()
   phone=request.form.get("phone","").strip()
   repeat_year=request.form.get("repeat_year","0")
   try: year=int(request.form.get("year","0"))
   except (TypeError,ValueError): year=0
   if not surname or not firstname or len(surname)>100 or len(firstname)>100 or year not in (1,2,3,4):
    flash("Surname and first name are required (max 100 characters) and a valid study year is required.")
   elif len(phone)>40 or repeat_year not in ("0","1"):
    flash("Phone must be at most 40 characters and repeat flag must be 0 or 1.")
   elif c.execute("SELECT 1 FROM residents WHERE lower(surname)=lower(?) AND lower(firstname)=lower(?) AND year=?",(surname,firstname,year)).fetchone():
    flash("A resident with the same name and study year already exists.")
   else:
    c.execute("INSERT INTO residents(surname,firstname,year,phone,repeat_year,status) VALUES(?,?,?,?,?,?)",(surname,firstname,year,phone,int(repeat_year),"active"))
    audit_event(c,"resident.created",surname+" "+firstname)
    c.commit(); flash(tr("done"))
 rows=c.execute("SELECT * FROM residents ORDER BY status DESC,year DESC,surname").fetchall(); c.close()
 return page("""<section><h3>{{t('add')}} — {{t('resident')}}</h3><form method="post"><div class="grid"><div><label>{{t('surname')}}</label><input name="surname" required></div><div><label>{{t('firstname')}}</label><input name="firstname" required></div><div><label>{{t('study_year')}}</label><select name="year">{% for y in range(1,5) %}<option>{{y}}</option>{% endfor %}</select></div><div><label>{{t('phone')}}</label><input name="phone"></div></div><label><input type="checkbox" name="repeat_year" value="1"> {{t('repeat')}}</label><button>{{t('add')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('save')}}</th></tr>{% for r in rows %}<tr><td><a href="/profile/{{r.id}}">{{r.surname}} {{r.firstname}}</a></td><td>{{r.year}}</td><td>{% if r.status=='graduated' %}<span>Graduated — DEMS archive</span>{% else %}<form method="post"><input type="hidden" name="action" value="edit"><input type="hidden" name="rid" value="{{r.id}}"><span>Year {{r.year}}</span><input type="hidden" name="year" value="{{r.year}}"><label><input type="checkbox" name="repeat_year" value="1" {% if r.repeat_year %}checked{% endif %}> {{t('repeat')}}</label><select name="status"><option value="active" {% if r.status=='active' %}selected{% endif %}>{{t('active')}}</option><option value="inactive" {% if r.status=='inactive' %}selected{% endif %}>{{t('inactive')}}</option></select><button>{{t('save')}}</button></form>{% endif %}</td></tr>{% endfor %}</table></div></section>""",rows=rows)
@app.route("/terrains",methods=["GET","POST"])
def terrains():
 if guard(): return redirect("/login")
 c=db()
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  action=request.form.get("action","create")
  try:
   if action=="edit":
    tid=int(request.form["tid"]); name=request.form["name"].strip(); kind=request.form["kind"]; mandatory_year=int(request.form["mandatory_year"])
    if not name or kind not in ("mandatory","basic","optional") or mandatory_year not in (0,1,2,3,4): raise ValueError("Invalid terrain fields")
    current=c.execute("SELECT * FROM terrains WHERE id=?",(tid,)).fetchone()
    if not current: raise ValueError("Terrain not found.")
    has_records=c.execute("SELECT (SELECT COUNT(*) FROM history WHERE terrain_id=?)+(SELECT COUNT(*) FROM assignments WHERE terrain_id=?)",(tid,tid)).fetchone()[0]>0
    if has_records and (kind!=current["kind"] or mandatory_year!=current["mandatory_year"]):
     raise ValueError("Terrain classification/year cannot be changed after rotation history or assignments exist. Preserve the audit trail and correct it before use or via a reviewed migration.")
    c.execute("UPDATE terrains SET name=?,kind=?,mandatory_year=? WHERE id=?",(name,kind,mandatory_year,tid))
    audit_event(c,"terrain.updated",f"terrain_id={tid}; name={name}; kind={kind}; mandatory_year={mandatory_year}")
   elif action in ("deactivate","activate"):
    tid=int(request.form["tid"]); active=0 if action=="deactivate" else 1
    current=c.execute("SELECT * FROM terrains WHERE id=?",(tid,)).fetchone()
    if not current: raise ValueError("Terrain not found.")
    c.execute("UPDATE terrains SET active=? WHERE id=?",(active,tid))
    audit_event(c,"terrain.deactivated" if not active else "terrain.activated",f"terrain_id={tid}")
   else:
    name=request.form["name"].strip(); kind=request.form["kind"]; mandatory_year=int(request.form["mandatory_year"])
    if not name or kind not in ("mandatory","basic","optional") or mandatory_year not in (0,1,2,3,4): raise ValueError("Invalid terrain fields")
    c.execute("INSERT INTO terrains(name,kind,mandatory_year) VALUES(?,?,?)",(name,kind,mandatory_year))
    audit_event(c,"terrain.created",name+" kind="+kind)
   c.commit(); flash(tr("done"))
  except sqlite3.IntegrityError: c.rollback(); flash("Already exists / Existe déjà / المصلحة موجودة")
  except ValueError as exc: c.rollback(); flash(str(exc))
 rows=c.execute("SELECT * FROM terrains ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); c.close()
 return page("""<section><h3>{{t('add')}} — {{t('terrain')}}</h3><form method="post"><input type="hidden" name="action" value="create"><label>{{t('terrain')}}</label><input name="name" required><div class="grid"><div><label>{{t('kind')}}</label><select name="kind"><option value="mandatory">{{t('mandatory')}}</option><option value="basic">{{t('basic')}}</option><option value="optional">{{t('optional')}}</option></select></div><div><label>{{t('mandatory_year')}}</label><select name="mandatory_year"><option value="0">{{t('any')}}</option>{% for y in range(1,5) %}<option value="{{y}}">{{y}}</option>{% endfor %}</select></div></div><button>{{t('save')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('terrain')}}</th><th>{{t('kind')}}</th><th>{{t('mandatory_year')}}</th><th>{{t('status')}}</th><th>Admin actions</th></tr>{% for r in rows %}<tr><td colspan="5"><form method="post" class="grid"><input type="hidden" name="action" value="edit"><input type="hidden" name="tid" value="{{r.id}}"><input name="name" value="{{r.name}}" required><select name="kind"><option value="mandatory" {% if r.kind=='mandatory' %}selected{% endif %}>{{t('mandatory')}}</option><option value="basic" {% if r.kind=='basic' %}selected{% endif %}>{{t('basic')}}</option><option value="optional" {% if r.kind=='optional' %}selected{% endif %}>{{t('optional')}}</option></select><select name="mandatory_year"><option value="0" {% if not r.mandatory_year %}selected{% endif %}>{{t('any')}}</option>{% for y in range(1,5) %}<option value="{{y}}" {% if r.mandatory_year==y %}selected{% endif %}>{{y}}</option>{% endfor %}</select><span>{{'Active' if r.active else 'Inactive'}}</span><button>{{t('save')}}</button></form>{% if role=='super_admin' %}<form method="post"><input type="hidden" name="tid" value="{{r.id}}"><button name="action" value="{{'deactivate' if r.active else 'activate'}}">{{'Deactivate' if r.active else 'Reactivate'}}</button></form>{% endif %}</td></tr>{% endfor %}</table></div><p>Deactivation hides a terrain from new assignments but preserves all past history.</p></section>""",rows=rows)
@app.route("/capacities",methods=["GET","POST"])
def capacities():
 if guard(): return redirect("/login")
 c=db(); y=setting(c,"year")
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  try:
   tid=int(request.form["terrain"]); pid=int(request.form["period"]); places=int(request.form["places"])
   if places<0: raise ValueError("Capacity cannot be negative.")
   terrain=c.execute("SELECT id FROM terrains WHERE id=? AND active=1",(tid,)).fetchone()
   period=c.execute("SELECT id FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
   if not terrain: raise ValueError("Choose an active terrain.")
   if not period: raise ValueError("Choose a rotation period from the selected academic year.")
   used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status IN ('confirmed','pending')",(tid,pid)).fetchone()[0]
   if places<used: flash(tr("full"))
   else:
    c.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,?) ON CONFLICT(year_label,terrain_id,period_id) DO UPDATE SET places=excluded.places",(y,tid,pid,places))
    audit_event(c,"capacity.changed",f"year={y}; terrain_id={tid}; period_id={pid}; places={places}")
    c.commit(); flash(tr("done"))
  except (ValueError,TypeError,KeyError):
   c.rollback(); flash("Invalid capacity input. Choose an active terrain, a period in the selected academic year, and a non-negative whole number of places.")
 ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); ps=c.execute("SELECT * FROM periods WHERE year_label=? ORDER BY id",(y,)).fetchall()
 rows=c.execute("SELECT t.name terrain,p.name period,cap.places,(SELECT COUNT(*) FROM assignments a WHERE a.terrain_id=t.id AND a.period_id=p.id AND a.status IN ('confirmed','pending')) used FROM capacities cap JOIN terrains t ON t.id=cap.terrain_id JOIN periods p ON p.id=cap.period_id WHERE cap.year_label=?",(y,)).fetchall(); c.close()
 return page("""<section><h3>{{t('capacities')}}</h3><form method="post"><div class="grid"><div><label>{{t('terrain')}}</label><select name="terrain">{% for x in ts %}<option value="{{x.id}}">{{x.name}}</option>{% endfor %}</select></div><div><label>{{t('period')}}</label><select name="period">{% for x in ps %}<option value="{{x.id}}">{{x.name}}</option>{% endfor %}</select></div><div><label>{{t('places')}}</label><input type="number" min="0" name="places" required></div></div><button>{{t('save')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('period')}}</th><th>{{t('terrain')}}</th><th>{{t('places')}}</th><th>{{t('used')}}</th><th>{{t('remaining')}}</th></tr>{% for r in rows %}<tr><td>{{r.period}}</td><td>{{r.terrain}}</td><td>{{r.places}}</td><td>{{r.used}}</td><td>{{r.places-r.used}}</td></tr>{% endfor %}</table></div></section>""",ts=ts,ps=ps,rows=rows)
def check(c,rid,tid,pid,y):
 r=c.execute("SELECT * FROM residents WHERE id=? AND status='active'",(rid,)).fetchone(); t=c.execute("SELECT * FROM terrains WHERE id=?",(tid,)).fetchone(); p=c.execute("SELECT * FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
 if not r or not t or not t["active"] or not p: return False,tr("none"),None
 start_day, end_day = rotation_period_bounds(p["year_label"], p["name"])
 external = c.execute("SELECT hospital,country,start_date,end_date FROM external_rotations WHERE resident_id=? AND status<>'cancelled' AND start_date<=? AND end_date>=? LIMIT 1",(rid,end_day.isoformat(),start_day.isoformat())).fetchone()
 if external: return False,("External rotation blocks local assignment: "+external["hospital"]+", "+external["country"]+" ("+external["start_date"]+" to "+external["end_date"]+")"),None
 dup=c.execute("SELECT t.name FROM assignments a JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=? AND a.period_id=? AND a.status IN ('confirmed','pending')",(rid,pid)).fetchone()
 if dup:return False,tr("duplicate")+dup["name"],None
 repeat = c.execute(
  "SELECT t.name,p.name AS period,p.year_label FROM assignments a JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.resident_id=? AND a.terrain_id=? AND a.period_id<>? AND a.status IN ('confirmed','pending') LIMIT 1",
  (rid,tid,pid)
 ).fetchone()
 total=c.execute(
  "SELECT COUNT(*) FROM terrains WHERE kind IN ('mandatory','basic') AND active=1"
 ).fetchone()[0]
 # Historical rows without an explicit completion date still need a completed period.
 # A future academic year with a blank date is not proof that the rotation happened.
 today=date.today()
 completed_terrains=set()
 history_rows=c.execute("""
  SELECT h.terrain_id,h.year_label,h.period,h.date_done
  FROM history h JOIN terrains t ON t.id=h.terrain_id
  WHERE h.resident_id=? AND t.kind IN ('mandatory','basic') AND t.active=1
 """,(rid,)).fetchall()
 for history_row in history_rows:
  done_date=str(history_row["date_done"] or "").strip()
  if done_date:
   try:
    if date.fromisoformat(done_date)<=today:
     completed_terrains.add(history_row["terrain_id"])
   except ValueError:
    continue
  else:
   try:
    _,history_period_end=rotation_period_bounds(history_row["year_label"],history_row["period"])
    if history_period_end<today:
     completed_terrains.add(history_row["terrain_id"])
   except (ValueError,TypeError):
    continue
 # A confirmed assignment is completion evidence only after its rotation period has ended.
 # Future/current assignments must not unlock repetition of mandatory terrains.
 completed_assignments=c.execute("""
  SELECT a.terrain_id,p.year_label,p.name FROM assignments a
  JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id
  WHERE a.resident_id=? AND a.status='confirmed'
    AND t.kind IN ('mandatory','basic') AND t.active=1
 """,(rid,)).fetchall()
 for completed in completed_assignments:
  _,period_end=rotation_period_bounds(completed["year_label"],completed["name"])
  if period_end<today:
   completed_terrains.add(completed["terrain_id"])
 done=len(completed_terrains)
 if repeat and done<total:
  return False,tr("same_terrain_year")+repeat["name"]+" — "+repeat["period"],None
 cap=c.execute("SELECT places FROM capacities WHERE year_label=? AND terrain_id=? AND period_id=?",(y,tid,pid)).fetchone()
 if not cap:return False,tr("capacity_first"),None
 used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status IN ('confirmed','pending')",(tid,pid)).fetchone()[0]
 if used>=cap["places"]:return False,tr("full"),None
 h=c.execute("SELECT year_label,period,date_done FROM history WHERE resident_id=? AND terrain_id=? ORDER BY id DESC LIMIT 1",(rid,tid)).fetchone()
 if h and done<total:return False,tr("history_block")+f" ({h['year_label']} / {h['period']} / {h['date_done']})",None
 if t["mandatory_year"] and t["mandatory_year"]!=r["year"]:return True,tr("mismatch"),"pending"
 return True,(tr("history_repeat") if h else tr("done")),"confirmed"
@app.route("/assign",methods=["GET","POST"])
def assign():
 if guard():return redirect("/login")
 c=db(); y=setting(c,"year")
 if request.method=="POST":
  if session.get("role")!="super_admin":
   c.close(); return ("Forbidden — Super Admin only",403)
  try:
   rid,tid,pid=map(int,(request.form["resident"],request.form["terrain"],request.form["period"])); ok,msg,status=check(c,rid,tid,pid,y)
   if not ok: flash(msg)
   else:
    old=c.execute("SELECT id,status FROM assignments WHERE resident_id=? AND period_id=?",(rid,pid)).fetchone()
    if old and old["status"] not in ("rejected","cancelled"): flash(tr("duplicate"))
    elif old:
     c.execute("UPDATE assignments SET terrain_id=?,status=? WHERE id=?",(tid,status,old["id"]))
     audit_event(c,"assignment.replaced",f"assignment_id={old['id']}; resident_id={rid}; terrain_id={tid}; status={status}")
     c.commit(); flash(msg)
    else:
     c.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,?)",(rid,tid,pid,status))
     audit_event(c,"assignment.created",f"resident_id={rid}; terrain_id={tid}; period_id={pid}; status={status}")
     c.commit(); flash(msg)
  except Exception: c.rollback(); flash("Invalid assignment / Affectation invalide / تعيين غير صالح")
 rs=c.execute("SELECT * FROM residents WHERE status='active' ORDER BY year DESC,surname").fetchall(); ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); ps=c.execute("SELECT * FROM periods WHERE year_label=? ORDER BY id",(y,)).fetchall()
 external_records=c.execute("SELECT resident_id,start_date,end_date FROM external_rotations WHERE status<>'cancelled'").fetchall()
 blocked_periods_by_resident={}
 for r in rs:
  blocked=[]
  for p in ps:
   pstart,pend=rotation_period_bounds(p["year_label"],p["name"])
   if any(e["resident_id"]==r["id"] and e["start_date"]<=pend.isoformat() and e["end_date"]>=pstart.isoformat() for e in external_records):
    blocked.append(str(p["id"]))
  blocked_periods_by_resident[r["id"]]=",".join(blocked)
 rows=c.execute("SELECT a.*,r.surname,r.firstname,r.year,t.name terrain,p.name period FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE p.year_label=? ORDER BY p.id,r.year DESC",(y,)).fetchall(); c.close()
 return page("""<section class="rt-assignment-hero"><h2>{% if lang=='ar' %}تخطيط التربصات وتعيين المقيمين{% elif lang=='fr' %}Planification et affectation des résidents{% else %}Rotation planning & resident assignment{% endif %}</h2><p>{{t('priority')}}</p><div class="rt-assignment-stats"><div><span>{% if lang=='ar' %}المقيمون النشطون{% elif lang=='fr' %}Résidents actifs{% else %}Active residents{% endif %}</span><strong>{{rs|length}}</strong></div><div><span>{% if lang=='ar' %}التخصصات / الأراضي{% elif lang=='fr' %}Terrains actifs{% else %}Active rotations{% endif %}</span><strong>{{ts|length}}</strong></div><div><span>{% if lang=='ar' %}الفترات المتاحة{% elif lang=='fr' %}Périodes{% else %}Periods{% endif %}</span><strong>{{ps|length}}</strong></div><div><span>{% if lang=='ar' %}التعيينات في السنة{% elif lang=='fr' %}Affectations cette année{% else %}Assignments this year{% endif %}</span><strong>{{rows|length}}</strong></div></div></section>{% if role=='super_admin' %}<section class="rt-assignment-form"><h3>{% if lang=='ar' %}تعيين جديد{% elif lang=='fr' %}Nouvelle affectation{% else %}Create an assignment{% endif %}</h3><p class="muted">{% if lang=='ar' %}الأولوية حسب السنة: الرابعة، ثم الثالثة، الثانية، فالأولى. يتم إظهار تعارضات التربص الخارجي باللون الأحمر.{% elif lang=='fr' %}Priorité : 4e année, puis 3e, 2e et 1re année. Les conflits de stages externes sont signalés en rouge.{% else %}Priority order: Year 4, then Year 3, Year 2, and Year 1. External rotation conflicts are marked in red.{% endif %}</p><label for="assignment-search">{% if lang=='ar' %}ابحث عن مقيم{% elif lang=='fr' %}Rechercher un résident{% else %}Find a resident{% endif %}</label><input id="assignment-search" type="search" placeholder="{% if lang=='ar' %}الاسم أو اللقب...{% elif lang=='fr' %}Nom ou prénom...{% else %}Name or surname...{% endif %}"><label for="assignment-year-filter">{% if lang=='ar' %}تصفية حسب السنة{% elif lang=='fr' %}Filtrer par année{% else %}Filter by year{% endif %}</label><select id="assignment-year-filter"><option value="all">{% if lang=='ar' %}كل السنوات{% elif lang=='fr' %}Toutes les années{% else %}All years{% endif %}</option>{% for y in [4,3,2,1] %}<option value="{{y}}">{% if lang=='ar' %}السنة{% elif lang=='fr' %}Année{% else %}Year{% endif %} {{y}}</option>{% endfor %}</select><form method="post" id="assignment-form"><label for="assignment-resident">{{t('resident')}}</label><select name="resident" id="assignment-resident" required>{% for r in rs %}<option value="{{r.id}}" data-year="{{r.year}}" data-blocked-periods="{{blocked_periods_by_resident[r.id]}}" data-base-label="{{r.year}} — {{r.surname}} {{r.firstname}}">{{r.year}} — {{r.surname}} {{r.firstname}}</option>{% endfor %}</select><p id="assignment-conflict" class="rt-assignment-notice" role="status" aria-live="polite"></p><label for="assignment-terrain">{{t('terrain')}}</label><select name="terrain" id="assignment-terrain" required>{% for x in ts %}<option value="{{x.id}}" data-kind="{{x.kind}}">{{x.name}}{% if x.kind=='optional' %} — {% if lang=='ar' %}اختياري{% elif lang=='fr' %}Optionnel{% else %}Optional{% endif %}{% endif %}</option>{% endfor %}</select><label for="assignment-period">{{t('period')}}</label><select name="period" id="assignment-period" required>{% for x in ps %}<option value="{{x.id}}">{{x.name}}</option>{% endfor %}</select><button type="submit">{{t('check_save')}}</button></form><script>(function(){const p=document.getElementById("assignment-period"),r=document.getElementById("assignment-resident"),search=document.getElementById("assignment-search"),year=document.getElementById("assignment-year-filter"),notice=document.getElementById("assignment-conflict");if(!p||!r)return;const options=Array.from(r.options);function update(){const pid=p.value,term=(search.value||"").toLocaleLowerCase(),yr=year.value;let visible=0;options.forEach(o=>{const blocked=(o.dataset.blockedPeriods||"").split(",").includes(pid),matchesName=(o.dataset.baseLabel||"").toLocaleLowerCase().includes(term),matchesYear=yr==="all"||o.dataset.year===yr;o.disabled=blocked||!matchesName||!matchesYear;o.hidden=!matchesName||!matchesYear;o.textContent=o.dataset.baseLabel+(blocked?" — ⚠ External rotation conflict":"");if(!o.disabled)visible++;});if(r.selectedOptions.length&&r.selectedOptions[0].disabled){const next=options.find(o=>!o.disabled);if(next)r.value=next.value;}const selected=r.selectedOptions[0],blocked=selected&&(selected.dataset.blockedPeriods||"").split(",").includes(pid);if(blocked){notice.className="rt-assignment-notice is-conflict";notice.textContent="{% if lang=='ar' %}تعارض: هذا المقيم لديه تربص خارجي في هذه الفترة.{% elif lang=='fr' %}Conflit : ce résident a un stage externe pendant cette période.{% else %}Conflict: this resident has an external rotation during this period.{% endif %}";}else{notice.className="rt-assignment-notice";notice.textContent=visible===0?"{% if lang=='ar' %}لا يوجد مقيم مطابق للفلاتر في هذه الفترة.{% elif lang=='fr' %}Aucun résident ne correspond aux filtres pour cette période.{% else %}No residents match these filters for this period.{% endif %}":"";}}p.addEventListener("change",update);search.addEventListener("input",update);year.addEventListener("change",update);update();})();</script></section>{% else %}<section><p>{% if lang=='ar' %}إنشاء التعيينات متاح للمسؤول العام فقط.{% elif lang=='fr' %}La création des affectations est réservée aux Super Admins.{% else %}Only Super Admins can create assignments.{% endif %}</p></section>{% endif %}<section><h3>{% if lang=='ar' %}التعيينات الحالية{% elif lang=='fr' %}Affectations actuelles{% else %}Current assignments{% endif %}</h3><div class="wrap"><table class="rt-assignment-table"><thead><tr><th>{{t('period')}}</th><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('terrain')}}</th><th>{{t('status')}}</th><th>{% if lang=='ar' %}إجراء إداري{% elif lang=='fr' %}Action admin{% else %}Admin action{% endif %}</th></tr></thead><tbody>{% for a in rows %}<tr class="{% if a.status=='pending' %}is-pending{% elif a.status=='rejected' or a.status=='cancelled' %}is-inactive{% else %}is-confirmed{% endif %}"><td>{{a.period}}</td><td>{{a.surname}} {{a.firstname}}</td><td><span class="rt-year-badge">{{a.year}}</span></td><td>{{a.terrain}}</td><td><span class="rt-status-pill status-{{a.status}}">{{t(a.status) if a.status in ['confirmed','pending','rejected'] else a.status}}</span></td><td>{% if role=='super_admin' and a.status not in ['cancelled','rejected'] %}<form method="post" action="/assignments/{{a.id}}/cancel"><button class="rt-danger-button" type="submit">{% if lang=='ar' %}إلغاء{% elif lang=='fr' %}Annuler{% else %}Cancel{% endif %}</button></form>{% else %}—{% endif %}</td></tr>{% else %}<tr><td colspan="6">{% if lang=='ar' %}لا توجد تعيينات بعد.{% elif lang=='fr' %}Aucune affectation pour le moment.{% else %}No assignments yet.{% endif %}</td></tr>{% endfor %}</tbody></table></div></section><style>.rt-assignment-hero h2{margin-top:0}.rt-assignment-stats{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px;margin-top:16px}.rt-assignment-stats>div{padding:12px;border:1px solid #155173;border-radius:13px;background:#071e33}.rt-assignment-stats span{display:block;font-size:12px;color:#a9ccdf}.rt-assignment-stats strong{display:block;font-size:25px;color:#27d8ce}.rt-assignment-notice{min-height:0;margin:0;padding:0;color:#a9ccdf}.rt-assignment-notice.is-conflict{margin:6px 0;padding:10px;border:1px solid #ff6d77;border-radius:10px;background:#531c2a;color:#ffe5e7;font-weight:700}.rt-assignment-table .is-pending td{border-top:2px solid #ffbd61}.rt-assignment-table .is-inactive td{opacity:.7}.rt-year-badge{display:inline-block;min-width:30px;padding:4px 8px;border-radius:8px;background:#0b4c69;color:#8ff5f4;text-align:center;font-weight:700}.rt-status-pill{display:inline-block;padding:4px 8px;border-radius:20px;font-size:12px;font-weight:700}.status-confirmed{background:#064c3e;color:#a7f3d0}.status-pending{background:#624013;color:#ffe0a3}.status-rejected,.status-cancelled{background:#5b1f2b;color:#ffc1c8}.rt-danger-button{background:#762536!important;border-color:#d64f67!important}@media(max-width:680px){.rt-assignment-stats{grid-template-columns:repeat(2,minmax(0,1fr))}.rt-assignment-stats strong{font-size:21px}} </style>""",rs=rs,ts=ts,ps=ps,rows=rows,blocked_periods_by_resident=blocked_periods_by_resident)
@app.post("/assignments/<int:aid>/cancel")
def cancel_assignment(aid):
 if super_guard(): return super_guard()
 c=db()
 row=c.execute("SELECT status FROM assignments WHERE id=?",(aid,)).fetchone()
 if not row:
  c.close(); return ("Assignment not found",404)
 if row["status"]=="cancelled":
  c.close(); flash("Assignment already cancelled."); return redirect("/assign")
 c.execute("UPDATE assignments SET status='cancelled' WHERE id=?",(aid,))
 audit_event(c,"assignment.cancelled",f"assignment_id={aid}; previous_status={row['status']}")
 c.commit(); c.close(); flash("Assignment cancelled; historical record retained.")
 return redirect("/assign")
@app.route("/history",methods=["GET","POST"])
def history():
 if guard():return redirect("/login")
 c=db()
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  try:
   rid=int(request.form["resident"]); tid=int(request.form["terrain"])
   year_label=request.form.get("year_label","").strip()
   period=request.form.get("period","").strip()
   date_done=request.form.get("date_done","").strip()
   if not re.fullmatch(r"\d{4}-\d{4}",year_label) or int(year_label.split("-")[1])!=int(year_label.split("-")[0])+1:
    raise ValueError("Academic year must be YYYY-YYYY with consecutive years.")
   valid_periods={"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
   if period and period not in valid_periods:
    raise ValueError("Choose one of the four official rotation periods.")
   if date_done:
    parsed_date=date.fromisoformat(date_done)
    if parsed_date>date.today():
     raise ValueError("A historical rotation date cannot be in the future.")
   if not c.execute("SELECT 1 FROM residents WHERE id=?",(rid,)).fetchone():
    raise ValueError("Resident does not exist.")
   if not c.execute("SELECT 1 FROM terrains WHERE id=?",(tid,)).fetchone():
    raise ValueError("Terrain does not exist.")
   duplicate=c.execute("SELECT 1 FROM history WHERE resident_id=? AND terrain_id=? AND year_label=? AND period=? AND COALESCE(date_done,'')=?",(rid,tid,year_label,period,date_done)).fetchone()
   if duplicate: raise ValueError("This historical rotation already exists.")
   c.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",(rid,tid,year_label,period,date_done))
   audit_event(c,"history.created",f"resident_id={rid}; terrain_id={tid}; year={year_label}; period={period}")
   c.commit(); flash(tr("done"))
  except (TypeError,ValueError,sqlite3.IntegrityError) as exc:
   c.rollback(); flash(str(exc) or "Invalid historical rotation.")
 rs=c.execute("SELECT * FROM residents ORDER BY year DESC,surname").fetchall(); ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); rows=c.execute("SELECT h.*,r.surname,r.firstname,t.name terrain FROM history h JOIN residents r ON r.id=h.resident_id JOIN terrains t ON t.id=h.terrain_id ORDER BY h.id DESC").fetchall(); c.close()
 return page("""<section><h3>{{t('history_add')}}</h3><form method="post"><label>{{t('resident')}}</label><select name="resident">{% for r in rs %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}} — {{r.year}}</option>{% endfor %}</select><label>{{t('terrain')}}</label><select name="terrain">{% for x in ts %}<option value="{{x.id}}">{{x.name}}</option>{% endfor %}</select><label>{{t('academic_year')}}</label><input name="year_label" required placeholder="2025-2026"><label>{{t('period')}}</label><select name="period"><option value="">—</option>{% for p in ["01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"] %}<option value="{{p}}">{{p}}</option>{% endfor %}</select><label>{{t('date')}}</label><input type="date" name="date_done"><button>{{t('save')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('terrain')}}</th><th>{{t('academic_year')}}</th><th>{{t('period')}}</th><th>{{t('date')}}</th></tr>{% for h in rows %}<tr><td>{{h.surname}} {{h.firstname}}</td><td>{{h.terrain}}</td><td>{{h.year_label}}</td><td>{{h.period}}</td><td>{{h.date_done}}</td></tr>{% endfor %}</table></div></section>""",rs=rs,ts=ts,rows=rows)
@app.get("/summary")
def summary():
 if guard():return redirect("/login")
 c=db(); y=setting(c,"year"); rs=c.execute("SELECT * FROM residents ORDER BY year DESC,surname").fetchall(); rows=[]
 for r in rs:
  a=c.execute("SELECT p.name period,t.name terrain,a.status FROM assignments a JOIN periods p ON p.id=a.period_id JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=? AND p.year_label=? ORDER BY p.id",(r["id"],y)).fetchall()
  h=c.execute("SELECT t.name terrain,h.year_label,h.period,h.date_done FROM history h JOIN terrains t ON t.id=h.terrain_id WHERE h.resident_id=?",(r["id"],)).fetchall(); rows.append((r,a,h))
 c.close()
 return page("""<section><h3>{{t('summary')}} — {{year}}</h3><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('assign')}}</th><th>{{t('history')}}</th></tr>{% for r,a,h in rows %}<tr><td>{{r.surname}} {{r.firstname}}</td><td>{{r.year}}</td><td>{% for x in a %}<div>{{x.period}} — {{x.terrain}} ({{x.status}})</div>{% else %}—{% endfor %}</td><td>{% for x in h %}<div>{{x.terrain}} — {{x.year_label}} / {{x.period}} / {{x.date_done}}</div>{% else %}—{% endfor %}</td></tr>{% endfor %}</table></div></section>""",rows=rows,year=y)
@app.route("/decisions",methods=["GET","POST"])
def decisions():
 if super_guard():return super_guard()
 c=db()
 if request.method=="POST":
  aid=request.form.get("aid",type=int); outcome=request.form.get("outcome",""); reason=request.form.get("reason","").strip()
  if outcome not in ("approve","reject"): c.close(); return ("Invalid decision",400)
  a=c.execute("SELECT a.*,r.surname,r.firstname,r.status resident_status,t.name terrain,t.active terrain_active,p.name period,p.year_label FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.id=? AND a.status='pending'",(aid,)).fetchone()
  if a and reason:
   if outcome=="approve":
    if a["resident_status"]!="active":
     flash("Cannot approve: the resident is no longer active.")
    elif not a["terrain_active"]:
     flash("Cannot approve: the terrain has been deactivated.")
    else:
     start_day,end_day=rotation_period_bounds(a["year_label"],a["period"])
     external=c.execute("SELECT hospital,country,start_date,end_date FROM external_rotations WHERE resident_id=? AND status<>'cancelled' AND start_date<=? AND end_date>=? LIMIT 1",(a["resident_id"],end_day.isoformat(),start_day.isoformat())).fetchone()
     if external:
      flash("Cannot approve: an external rotation overlaps this period.")
     else:
      cap=c.execute("SELECT places FROM capacities WHERE year_label=? AND terrain_id=? AND period_id=?",(a["year_label"],a["terrain_id"],a["period_id"])).fetchone()
      used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status='confirmed'",(a["terrain_id"],a["period_id"])).fetchone()[0]
      dup=c.execute("SELECT id FROM assignments WHERE resident_id=? AND period_id=? AND id<>? AND status IN ('confirmed','pending')",(a["resident_id"],a["period_id"],aid)).fetchone()
      if not cap or used>=cap["places"] or dup: flash(tr("full") if not dup else tr("duplicate"))
      else:
       c.execute("UPDATE assignments SET status='confirmed' WHERE id=?",(aid,)); c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(a["surname"]+" "+a["firstname"]+" — "+a["terrain"],tr("approve"),reason,session["user"]))
       audit_event(c,"assignment.override_approved",f"assignment_id={aid}; reason={reason}")
       c.commit(); flash(tr("done"))
   else:
    c.execute("UPDATE assignments SET status='rejected' WHERE id=?",(aid,)); c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(a["surname"]+" "+a["firstname"]+" — "+a["terrain"],tr("reject"),reason,session["user"]))
    audit_event(c,"assignment.override_rejected",f"assignment_id={aid}; reason={reason}")
    c.commit(); flash(tr("done"))
 pending=c.execute("SELECT a.id,r.surname,r.firstname,r.year,t.name terrain,p.name period FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.status='pending'").fetchall(); ds=c.execute("SELECT * FROM decisions ORDER BY id DESC").fetchall(); c.close()
 return page("""<section><h3>{{t('decisions')}}</h3>{% for a in pending %}<p><b>{{a.surname}} {{a.firstname}}</b> — {{a.terrain}} / {{a.period}}</p><form method="post"><input type="hidden" name="aid" value="{{a.id}}"><input name="reason" placeholder="{{t('reason')}}" required><button name="outcome" value="approve">{{t('approve')}}</button> <button name="outcome" value="reject">{{t('reject')}}</button></form>{% else %}<p>{{t('none')}}</p>{% endfor %}</section><section><div class="wrap"><table><tr><th>Date</th><th>{{t('assign')}}</th><th>{{t('status')}}</th><th>{{t('reason')}}</th></tr>{% for d in ds %}<tr><td>{{d.created_at}}</td><td>{{d.details}}</td><td>{{d.decision}}</td><td>{{d.reason}}</td></tr>{% endfor %}</table></div></section>""",pending=pending,ds=ds)
@app.route("/settings",methods=["GET","POST"])
def settings():
 if super_guard():return super_guard()
 c=db()
 if request.method=="POST":
  lang=request.form.get("lang","ar")
  if lang in LANGS:session["lang"]=lang
  y=request.form.get("year","").strip()
  if y and (not re.fullmatch(r"\d{4}-\d{4}",y) or int(y.split("-")[1])!=int(y.split("-")[0])+1):
   flash("Academic year must be formatted as YYYY-YYYY with consecutive years.")
   y=setting(c,"year")
  if y:
   c.execute("INSERT INTO settings(key,value) VALUES('year',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(y,))
   for n in ["01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"]:
    c.execute("INSERT OR IGNORE INTO periods(year_label,name) VALUES(?,?)",(y,n))
  d=request.form.get("director","").strip()
  if d:c.execute("INSERT INTO settings(key,value) VALUES('director',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(d,))
  pw=request.form.get("new_password","")
  if pw and len(pw)<12:
   flash("New password must be at least 12 characters.")
   pw=""
  if pw:
   secure_hash=hash_password(pw)
   c.execute("INSERT INTO settings(key,value) VALUES('password',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(secure_hash,))
   c.execute("UPDATE users SET password_hash=? WHERE username=?",(secure_hash,session.get("user","admin")))
   audit_event(c,"security.password_changed","Administrator password changed")
  audit_event(c,"settings.updated",f"year={y}; director_updated={bool(d)}")
  c.commit(); flash(tr("done"))
 vals={r["key"]:r["value"] for r in c.execute("SELECT key,value FROM settings")}; c.close()
 return page("""<section><h3>{{t('settings')}}</h3><form method="post"><label>{{t('language')}}</label><select name="lang"><option value="ar" {% if lang=='ar' %}selected{% endif %}>العربية</option><option value="fr" {% if lang=='fr' %}selected{% endif %}>Français</option><option value="en" {% if lang=='en' %}selected{% endif %}>English</option></select><label>{{t('year')}}</label><input name="year" value="{{vals.get('year','')}}" required><label>{{t('director')}}</label><input name="director" value="{{vals.get('director','')}}"><label>{{t('new_password')}}</label><input type="password" name="new_password" minlength="12" autocomplete="new-password"><button>{{t('change')}}</button></form><p class="muted">{{t('note')}}</p></section>""",vals=vals,lang=session.get("lang","ar"))
@app.route("/users",methods=["GET","POST"])
def manage_users():
 if super_guard(): return super_guard()
 c=db()
 if request.method=="POST":
  action=request.form.get("action","create")
  if action in ("deactivate","activate"):
   uid=request.form.get("user_id",type=int)
   target=c.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
   if not target: flash("User not found.")
   elif action=="deactivate" and target["username"].lower() in ("benhocine","drdris"): flash("Official accounts cannot be deactivated.")
   elif action=="deactivate" and target["username"]==session.get("user"): flash("You cannot deactivate your own account while signed in.")
   elif action=="deactivate" and target["role"]=="super_admin" and target["active"] and c.execute("SELECT COUNT(*) FROM users WHERE role='super_admin' AND active=1").fetchone()[0]<=1: flash("At least one active Super Admin must remain.")
   else:
    active=0 if action=="deactivate" else 1
    c.execute("UPDATE users SET active=? WHERE id=?",(active,uid))
    audit_event(c,"user.deactivated" if not active else "user.activated",f"username={target['username']}")
    c.commit(); flash(tr("done"))
  else:
   username=request.form.get("username","").strip().lower()
   password=request.form.get("password","")
   role=request.form.get("role","user")
   if username in ("benhocine","drdris") and role!="super_admin":
    flash("Official accounts must be Super Admins.")
   elif len(username)<3 or len(password)<12 or role not in ("super_admin","user"):
    flash("Username must have 3+ characters and password 12+ characters.")
   else:
    try:
     c.execute("INSERT INTO users(username,password_hash,role) VALUES(?,?,?)",(username,hash_password(password),role))
     audit_event(c,"user.created",f"username={username}; role={role}")
     c.commit(); flash(tr("done"))
    except sqlite3.IntegrityError:
     flash("Username already exists.")
 rows=c.execute("SELECT id,username,role,active,created_at FROM users ORDER BY role DESC,username").fetchall(); c.close()
 return page("""<section><h3>{{t('users')}}</h3><p>{{t('users_help')}}</p><form method="post"><input type="hidden" name="action" value="create"><label>{{t('username')}}</label><input name="username" minlength="3" required><label>{{t('password')}} (12+)</label><input name="password" type="password" minlength="12" required><label>{{t('role')}}</label><select name="role"><option value="user">{{t('standard_user')}}</option><option value="super_admin">{{t('super_admin')}}</option></select><button>{{t('add')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('username')}}</th><th>{{t('role')}}</th><th>{{t('active_col')}}</th><th>{{t('created')}}</th><th>Admin action</th></tr>{% for u in rows %}<tr><td>{{u.username}}</td><td>{{u.role}}</td><td>{{u.active}}</td><td>{{u.created_at}}</td><td>{% if u.username!=session.get('user') or u.active==0 %}<form method="post"><input type="hidden" name="user_id" value="{{u.id}}"><button name="action" value="{{'deactivate' if u.active else 'activate'}}">{{'Deactivate' if u.active else 'Activate'}}</button></form>{% else %}Current account{% endif %}</td></tr>{% endfor %}</table></div></section>""",rows=rows)
@app.get("/audit")
def audit_log():
 if super_guard(): return super_guard()
 c=db(); rows=c.execute("SELECT * FROM audit ORDER BY id DESC LIMIT 300").fetchall(); c.close()
 return page("""<section><h3>{{t('audit')}}</h3><div class="wrap"><table><tr><th>{{t('date')}}</th><th>{{t('actor')}}</th><th>{{t('action')}}</th><th>{{t('details')}}</th></tr>{% for r in rows %}<tr><td>{{r.created_at}}</td><td>{{r.actor}}</td><td>{{r.action}}</td><td>{{r.details}}</td></tr>{% else %}<tr><td colspan="4">{{t('none')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows)
def normalize_header(value):
 value=str(value or "").strip().lower()
 value=re.sub(r"[\s_\-]+"," ",value)
 aliases={
  "surname":"surname","nom":"surname","family name":"surname","last name":"surname","اللقب":"surname",
  "firstname":"firstname","first name":"firstname","prénom":"firstname","prenom":"firstname","given name":"firstname","الاسم":"firstname",
  "year":"year","study year":"year","année":"year","annee":"year","niveau":"year","السنة":"year",
  "phone":"phone","telephone":"phone","téléphone":"phone","mobile":"phone","الهاتف":"phone",
  "birth place":"birth_place","lieu de naissance":"birth_place","مكان الميلاد":"birth_place",
  "birth date":"birth_date","date de naissance":"birth_date","تاريخ الميلاد":"birth_date",
  "resident":"resident","resident name":"resident","resident full name":"resident","nom du résident":"resident","nom resident":"resident","المقيم":"resident",
  "terrain":"terrain","service":"terrain","stage":"terrain","rotation":"terrain","service de stage":"terrain","المصلحة":"terrain","التربص":"terrain",
  "academic year":"year_label","year label":"year_label","academic year label":"year_label","année universitaire":"year_label","annee universitaire":"year_label","السنة الجامعية":"year_label",
  "period":"period","rotation period":"period","période":"period","periode":"period","الفترة":"period",
  "date done":"date_done","completion date":"date_done","date completed":"date_done","date de réalisation":"date_done","date de fin":"date_done","تاريخ الإنجاز":"date_done"
 }
 return aliases.get(value,value.replace(" ","_"))
def parse_import_file(path, extension, import_type='residents'):
 rows=[]
 if extension==".csv":
  with open(path,"r",encoding="utf-8-sig",newline="") as handle:
   sample=handle.read(4096); handle.seek(0)
   try: dialect=csv.Sniffer().sniff(sample,delimiters=",;\t")
   except csv.Error: dialect=csv.excel
   rows=list(csv.reader(handle,dialect))
 elif extension in (".xlsx",".xls"):
  if extension==".xlsx":
   from openpyxl import load_workbook
   book=load_workbook(path,read_only=True,data_only=True)
   rows=[list(row) for row in book.active.iter_rows(values_only=True)]
   book.close()
  else:
   import xlrd
   book=xlrd.open_workbook(path)
   sheet=book.sheet_by_index(0)
   rows=[[sheet.cell_value(r,c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
 elif extension==".docx":
  from docx import Document
  doc=Document(path)
  if doc.tables:
   table=doc.tables[0]
   rows=[[cell.text.strip() for cell in row.cells] for row in table.rows]
  else:
   rows=[re.split(r"[,;\t]+",p.text.strip()) for p in doc.paragraphs if p.text.strip()]
 elif extension==".pdf":
  from pypdf import PdfReader
  text_content="\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
  if not text_content.strip(): raise ValueError("The PDF has no extractable text (it may be scanned). Please provide a text-based PDF or spreadsheet.")
  lines=[line.strip() for line in text_content.splitlines() if line.strip()]
  delimiter=";" if sum(x.count(";") for x in lines)>sum(x.count(",") for x in lines) else ("\t" if any("\t" in x for x in lines) else ",")
  rows=[next(csv.reader([line],delimiter=delimiter)) for line in lines]
 elif extension==".doc":
  try:
   result=subprocess.run(["antiword",path],capture_output=True,text=True,check=True,timeout=20)
   document_text=result.stdout
  except FileNotFoundError:
   try:
    with tempfile.TemporaryDirectory(prefix="rotation_doc_") as converted_dir:
     subprocess.run(["libreoffice","--headless","--convert-to","txt:Text","--outdir",converted_dir,path],capture_output=True,text=True,check=True,timeout=30)
     converted_path=os.path.join(converted_dir,os.path.splitext(os.path.basename(path))[0]+".txt")
     with open(converted_path,"r",encoding="utf-8",errors="replace") as handle: document_text=handle.read()
   except (FileNotFoundError,subprocess.CalledProcessError,subprocess.TimeoutExpired) as exc:
    raise ValueError("Legacy .doc requires antiword or LibreOffice; convert it to .docx if neither is installed.") from exc
  lines=[line.strip() for line in document_text.splitlines() if line.strip()]
  rows=[re.split(r"\s*[,;\t]\s*",line) for line in lines]
 if not rows: raise ValueError("No table rows could be read from this file.")
 header=[normalize_header(x) for x in rows[0]]
 required={"resident","terrain","year_label"} if import_type=="history" else {"surname","firstname","year"}
 if not required.issubset(set(header)):
  if import_type=="history":
   raise ValueError("History import requires resident name, terrain/service, and academic year (YYYY-YYYY) columns; period and completion date are optional.")
  raise ValueError("Missing required columns. Include surname/nom/اللقب, firstname/prénom/الاسم, and year/année/السنة in the first row.")
 parsed=[]
 for line_number,values in enumerate(rows[1:],start=2):
  if not any(str(x or "").strip() for x in values): continue
  item={}
  for i,key in enumerate(header):
   if key and i<len(values): item[key]=str(values[i] if values[i] is not None else "").strip()
  item["row_number"]=line_number
  parsed.append(item)
 if not parsed: raise ValueError("The file contains headers but no resident rows.")
 return parsed
@app.route("/import",methods=["GET","POST"])
def import_residents():
 if super_guard(): return super_guard()
 c=db()
 c.execute("DELETE FROM import_previews WHERE created_at < datetime('now','-1 day')"); c.commit()
 if request.method=="POST":
  action=request.form.get("action","upload")
  if action in ("confirm","cancel"):
   preview_id=request.form.get("preview_id",type=int)
   preview=c.execute("SELECT * FROM import_previews WHERE id=? AND owner=?",(preview_id,session["user"])).fetchone()
   if not preview:
    c.close(); flash("Import preview expired or not found."); return redirect("/import")
   if action=="cancel":
    c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close(); flash("Import cancelled."); return redirect("/import")
   rows=json.loads(preview["rows_json"]); errors=json.loads(preview["errors_json"])
   if errors:
    c.close(); flash("Fix the invalid rows in the source file before importing."); return redirect("/import")
   # Re-check against live data at confirmation: another administrator may have
   # created one of these residents after the preview was generated.
   seen=set(); confirmation_errors=[]
   for row in rows:
    surname=str(row.get("surname","")).strip()
    firstname=str(row.get("firstname","")).strip()
    try:
     year=int(row.get("year"))
     if year not in (1,2,3,4): raise ValueError()
    except (ValueError,TypeError):
     confirmation_errors.append({"row":row.get("row_number"),"message":"Year must be 1, 2, 3 or 4."})
     continue
    key=(surname.casefold(),firstname.casefold(),year)
    if not surname or not firstname:
     confirmation_errors.append({"row":row.get("row_number"),"message":"Surname and first name are required."})
    elif key in seen:
     confirmation_errors.append({"row":row.get("row_number"),"message":"Duplicate resident row in this file."})
    elif c.execute("SELECT 1 FROM residents WHERE lower(surname)=lower(?) AND lower(firstname)=lower(?) AND year=?",(surname,firstname,year)).fetchone():
     confirmation_errors.append({"row":row.get("row_number"),"message":"Resident with same name and year already exists; upload a fresh preview."})
    seen.add(key)
   if confirmation_errors:
    c.execute("UPDATE import_previews SET errors_json=? WHERE id=?",(json.dumps(confirmation_errors,ensure_ascii=False),preview_id))
    c.commit(); c.close()
    flash("The resident list changed after preview. Nothing was imported; upload the file again to review the current data.")
    return redirect("/import")
   inserted=0
   try:
    for row in rows:
     c.execute("INSERT INTO residents(surname,firstname,year,phone,birth_place,birth_date,status) VALUES(?,?,?,?,?,?, 'active')",
       (row["surname"],row["firstname"],int(row["year"]),row.get("phone",""),row.get("birth_place",""),row.get("birth_date","")))
     inserted+=1
   except sqlite3.IntegrityError:
    c.rollback(); c.close()
    flash("Import cancelled because the resident data changed during confirmation. Nothing was imported; upload a fresh preview.")
    return redirect("/import")
   audit_event(c,"residents.imported",f"filename={preview['filename']}; imported={inserted}; history_not_imported=true")
   c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close()
   flash(f"Imported {inserted} residents. No past rotations were invented or imported.")
   return redirect("/residents")
  upload=request.files.get("file")
  if not upload or not upload.filename:
   c.close(); flash("Choose a file first."); return redirect("/import")
  filename=os.path.basename(upload.filename); extension=os.path.splitext(filename)[1].lower()
  allowed={".csv",".xlsx",".xls",".docx",".doc",".pdf"}
  if extension not in allowed:
   c.close(); flash("Supported formats: CSV, XLS/XLSX, DOC/DOCX and text-based PDF."); return redirect("/import")
  staged=None
  try:
   with tempfile.NamedTemporaryFile(prefix="rotation_import_",suffix=extension,dir=os.path.dirname(os.path.abspath(DB)),delete=False) as temp:
    staged=temp.name; upload.save(staged)
   rows=parse_import_file(staged,extension)
   errors=[]; seen=set()
   for row in rows:
    row_errors=[]
    surname=row.get("surname","").strip(); firstname=row.get("firstname","").strip()
    if not surname or not firstname: row_errors.append("Surname and first name are required.")
    try:
     numeric_year=float(row.get("year",""))
     if not numeric_year.is_integer(): raise ValueError()
     year=int(numeric_year)
     if year not in (1,2,3,4): raise ValueError()
     row["year"]=year
    except (ValueError,TypeError): row_errors.append("Year must be 1, 2, 3 or 4.")
    row["surname"]=surname; row["firstname"]=firstname
    key=(surname.casefold(),firstname.casefold(),row.get("year"))
    if key in seen: row_errors.append("Duplicate resident row in this file.")
    seen.add(key)
    exists=c.execute("SELECT 1 FROM residents WHERE lower(surname)=lower(?) AND lower(firstname)=lower(?) AND year=?",(surname,firstname,row.get("year",0))).fetchone()
    if exists: row_errors.append("Resident with same name and year already exists.")
    if row_errors: errors.append({"row":row.get("row_number"),"message":" ".join(row_errors)})
   rows.sort(key=lambda row:(row.get("year",99),row.get("surname","").casefold(),row.get("firstname","").casefold()))
   c.execute("DELETE FROM import_previews WHERE created_at < datetime('now','-1 day')")
   c.execute("INSERT INTO import_previews(owner,rows_json,errors_json,filename) VALUES(?,?,?,?)",(session["user"],json.dumps(rows,ensure_ascii=False),json.dumps(errors,ensure_ascii=False),filename))
   preview_id=c.execute("SELECT last_insert_rowid()").fetchone()[0]; c.commit()
   preview=c.execute("SELECT * FROM import_previews WHERE id=?",(preview_id,)).fetchone(); c.close()
   return page("""<section><h3>{{t('import')}} — {{t('preview')}}</h3><p>File: {{filename}} — {{rows|length}} rows. Past stages/history will not be fabricated.</p>{% if errors %}<p class="bad">Validation errors: {{errors|length}}. Correct the source file and upload again; no rows can be imported while errors remain.</p>{% endif %}<div class="wrap"><table><tr><th>Row</th><th>Surname</th><th>First name</th><th>Year</th><th>Phone</th><th>Validation</th></tr>{% for row in rows %}<tr><td>{{row.row_number}}</td><td>{{row.surname}}</td><td>{{row.firstname}}</td><td>{{row.year}}</td><td>{{row.phone or ''}}</td><td>{% for e in errors if e.row==row.row_number %}<span class="bad">{{e.message}}</span>{% else %}<span class="good">OK</span>{% endfor %}</td></tr>{% endfor %}</table></div><form method="post"><input type="hidden" name="preview_id" value="{{preview_id}}"><button name="action" value="confirm" {% if errors %}disabled{% endif %}>{{t('confirm_import')}}</button> <button name="action" value="cancel">{{t('cancel')}}</button></form></section>""",filename=filename,rows=rows,errors=errors,preview_id=preview_id)
  except Exception as exc:
   c.close(); flash("Could not read/validate the file: "+str(exc)[:250]); return redirect("/import")
  finally:
   if staged and os.path.exists(staged): os.remove(staged)
 c.close()
 return page("""<section><h3>{{t('import')}}</h3><p>{{t('import_help')}} Supported: CSV, XLS/XLSX, DOC/DOCX and text-based PDF.</p><form method="post" enctype="multipart/form-data"><label>{{t('file')}}</label><input type="file" name="file" accept=".csv,.xls,.xlsx,.doc,.docx,.pdf" required><button>{{t('upload_preview')}}</button></form><p><a href="/import-history">Import documented past stages / Importer l’historique documenté / استيراد التربصات السابقة الموثقة</a></p></section>""")

@app.route("/import-history",methods=["GET","POST"])
def import_history():
 if super_guard(): return super_guard()
 c=db()
 c.execute("DELETE FROM import_previews WHERE created_at < datetime('now','-1 day')"); c.commit()
 if request.method=="POST":
  action=request.form.get("action","upload")
  if action in ("confirm","cancel"):
   preview_id=request.form.get("preview_id",type=int)
   preview=c.execute("SELECT * FROM import_previews WHERE id=? AND owner=? AND import_type='history'",(preview_id,session["user"])).fetchone()
   if not preview:
    c.close(); flash("History import preview expired or not found."); return redirect("/import-history")
   if action=="cancel":
    c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close(); flash("History import cancelled."); return redirect("/import-history")
   rows=json.loads(preview["rows_json"]); errors=json.loads(preview["errors_json"])
   if errors:
    c.close(); flash("Fix all history validation errors before importing."); return redirect("/import-history")
   seen=set(); validation=[]
   for row in rows:
    resident_name=str(row.get("resident","")).strip()
    terrain_name=str(row.get("terrain","")).strip()
    year_label=str(row.get("year_label","")).strip()
    matches=c.execute("SELECT id,surname,firstname FROM residents WHERE lower(trim(surname||' '||firstname))=lower(?) OR lower(trim(firstname||' '||surname))=lower(?) OR lower(trim(surname||', '||firstname))=lower(?)",(resident_name,resident_name,resident_name)).fetchall()
    terrains=c.execute("SELECT id,name FROM terrains WHERE lower(trim(name))=lower(?)",(terrain_name,)).fetchall()
    if len(matches)!=1: validation.append({"row":row.get("row_number"),"message":"Resident no longer matches exactly one existing resident."}); continue
    if len(terrains)!=1: validation.append({"row":row.get("row_number"),"message":"Terrain no longer matches exactly one existing terrain."}); continue
    try:
     if not re.fullmatch(r"\d{4}-\d{4}",year_label) or int(year_label.split("-")[1])!=int(year_label.split("-")[0])+1: raise ValueError("Academic year must use YYYY-YYYY with consecutive years.")
     period=str(row.get("period","")).strip()
     valid_periods={"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
     if period and period not in valid_periods: raise ValueError("Choose one of the four official periods or leave period blank.")
     date_done=str(row.get("date_done","")).strip()
     if date_done:
      parsed=date.fromisoformat(date_done)
      if parsed>date.today(): raise ValueError("A historical rotation date cannot be in the future.")
    except ValueError as exc:
     validation.append({"row":row.get("row_number"),"message":str(exc)}); continue
    rid=matches[0]["id"]; tid=terrains[0]["id"]
    key=(rid,tid,year_label,period,date_done)
    if key in seen: validation.append({"row":row.get("row_number"),"message":"Duplicate history row in this file."})
    elif c.execute("SELECT 1 FROM history WHERE resident_id=? AND terrain_id=? AND year_label=? AND COALESCE(period,'')=? AND COALESCE(date_done,'')=?",(rid,tid,year_label,period,date_done)).fetchone():
     validation.append({"row":row.get("row_number"),"message":"This exact history record already exists; upload a fresh preview."})
    seen.add(key)
   if validation:
    c.execute("UPDATE import_previews SET errors_json=? WHERE id=?",(json.dumps(validation,ensure_ascii=False),preview_id)); c.commit(); c.close()
    flash("History changed after preview. Nothing was imported; upload the file again.")
    return redirect("/import-history")
   inserted=0
   try:
    for row in rows:
     resident_name=str(row["resident"]).strip(); terrain_name=str(row["terrain"]).strip()
     match=c.execute("SELECT id FROM residents WHERE lower(trim(surname||' '||firstname))=lower(?) OR lower(trim(firstname||' '||surname))=lower(?) OR lower(trim(surname||', '||firstname))=lower(?)",(resident_name,resident_name,resident_name)).fetchone()
     terrain=c.execute("SELECT id FROM terrains WHERE lower(trim(name))=lower(?)",(terrain_name,)).fetchone()
     c.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",(match["id"],terrain["id"],row["year_label"],row.get("period",""),row.get("date_done","")))
     inserted+=1
   except sqlite3.IntegrityError:
    c.rollback(); c.close(); flash("Import cancelled because history changed during confirmation. Nothing was imported; upload a fresh preview."); return redirect("/import-history")
   audit_event(c,"history.imported",f"filename={preview['filename']}; imported={inserted}; documented_rows_only=true")
   c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close()
   flash(f"Imported {inserted} documented past-stage records. No history was invented.")
   return redirect("/history")
  upload=request.files.get("file")
  if not upload or not upload.filename:
   c.close(); flash("Choose a file first."); return redirect("/import-history")
  filename=os.path.basename(upload.filename); extension=os.path.splitext(filename)[1].lower()
  if extension not in {".csv",".xlsx",".xls",".docx",".doc",".pdf"}:
   c.close(); flash("Supported formats: CSV, XLS/XLSX, DOC/DOCX and text-based PDF."); return redirect("/import-history")
  staged=None
  try:
   with tempfile.NamedTemporaryFile(prefix="rotation_history_import_",suffix=extension,dir=os.path.dirname(os.path.abspath(DB)),delete=False) as temp:
    staged=temp.name; upload.save(staged)
   rows=parse_import_file(staged,extension,"history")
   errors=[]; seen=set(); valid_periods={"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
   residents_pool=c.execute("SELECT id,surname,firstname FROM residents").fetchall()
   terrain_pool=c.execute("SELECT id,name FROM terrains").fetchall()
   for row in rows:
    row_errors=[]
    resident_name=str(row.get("resident","")).strip(); terrain_name=str(row.get("terrain","")).strip(); year_label=str(row.get("year_label","")).strip()
    row["resident"]=resident_name; row["terrain"]=terrain_name; row["year_label"]=year_label
    if not resident_name: row_errors.append("Resident name is required.")
    if not terrain_name: row_errors.append("Terrain name is required.")
    if not re.fullmatch(r"\d{4}-\d{4}",year_label) or int(year_label.split("-")[1])!=int(year_label.split("-")[0])+1:
     row_errors.append("Academic year must use YYYY-YYYY with consecutive years.")
    resident_matches=[r for r in residents_pool if resident_name.casefold() in {(r["surname"]+" "+r["firstname"]).casefold().strip(),(r["firstname"]+" "+r["surname"]).casefold().strip(),(r["surname"]+", "+r["firstname"]).casefold().strip()}]
    terrain_matches=[t for t in terrain_pool if t["name"].casefold().strip()==terrain_name.casefold().strip()]
    if len(resident_matches)!=1: row_errors.append("Resident must match exactly one existing resident; use 'Surname Firstname'.")
    else: row["resident_id"]=resident_matches[0]["id"]
    if len(terrain_matches)!=1: row_errors.append("Terrain must match exactly one existing terrain.")
    else: row["terrain_id"]=terrain_matches[0]["id"]
    period=str(row.get("period","")).strip(); date_done=str(row.get("date_done","")).strip()
    if period and period not in valid_periods: row_errors.append("Choose one of the four official periods or leave period blank.")
    if date_done:
     try:
      parsed=date.fromisoformat(date_done)
      if parsed>date.today(): row_errors.append("A historical rotation date cannot be in the future.")
     except ValueError: row_errors.append("Date must be a real calendar date in YYYY-MM-DD format.")
    row["period"]=period; row["date_done"]=date_done
    key=(row.get("resident_id"),row.get("terrain_id"),year_label,period,date_done)
    if key in seen: row_errors.append("Duplicate history row in this file.")
    seen.add(key)
    if row.get("resident_id") and row.get("terrain_id") and c.execute("SELECT 1 FROM history WHERE resident_id=? AND terrain_id=? AND year_label=? AND COALESCE(period,'')=? AND COALESCE(date_done,'')=?",(row["resident_id"],row["terrain_id"],year_label,period,date_done)).fetchone():
     row_errors.append("This exact history record already exists.")
    if row_errors: errors.append({"row":row.get("row_number"),"message":" ".join(row_errors)})
   rows.sort(key=lambda row:(row.get("year_label",""),row.get("resident","").casefold(),row.get("terrain","").casefold()))
   c.execute("INSERT INTO import_previews(owner,rows_json,errors_json,filename,import_type) VALUES(?,?,?,?,?)",(session["user"],json.dumps(rows,ensure_ascii=False),json.dumps(errors,ensure_ascii=False),filename,"history"))
   preview_id=c.execute("SELECT last_insert_rowid()").fetchone()[0]; c.commit(); c.close()
   return page("""<section><h3>Import documented stage history / Importer l’historique / استيراد التربصات السابقة</h3><p>File: {{filename}} — {{rows|length}} rows. Only explicitly documented stages are imported; no stages are inferred.</p>{% if errors %}<p class="bad">Validation errors: {{errors|length}}. Nothing can be imported until the source is corrected.</p>{% endif %}<div class="wrap"><table><tr><th>Row</th><th>Resident</th><th>Terrain</th><th>Academic year</th><th>Period</th><th>Completion date</th><th>Validation</th></tr>{% for row in rows %}<tr><td>{{row.row_number}}</td><td>{{row.resident}}</td><td>{{row.terrain}}</td><td>{{row.year_label}}</td><td>{{row.period}}</td><td>{{row.date_done}}</td><td>{% for e in errors if e.row==row.row_number %}<span class="bad">{{e.message}}</span>{% else %}<span class="good">OK</span>{% endfor %}</td></tr>{% endfor %}</table></div><form method="post"><input type="hidden" name="preview_id" value="{{preview_id}}"><button name="action" value="confirm" {% if errors %}disabled{% endif %}>Confirm import</button> <button name="action" value="cancel">Cancel</button></form></section>""",filename=filename,rows=rows,errors=errors,preview_id=preview_id)
  except Exception as exc:
   c.close(); flash("Could not read/validate history file: "+str(exc)[:220]); return redirect("/import-history")
  finally:
   if staged and os.path.exists(staged): os.remove(staged)
 c.close()
 return page("""<section><h3>Import documented stage history / Importer l’historique / استيراد التربصات السابقة</h3><p>Upload CSV, XLS/XLSX, DOC/DOCX or text-based PDF with resident, terrain/service, academic year columns; period and completion date are optional. Resident and terrain must match existing records exactly.</p><form method="post" enctype="multipart/form-data"><label>History file</label><input type="file" name="file" accept=".csv,.xls,.xlsx,.doc,.docx,.pdf" required><button>Upload and preview</button></form><p><a href="/import">Back to resident import</a></p></section>""")

@app.get("/photos/<filename>")
def resident_photo(filename):
 if guard(): return redirect("/login")
 safe=secure_filename(filename)
 if safe!=filename or not safe.startswith("resident_") or os.path.splitext(safe)[1].lower() not in (".png",".jpg",".jpeg",".webp"): return ("Not found",404)
 directory=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos")
 return send_from_directory(directory,safe)
@app.route("/profile/<int:rid>",methods=["GET","POST"])
def resident_profile(rid):
 if guard(): return redirect("/login")
 c=db(); r=c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone()
 if not r: c.close(); return ("Resident not found",404)
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  surname=request.form.get("surname",r["surname"]).strip()
  firstname=request.form.get("firstname",r["firstname"]).strip()
  birth_place=request.form.get("birth_place",r["birth_place"] or "").strip()
  birth_date=request.form.get("birth_date",r["birth_date"] or "").strip()
  phone=request.form.get("phone",r["phone"] or "").strip()
  if not surname or not firstname or len(surname)>100 or len(firstname)>100:
   c.close(); flash("Surname and first name are required and must be at most 100 characters."); return redirect("/profile/"+str(rid))
  if len(birth_place)>120 or len(phone)>40:
   c.close(); flash("Birth place must be at most 120 characters and phone at most 40 characters."); return redirect("/profile/"+str(rid))
  if birth_date:
   try: parsed_birth_date=date.fromisoformat(birth_date)
   except ValueError:
    c.close(); flash("Enter a valid birth date."); return redirect("/profile/"+str(rid))
   if parsed_birth_date>date.today():
    c.close(); flash("Birth date cannot be in the future."); return redirect("/profile/"+str(rid))
  # Keep the stored image unless a new valid upload is explicitly selected.
  photo_path=(r["photo_path"] or "").strip()
  photo_file=request.files.get("photo")
  if photo_file and photo_file.filename:
   ext=os.path.splitext(secure_filename(photo_file.filename))[1].lower()
   expected_formats={".png":"PNG",".jpg":"JPEG",".jpeg":"JPEG",".webp":"WEBP"}
   if ext not in expected_formats:
    c.close(); flash("Photo must be PNG, JPG or WEBP."); return redirect("/profile/"+str(rid))
   photo_bytes=photo_file.stream.read(4*1024*1024+1); photo_file.stream.seek(0)
   if not photo_bytes or len(photo_bytes)>4*1024*1024:
    c.close(); flash("Photo must be non-empty and 4 MB or smaller."); return redirect("/profile/"+str(rid))
   try:
    with warnings.catch_warnings():
     warnings.simplefilter("error", Image.DecompressionBombWarning)
     with Image.open(io.BytesIO(photo_bytes)) as probe:
      if probe.format != expected_formats[ext]:
       raise ValueError("Image format does not match its filename.")
      probe.verify()
     with Image.open(io.BytesIO(photo_bytes)) as source:
      image=ImageOps.exif_transpose(source)
      image.load()
      image.thumbnail((1200,1200))
      output=io.BytesIO()
      has_alpha="A" in image.getbands()
      if expected_formats[ext]=="JPEG":
       image.convert("RGB").save(output,format="JPEG",quality=90,optimize=True)
      elif expected_formats[ext]=="PNG":
       image.convert("RGBA" if has_alpha else "RGB").save(output,format="PNG",optimize=True)
      else:
       image.convert("RGBA" if has_alpha else "RGB").save(output,format="WEBP",quality=90,method=6)
      normalized_photo=output.getvalue()
   except Exception:
    c.close(); flash("The selected file is not a readable PNG/JPG/WEBP image. Please choose a real photo."); return redirect("/profile/"+str(rid))
   photo_dir=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos"); os.makedirs(photo_dir,exist_ok=True)
   photo_name="resident_"+str(rid)+"_"+str(int(time.time()))+"_"+secrets.token_hex(5)+ext
   target=os.path.join(photo_dir,photo_name); temporary=target+".tmp"
   try:
    with open(temporary,"wb") as handle: handle.write(normalized_photo)
    os.replace(temporary,target)
   finally:
    if os.path.exists(temporary): os.remove(temporary)
   photo_path="/photos/"+photo_name
  c.execute("UPDATE residents SET surname=?,firstname=?,birth_place=?,birth_date=?,photo_path=?,phone=? WHERE id=?",
   (surname,firstname,birth_place,birth_date,photo_path,phone,rid))
  audit_event(c,"resident.profile_updated",f"resident_id={rid}; name={surname} {firstname}; photo_uploaded={bool(photo_file and photo_file.filename)}")
  c.commit(); r=c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone(); flash(tr("done"))
 history_rows=c.execute("SELECT h.*,t.name terrain FROM history h JOIN terrains t ON t.id=h.terrain_id WHERE h.resident_id=? ORDER BY h.id DESC",(rid,)).fetchall()
 assignment_rows=c.execute("SELECT a.status,p.name period,p.year_label,t.name terrain FROM assignments a JOIN periods p ON p.id=a.period_id JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=? ORDER BY p.year_label DESC,p.id",(rid,)).fetchall()
 objectives=c.execute("SELECT o.*,t.name terrain FROM objectives o JOIN terrains t ON t.id=o.terrain_id ORDER BY o.study_year,t.name").fetchall()
 exam_rows=c.execute("SELECT * FROM exam_records WHERE resident_id=? ORDER BY study_year,module",(rid,)).fetchall()
 done={x[0] for x in c.execute("SELECT terrain_id FROM history WHERE resident_id=? UNION SELECT terrain_id FROM assignments WHERE resident_id=? AND status='confirmed'",(rid,rid)).fetchall()}
 terrains=c.execute("SELECT * FROM terrains ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END,name").fetchall(); c.close()
 return page("""<section class="rt-profile-panel"><div class="rt-profile-head"><div class="rt-avatar-wrap">{% if r.photo_path %}<img id="resident-photo-preview" class="rt-profile-avatar" src="{{r.photo_path}}" alt="Resident photo">{% else %}<div id="resident-photo-placeholder" class="rt-profile-avatar rt-avatar-placeholder" aria-label="No resident photo">👤</div><img id="resident-photo-preview" class="rt-profile-avatar" alt="Selected resident photo preview" hidden>{% endif %}</div><div class="rt-profile-heading"><span class="rt-eyebrow">RESIMANAGE · RESIDENT FILE</span><h2>{{r.surname}} {{r.firstname}}</h2><p><span class="rt-chip">Year {{r.year}}</span><span class="rt-chip {{'rt-chip-muted' if r.status!='active' else 'rt-chip-active'}}">{{r.status}}</span></p><p class="muted">Profile, rotation progress, exam results and objectives in one place.</p></div></div><form class="rt-profile-form" method="post" enctype="multipart/form-data" accept-charset="UTF-8"><div class="rt-field-grid"><label>{{t('surname')}}<input name="surname" value="{{r.surname}}" autocomplete="family-name" maxlength="100" required></label><label>{{t('firstname')}}<input name="firstname" value="{{r.firstname}}" autocomplete="given-name" maxlength="100" required></label><label>Birth place / Lieu de naissance / مكان الميلاد<input name="birth_place" value="{{r.birth_place or ''}}" maxlength="120"></label><label>Birth date / Date de naissance / تاريخ الميلاد<input type="date" name="birth_date" value="{{r.birth_date or ''}}"></label><label>Phone / Téléphone / الهاتف<input name="phone" value="{{r.phone or ''}}" type="tel" autocomplete="tel" maxlength="40"></label></div><div class="rt-photo-picker"><label for="resident-photo-input">Resident photo / Photo du résident / صورة المقيم</label><p class="muted">Choose a real PNG, JPG or WEBP image (maximum 4 MB). A preview appears before saving.</p><input id="resident-photo-input" type="file" name="photo" accept=".png,.jpg,.jpeg,.webp,image/png,image/jpeg,image/webp" aria-describedby="resident-photo-help"><small id="resident-photo-help">The existing photo is kept if you do not select a replacement.</small><small id="resident-photo-status" aria-live="polite"></small></div><div class="rt-savebar"><button type="submit">Save profile / Enregistrer / حفظ الملف</button><a class="rt-secondary-action" href="/residents">Back to residents / Retour / رجوع</a></div></form><p class="rt-profile-links"><a href="/exams?resident={{r.id}}">Exams / Examens / الامتحانات</a><a href="/dems">DEMS archive</a></p></section><section><h3>Terrain progress / Avancement / تقدم التربصات</h3><div class="wrap"><table><tr><th>Terrain</th><th>Status</th></tr>{% for t in terrains %}<tr><td>{{t.name}}</td><td class="{{'bad' if t.id in done else 'good'}}">{{'Completed / Effectué / مكتمل' if t.id in done else 'Remaining / À faire / متبقٍ'}}</td></tr>{% endfor %}</table></div></section><section><h3>Assignments and history / Affectations et historique / التوزيعات والتاريخ</h3>{% for a in assignment_rows %}<p>{{a.year_label}} — {{a.period}} — {{a.terrain}} ({{a.status}})</p>{% endfor %}{% for h in history_rows %}<p>{{h.year_label}} — {{h.period}} — {{h.terrain}} — {{h.date_done}}</p>{% endfor %}</section><section><h3>Exam scores / Notes d’examen / علامات الامتحان</h3><div class="wrap"><table><tr><th>Study year</th><th>Module</th><th>Theory</th><th>Practical</th></tr>{% for e in exam_rows %}<tr><td>{{e.study_year}}</td><td>{{e.module}}</td><td>{{e.theory}}/20</td><td>{{e.practical}}/20</td></tr>{% else %}<tr><td colspan="4">{{t('none')}}</td></tr>{% endfor %}</table></div></section><section><h3>Objectives / Objectifs / الأهداف</h3>{% for o in objectives %}<p>Year {{o.study_year}} — {{o.terrain}}: {{o.objective}}</p>{% else %}<p>{{t('none')}}</p>{% endfor %}</section><script>(function(){const input=document.getElementById('resident-photo-input');const preview=document.getElementById('resident-photo-preview');const placeholder=document.getElementById('resident-photo-placeholder');if(!input||!preview)return;const status=document.getElementById('resident-photo-status');input.addEventListener('change',function(){const file=input.files&&input.files[0];if(!file)return;const allowed=['image/png','image/jpeg','image/webp'];const ext=(file.name.split('.').pop()||'').toLowerCase();if((file.type&&!allowed.includes(file.type))||!['png','jpg','jpeg','webp'].includes(ext)){if(status)status.textContent='Please choose a PNG, JPG or WEBP photo.';input.value='';return;}if(file.size>4*1024*1024){if(status)status.textContent='The photo must be 4 MB or smaller.';input.value='';return;}const reader=new FileReader();reader.onload=function(){preview.src=reader.result;preview.hidden=false;if(placeholder)placeholder.hidden=true;if(status)status.textContent='Preview ready. Save the profile to upload the photo.';};reader.readAsDataURL(file);});})();</script>""",r=r,terrains=terrains,done=done,assignment_rows=assignment_rows,history_rows=history_rows,objectives=objectives,exam_rows=exam_rows)
@app.route("/objectives",methods=["GET","POST"])
def manage_objectives():
 if super_guard(): return super_guard()
 c=db()
 if request.method=="POST":
  try:
   tid=int(request.form["terrain"]); sy=int(request.form["study_year"]); obj=request.form.get("objective","").strip()
   terrain=c.execute("SELECT id FROM terrains WHERE id=? AND active=1",(tid,)).fetchone()
   if not terrain: raise ValueError("Choose an active terrain.")
   if sy not in (1,2,3,4): raise ValueError("Objective study year must be 1, 2, 3 or 4.")
   if not obj or len(obj)>1000: raise ValueError("Objective text is required and must be at most 1000 characters.")
   duplicate=c.execute("SELECT 1 FROM objectives WHERE terrain_id=? AND study_year=? AND lower(objective)=lower(?)",(tid,sy,obj)).fetchone()
   if duplicate: raise ValueError("This objective already exists for the selected terrain and year.")
   c.execute("INSERT INTO objectives(terrain_id,study_year,objective) VALUES(?,?,?)",(tid,sy,obj))
   audit_event(c,"objective.created",f"terrain_id={tid}; study_year={sy}")
   c.commit(); flash(tr("done"))
  except (ValueError,TypeError,KeyError,sqlite3.IntegrityError) as exc:
   c.rollback(); flash(str(exc) or "Invalid objective input.")
 ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY name").fetchall()
 rows=c.execute("SELECT o.*,t.name terrain FROM objectives o JOIN terrains t ON t.id=o.terrain_id ORDER BY study_year,t.name").fetchall(); c.close()
 return page("""<section><h3>{{t('objectives')}}</h3><form method="post"><label>{{t('terrain')}}</label><select name="terrain">{% for t in ts %}<option value="{{t.id}}">{{t.name}}</option>{% endfor %}</select><label>{{t('study_year')}}</label><select name="study_year">{% for y in range(1,5) %}<option>{{y}}</option>{% endfor %}</select><label>{{t('objective')}}</label><input name="objective" required><button>{{t('save')}}</button></form></section><section>{% for o in rows %}<p>Year {{o.study_year}} — {{o.terrain}}: {{o.objective}}</p>{% else %}<p>{{t('none')}}</p>{% endfor %}</section>""",ts=ts,rows=rows)
@app.route("/exams",methods=["GET","POST"])
def exams():
 if guard(): return redirect("/login")
 c=db()
 if request.method=="POST":
  action=request.form.get("action","record")
  try:
   rid=int(request.form.get("resident",""))
   sy=int(request.form.get("study_year",""))
  except (TypeError,ValueError):
   c.close(); flash("Choose a valid resident and study year."); return redirect("/exams")
  resident=c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone()
  if not resident:
   flash("Resident not found")
  elif action=="record":
   if session.get("role")!="super_admin":
    c.close(); return ("Forbidden — Super Admin only",403)
   try:
    module=request.form.get("module","").strip()
    theory=float(request.form.get("theory",""))
    practical=float(request.form.get("practical",""))
   except (TypeError,ValueError):
    c.close(); flash("Module and numeric marks from 0 to 20 are required."); return redirect("/exams")
   if resident["status"]=="graduated" or sy>resident["year"]:
    flash("Cannot record marks for a future study year or modify a DEMS-graduated resident's archived record.")
   elif sy not in (1,2,3,4) or not module or not (0<=theory<=20 and 0<=practical<=20):
    flash("Module and marks from 0 to 20 are required")
   else:
    existing=c.execute("SELECT id FROM exam_records WHERE resident_id=? AND study_year=? AND lower(module)=lower(?)",(rid,sy,module)).fetchone()
    if existing:
     c.execute("UPDATE exam_records SET theory=?,practical=?,created_at=CURRENT_TIMESTAMP WHERE id=?",(theory,practical,existing["id"]))
     audit_event(c,"exam.updated",f"resident_id={rid}; year={sy}; module={module}")
    else:
     c.execute("INSERT INTO exam_records(resident_id,study_year,module,theory,practical) VALUES(?,?,?,?,?)",(rid,sy,module,theory,practical))
     audit_event(c,"exam.recorded",f"resident_id={rid}; year={sy}; module={module}")
    c.commit(); flash(tr("done"))
  elif action=="evaluate_first_year" and request.form.get("all_modules_confirmed")!="yes":
   flash("Confirm that all first-year module results have been entered before evaluating progression.")
  elif action=="evaluate_first_year":
   if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
   if resident["year"]!=1:
    flash("Only current first-year residents can be evaluated for first-year progression.")
   elif resident["status"]!="active":
    flash("The resident is inactive; progression is not available.")
   else:
    records=c.execute("SELECT theory,practical FROM exam_records WHERE resident_id=? AND study_year=1",(rid,)).fetchall()
    if not records:
     flash("Enter first-year exam modules before evaluation")
    else:
     average=sum((x["theory"]+x["practical"])/2 for x in records)/len(records)
     passed=average>=10 and all(x["theory"]>=10 and x["practical"]>=10 for x in records)
     c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(f"{resident['surname']} {resident['firstname']} — first-year exams","Passed" if passed else "Failed",f"Average={average:.2f}/20; all theory/practical marks >=10: {all(x['theory']>=10 and x['practical']>=10 for x in records)}",session["user"]))
     audit_event(c,"exam.first_year_evaluated",f"resident_id={rid}; average={average:.2f}; passed={passed}")
     advanced=False
     if passed and resident["year"]==1 and resident["status"]=="active":
      c.execute("UPDATE residents SET year=2 WHERE id=? AND year=1 AND status='active'",(rid,))
      advanced=c.execute("SELECT changes()").fetchone()[0]==1
      if advanced:
       audit_event(c,"resident.year_advanced",f"resident_id={rid}; from=1; to=2")
     c.commit()
     if not passed:
      flash("Failed: average must be at least 10/20 and every theory/practical mark must be at least 10/20")
     elif resident["year"]!=1:
      flash("Exam criteria passed; no year change because the resident is not currently in year 1.")
     elif resident["status"]!="active":
      flash("Exam criteria passed, but the resident is inactive; no progression was applied.")
     elif advanced:
      flash("Passed; resident advanced to year 2.")
     else:
      flash("Exam criteria passed, but progression was not applied; review the resident record.")
  elif action=="approve_advance":
   actor=session.get("user","").lower()
   outcome=request.form.get("outcome","approve")
   reason=request.form.get("reason","").strip()
   if session.get("role")!="super_admin" or actor not in ("benhocine","drdris"):
    c.close(); return ("Only the official Super Admins Dr Dris and Pr Benhocine may vote on progression.",403)
   if outcome not in ("approve","reject") or not reason:
    flash("Choose approve/reject and enter the decision reason.")
   elif resident["status"]!="active":
    flash("Inactive residents cannot be advanced; no progression vote was recorded.")
   elif sy not in (2,3) or resident["year"]!=sy:
    flash("Only active year 2 or 3 residents can use the dual-official progression vote.")
   else:
    c.execute("INSERT INTO promotion_votes(resident_id,from_year,decision,reason,actor) VALUES(?,?,?,?,?) ON CONFLICT(resident_id,from_year,actor) DO UPDATE SET decision=excluded.decision,reason=excluded.reason,created_at=CURRENT_TIMESTAMP",(rid,sy,outcome,reason,actor))
    audit_event(c,"promotion.vote_recorded",f"resident_id={rid}; from_year={sy}; decision={outcome}; reason={reason}")
    votes=c.execute("SELECT actor,decision,reason FROM promotion_votes WHERE resident_id=? AND from_year=? AND actor IN ('benhocine','drdris')",(rid,sy)).fetchall()
    if len(votes)<2:
     c.commit(); flash("Vote recorded. Waiting for the other official administrator.")
    else:
     combined_reason=" | ".join(v["actor"]+": "+v["decision"]+" — "+v["reason"] for v in votes)
     actor_names="+".join(sorted(v["actor"] for v in votes))
     approved=all(v["decision"]=="approve" for v in votes)
     if approved:
      c.execute("UPDATE residents SET year=? WHERE id=? AND year=? AND status='active'",(sy+1,rid,sy))
      if c.execute("SELECT changes()").fetchone()[0]:
       c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(f"{resident['surname']} {resident['firstname']} — year {sy}","Approved by both official administrators",combined_reason,actor_names))
       audit_event(c,"resident.year_advanced",f"resident_id={rid}; from={sy}; to={sy+1}; officials={actor_names}")
       flash("Approved by both official administrators; resident advanced.")
      else:
       flash("Resident status changed before the vote was finalized; no progression was applied.")
     else:
      c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(f"{resident['surname']} {resident['firstname']} — year {sy}","Rejected by official vote",combined_reason,actor_names))
      audit_event(c,"promotion.vote_rejected",f"resident_id={rid}; from_year={sy}; officials={actor_names}")
      flash("Progression not approved; both official votes are recorded.")
     c.execute("DELETE FROM promotion_votes WHERE resident_id=? AND from_year=?",(rid,sy))
     c.commit()
 rows=c.execute("SELECT e.*,r.surname,r.firstname FROM exam_records e JOIN residents r ON r.id=e.resident_id ORDER BY e.created_at DESC,e.id DESC").fetchall()
 rs=c.execute("SELECT * FROM residents WHERE status='active' ORDER BY year DESC,surname").fetchall()
 selected=request.args.get("resident",type=int)
 pending_votes=c.execute("SELECT v.*,r.surname,r.firstname FROM promotion_votes v JOIN residents r ON r.id=v.resident_id ORDER BY v.created_at DESC").fetchall()
 c.close()
 return page("""<section><h3>{{t('exams')}}</h3><form method="post"><input type="hidden" name="action" value="record"><label>Resident</label><select name="resident">{% for r in rs %}<option value="{{r.id}}" {% if selected==r.id %}selected{% endif %}>{{r.year}} — {{r.surname}} {{r.firstname}}</option>{% endfor %}</select><label>Study year</label><select name="study_year">{% for y in range(1,5) %}<option>{{y}}</option>{% endfor %}</select><label>{{t('module')}}</label><input name="module" required><div class="grid"><div><label>{{t('theory')}}</label><input type="number" name="theory" min="0" max="20" step="0.01" required></div><div><label>{{t('practical')}}</label><input type="number" name="practical" min="0" max="20" step="0.01" required></div></div><button>Record exam / Enregistrer / تسجيل</button></form></section><section><h4>Evaluate first year / Évaluer la 1re année / تقييم السنة الأولى</h4><form method="post"><input type="hidden" name="action" value="evaluate_first_year"><label>Resident</label><select name="resident">{% for r in rs if r.year==1 %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}}</option>{% endfor %}</select><input type="hidden" name="study_year" value="1"><label><input type="checkbox" name="all_modules_confirmed" value="yes" required> Confirm all module results are entered / Confirmer toutes les notes / تأكيد إدخال جميع العلامات</label><button>Evaluate and advance if eligible</button></form><h4>Progression vote — years 2–3 (Dr Dris + Pr Benhocine)</h4><form method="post"><input type="hidden" name="action" value="approve_advance"><label>Resident</label><select name="resident" required>{% for r in rs if r.year in [2,3] %}<option value="{{r.id}}">{{r.year}} — {{r.surname}} {{r.firstname}}</option>{% endfor %}</select><label>Study year</label><select name="study_year"><option>2</option><option>3</option></select><label>Vote</label><select name="outcome"><option value="approve">Approve</option><option value="reject">Reject</option></select><label>Reason</label><input name="reason" required maxlength="500"><button>Record official vote</button></form><h4>Pending official votes</h4><div class="wrap"><table><tr><th>Resident</th><th>From year</th><th>Official</th><th>Vote</th><th>Reason</th></tr>{% for v in pending_votes %}<tr><td>{{v.surname}} {{v.firstname}}</td><td>{{v.from_year}}</td><td>{{v.actor}}</td><td>{{v.decision}}</td><td>{{v.reason}}</td></tr>{% else %}<tr><td colspan="5">No pending official votes</td></tr>{% endfor %}</table></div></section><section><div class="wrap"><table><tr><th>Resident</th><th>Year</th><th>Module</th><th>Theory</th><th>Practical</th></tr>{% for e in rows %}<tr><td>{{e.surname}} {{e.firstname}}</td><td>{{e.study_year}}</td><td>{{e.module}}</td><td>{{e.theory}}</td><td>{{e.practical}}</td></tr>{% else %}<tr><td colspan="5">{{t('none')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows,rs=rs,selected=selected,pending_votes=pending_votes)
@app.route("/dems",methods=["GET","POST"])
def dems():
 if super_guard(): return super_guard()
 c=db()
 if request.method=="POST":
  actor=session.get("user","").lower()
  if actor not in ("benhocine","drdris"):
   c.close(); return ("Only Dr Dris and Pr Benhocine may vote on DEMS success.",403)
  rid=request.form.get("resident",type=int)
  outcome=request.form.get("outcome","approve")
  reason=request.form.get("reason","").strip()
  try: success_date=date.fromisoformat(request.form.get("success_date","").strip()).isoformat()
  except ValueError:
   c.close(); flash("Enter a valid DEMS exam date."); return redirect("/dems")
  if date.fromisoformat(success_date)>date.today():
   c.close(); flash("The DEMS exam date cannot be in the future."); return redirect("/dems")
  r=c.execute("SELECT * FROM residents WHERE id=? AND year=4 AND status='active'",(rid,)).fetchone()
  if not r or outcome not in ("approve","reject") or not reason:
   flash("Choose an active year-4 resident, a valid vote, a date and a reason.")
  else:
   c.execute("INSERT INTO dems_votes(resident_id,success_date,decision,reason,actor) VALUES(?,?,?,?,?) ON CONFLICT(resident_id,success_date,actor) DO UPDATE SET decision=excluded.decision,reason=excluded.reason,created_at=CURRENT_TIMESTAMP",(rid,success_date,outcome,reason,actor))
   audit_event(c,"dems.vote_recorded",f"resident_id={rid}; success_date={success_date}; decision={outcome}; reason={reason}")
   votes=c.execute("SELECT actor,decision,reason FROM dems_votes WHERE resident_id=? AND success_date=? AND actor IN ('benhocine','drdris')",(rid,success_date)).fetchall()
   if len(votes)<2:
    c.commit(); flash("DEMS vote recorded. Waiting for the other official administrator.")
   else:
    combined_reason=" | ".join(v["actor"]+": "+v["decision"]+" — "+v["reason"] for v in votes)
    actor_names="+".join(sorted(v["actor"] for v in votes))
    approved=all(v["decision"]=="approve" for v in votes)
    if approved:
     hist=c.execute("SELECT h.*,t.name terrain FROM history h JOIN terrains t ON t.id=h.terrain_id WHERE h.resident_id=?",(rid,)).fetchall()
     asgs=c.execute("SELECT a.*,p.name period,p.year_label,t.name terrain FROM assignments a JOIN periods p ON p.id=a.period_id JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=?",(rid,)).fetchall()
     exams_snapshot=c.execute("SELECT * FROM exam_records WHERE resident_id=? ORDER BY study_year,module",(rid,)).fetchall()
     profile={k:r[k] for k in r.keys()}
     c.execute("INSERT INTO dems_archive(resident_id,success_date,profile_json,history_json) VALUES(?,?,?,?)",(rid,success_date,json.dumps(profile,ensure_ascii=False),json.dumps({"history":[dict(x) for x in hist],"assignments":[dict(x) for x in asgs],"exams":[dict(x) for x in exams_snapshot]},ensure_ascii=False)))
     c.execute("UPDATE residents SET status='graduated',dems_status='success' WHERE id=?",(rid,))
     c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(f"{r['surname']} {r['firstname']} — DEMS","Success approved by both official administrators",combined_reason,actor_names))
     audit_event(c,"dems.archived",f"resident_id={rid}; success_date={success_date}; officials={actor_names}")
     flash("DEMS success approved by both officials; profile and history archived.")
    else:
     c.execute("INSERT INTO decisions(details,decision,reason,actor) VALUES(?,?,?,?)",(f"{r['surname']} {r['firstname']} — DEMS","Not approved by official vote",combined_reason,actor_names))
     audit_event(c,"dems.vote_rejected",f"resident_id={rid}; success_date={success_date}; officials={actor_names}")
     flash("DEMS success was not approved; both official votes are recorded.")
    c.execute("DELETE FROM dems_votes WHERE resident_id=? AND success_date=?",(rid,success_date))
    c.commit()
 residentsrow_numbers=c.execute("SELECT * FROM residents WHERE year=4 AND status='active' ORDER BY surname").fetchall()
 archiverow_numbers=c.execute("SELECT d.*,r.surname,r.firstname FROM dems_archive d JOIN residents r ON r.id=d.resident_id ORDER BY d.id DESC").fetchall()
 pending_votes=c.execute("SELECT v.*,r.surname,r.firstname FROM dems_votes v JOIN residents r ON r.id=v.resident_id ORDER BY v.created_at DESC").fetchall()
 c.close()
 return page("""<section><h3>{{t('dems')}}</h3><p>DEMS success is finalized only after both Dr Dris and Pr Benhocine approve.</p><form method="post"><label>Year-4 resident</label><select name="resident" required>{% for r in residentsrow_numbers %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}}</option>{% endfor %}</select><label>{{t('success_date')}}</label><input type="date" name="success_date" required><label>Vote</label><select name="outcome"><option value="approve">Approve success</option><option value="reject">Reject success</option></select><label>Reason</label><input name="reason" required maxlength="500"><button>Record official vote</button></form></section><section><h4>Pending official votes</h4><div class="wrap"><table><tr><th>Resident</th><th>Exam date</th><th>Official</th><th>Vote</th><th>Reason</th></tr>{% for v in pending_votes %}<tr><td>{{v.surname}} {{v.firstname}}</td><td>{{v.success_date}}</td><td>{{v.actor}}</td><td>{{v.decision}}</td><td>{{v.reason}}</td></tr>{% else %}<tr><td colspan="5">No pending official votes</td></tr>{% endfor %}</table></div></section><section><h4>Archived graduates</h4>{% for a in archiverow_numbers %}<p>{{a.surname}} {{a.firstname}} — DEMS success: {{a.success_date}} — archived {{a.archived_at}}</p>{% else %}<p>{{t('none')}}</p>{% endfor %}</section>""",residentsrow_numbers=residentsrow_numbers,archiverow_numbers=archiverow_numbers,pending_votes=pending_votes)
@app.get("/export")
def export():
 if super_guard():return super_guard()
 c=db(); out=io.StringIO(); w=csv.writer(out)
 for table in ["residents","terrains","periods","capacities","assignments","history","decisions","exam_records","objectives","dems_archive","leave_years","leave_entries","external_rotations","promotion_votes","dems_votes","audit"]:
  w.writerow([table])
  w.writerow([x[1] for x in c.execute('PRAGMA table_info("'+table+'")')])
  for row in c.execute('SELECT * FROM "'+table+'"'):w.writerow(list(row))
 c.close(); f=io.BytesIO(("\ufeff"+out.getvalue()).encode("utf-8")); return send_file(f,mimetype="text/csv",as_attachment=True,download_name="rotation_export.csv")
def create_full_backup_bytes(directory):
 with tempfile.TemporaryDirectory(prefix="resimanage_backup_",dir=directory) as workdir:
  snapshot=os.path.join(workdir,"rotation.db")
  source=sqlite3.connect(DB,timeout=20)
  destination=sqlite3.connect(snapshot)
  try:
   source.backup(destination); destination.commit()
  finally:
   destination.close(); source.close()
  output=io.BytesIO()
  with zipfile.ZipFile(output,"w",compression=zipfile.ZIP_DEFLATED) as bundle:
   bundle.write(snapshot,"rotation.db")
   bundle.writestr("resident_photos/","")
   photo_dir=os.path.join(directory,"resident_photos")
   if os.path.isdir(photo_dir):
    for name in sorted(os.listdir(photo_dir)):
     path=os.path.join(photo_dir,name)
     if os.path.isfile(path) and not os.path.islink(path) and os.path.basename(name)==name:
      bundle.write(path,"resident_photos/"+name)
  return output.getvalue()

@app.get("/backup")
def backup():
 if super_guard(): return super_guard()
 directory=os.path.dirname(os.path.abspath(DB))
 data=create_full_backup_bytes(directory)
 return send_file(io.BytesIO(data),mimetype="application/zip",as_attachment=True,download_name="resimanage_full_backup.zip")

@app.route("/restore",methods=["GET","POST"])
def restore():
 if super_guard(): return super_guard()
 if request.method=="POST":
  uploaded=request.files.get("backup_file")
  if not uploaded or not uploaded.filename:
   flash("Choose a .zip full backup or a legacy .db/.sqlite backup.")
   return redirect("/restore")
  extension=os.path.splitext(uploaded.filename.lower())[1]
  if extension not in (".zip",".db",".sqlite",".sqlite3"):
   flash("Unsupported backup type. Choose .zip, .db or .sqlite.")
   return redirect("/restore")
  directory=os.path.dirname(os.path.abspath(DB))
  workdir=None; old_db_path=None; old_photo_path=None; staged_db=None; staged_photos=None
  db_replaced=False; photos_replaced=False
  try:
   workdir=tempfile.mkdtemp(prefix="resimanage_restore_",dir=directory)
   uploaded_path=os.path.join(workdir,"uploaded"+extension)
   uploaded.save(uploaded_path)
   staged_db=os.path.join(workdir,"rotation.db")
   restore_photos=False
   if extension==".zip":
    with zipfile.ZipFile(uploaded_path,"r") as bundle:
     infos=bundle.infolist()
     names=[info.filename for info in infos]
     if names.count("rotation.db")!=1:
      raise ValueError("Full backup must contain exactly one rotation.db.")
     if sum(info.file_size for info in infos)>100*1024*1024:
      raise ValueError("Backup expands beyond the 100 MB safety limit.")
     for info in infos:
      name=info.filename
      if name=="rotation.db" or name=="resident_photos/" or (name.startswith("resident_photos/") and not name.endswith("/")):
       if ".." in name.split("/") or name.startswith("/"):
        raise ValueError("Backup contains an unsafe path.")
      else:
       raise ValueError("Backup contains an unsupported file path.")
     with bundle.open("rotation.db") as source, open(staged_db,"wb") as destination:
      shutil.copyfileobj(source,destination)
     restore_photos="resident_photos/" in names or any(name.startswith("resident_photos/") for name in names)
     if restore_photos:
      staged_photos=os.path.join(workdir,"resident_photos")
      os.makedirs(staged_photos,exist_ok=True)
      for info in infos:
       name=info.filename
       if name.startswith("resident_photos/") and not name.endswith("/"):
        photo_name=name[len("resident_photos/"):]
        if not photo_name or "/" in photo_name or photo_name in (".","..") or info.file_size>8*1024*1024:
         raise ValueError("Backup contains an invalid or oversized resident photo.")
        with bundle.open(info) as source, open(os.path.join(staged_photos,photo_name),"wb") as destination:
         shutil.copyfileobj(source,destination)
   else:
    shutil.copyfile(uploaded_path,staged_db)
   check=sqlite3.connect("file:"+staged_db+"?mode=ro",uri=True,timeout=10)
   integrity=check.execute("PRAGMA integrity_check").fetchone()[0]
   tables={row[0] for row in check.execute("SELECT name FROM sqlite_master WHERE type='table'")}
   check.close()
   required={"residents","terrains","periods","assignments","history","settings"}
   if integrity!="ok" or not required.issubset(tables):
    raise ValueError("The uploaded file is not a valid Rotation App database backup.")
   timestamp=datetime.now().strftime("%Y%m%d_%H%M%S")
   full_before=os.path.join(directory,"rotation_backup_before_restore_"+timestamp+".zip")
   with open(full_before,"wb") as handle:
    handle.write(create_full_backup_bytes(directory))
   old_db_path=os.path.join(directory,"rotation_backup_before_restore_"+timestamp+".db")
   shutil.copy2(DB,old_db_path)
   photo_dir=os.path.join(directory,"resident_photos")
   if restore_photos and os.path.exists(photo_dir):
    old_photo_path=os.path.join(directory,"resident_photos_before_restore_"+timestamp)
    os.replace(photo_dir,old_photo_path)
   os.replace(staged_db,DB); staged_db=None; db_replaced=True
   if restore_photos:
    os.replace(staged_photos,photo_dir); staged_photos=None; photos_replaced=True
   app._ready=False
   init()
   c=db(); audit_event(c,"database.restored",f"Restored from {extension} backup; pre-restore full backup={os.path.basename(full_before)}"); c.commit(); c.close()
   flash("Restore completed. Pre-restore full backup saved as "+os.path.basename(full_before))
  except Exception as exc:
   if db_replaced and old_db_path and os.path.exists(old_db_path):
    try: shutil.copy2(old_db_path,DB); app._ready=False
    except Exception: pass
   if photos_replaced:
    photo_dir=os.path.join(directory,"resident_photos")
    try:
     if os.path.isdir(photo_dir): shutil.rmtree(photo_dir)
    except Exception: pass
   if old_photo_path and os.path.exists(old_photo_path):
    try: os.replace(old_photo_path,os.path.join(directory,"resident_photos"))
    except Exception: pass
   flash("Restore failed; a pre-restore full backup was saved if validation had completed. Details: "+str(exc)[:220])
  finally:
   if workdir and os.path.exists(workdir): shutil.rmtree(workdir,ignore_errors=True)
  return redirect("/restore")
 return page("""<section><h3>{{t('restore')}}</h3><p>{{t('restore_help')}} Upload a full .zip backup (database and resident photos) or a legacy .db/.sqlite database backup.</p><form method="post" enctype="multipart/form-data"><label>{{t('file')}} (.zip/.db/.sqlite)</label><input type="file" name="backup_file" accept=".zip,.db,.sqlite,.sqlite3" required><button>{{t('restore')}}</button></form></section>""")
def rotation_period_bounds(year_label, period_name):
    try:
        y = int(str(year_label).split("-")[0])
    except (ValueError, IndexError):
        y = date.today().year
    n = str(period_name).lower()
    if any(x in n for x in ("fév", "fev", "feb")): return date(y,2,1), date(y,4,30)
    if any(x in n for x in ("mai", "may")): return date(y,5,1), date(y,7,30)
    if any(x in n for x in ("août", "aout", "aug")): return date(y,8,1), date(y,10,30)
    if "nov" in n: return date(y,11,1), date(y+1,1,30)
    return date(y,2,1), date(y+1,1,30)

def reject_external_overlap_with_local_assignment(c, resident_id, start, end):
    local_rows = c.execute(
        """SELECT a.id,t.name terrain,p.year_label,p.name period
           FROM assignments a
           JOIN terrains t ON t.id=a.terrain_id
           JOIN periods p ON p.id=a.period_id
           WHERE a.resident_id=? AND a.status IN ('confirmed','pending')""",
        (resident_id,),
    ).fetchall()
    for assignment in local_rows:
        local_start, local_end = rotation_period_bounds(
            assignment["year_label"], assignment["period"]
        )
        if start <= local_end and end >= local_start:
            raise ValueError(
                "External rotation overlaps an existing local assignment: "
                + assignment["terrain"] + " / " + assignment["period"] + "."
            )


@app.route("/leave", methods=["GET","POST"])
def leave_management():
    if guard(): return redirect("/login")
    c=db()
    if request.method=="POST":
        if super_guard(): c.close(); return ("Forbidden — Super Admin only",403)
        try:
            rid=int(request.form["resident_id"])
            start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
            if end<start: raise ValueError("End date must not precede start date.")
            days=(end-start).days+1
            overlap=c.execute("SELECT id FROM leave_entries WHERE resident_id=? AND start_date<=? AND end_date>=?",(rid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError("Leave dates overlap an existing entry.")
            for yr in range(start.year,end.year+1):
                c.execute("INSERT OR IGNORE INTO leave_years(resident_id,year,earned_days,created_by) VALUES(?,?,30,?)",(rid,yr,session.get("user","")))
            earned=c.execute("SELECT COALESCE(SUM(earned_days),0) FROM leave_years WHERE resident_id=? AND year<=?",(rid,end.year)).fetchone()[0]
            used=c.execute("SELECT COALESCE(SUM(days),0) FROM leave_entries WHERE resident_id=? AND end_date<=?",(rid,date(end.year,12,31).isoformat())).fetchone()[0]
            if used+days>earned: raise ValueError("Insufficient accrued leave balance, including carry-forward.")
            c.execute("INSERT INTO leave_entries(resident_id,start_date,end_date,days,notes,created_by) VALUES(?,?,?,?,?,?)",(rid,start.isoformat(),end.isoformat(),days,request.form.get("notes","").strip(),session.get("user","")))
            audit_event(c,"leave.entry_created",f"resident_id={rid}; start={start}; end={end}; days={days}")
            c.commit(); flash(f"Leave recorded: {days} calendar days. Unused entitlement carries forward.")
        except Exception as exc: c.rollback(); flash("Leave not saved: "+str(exc)[:220])
        c.close(); return redirect("/leave")
    residents=c.execute("SELECT id,surname,firstname FROM residents ORDER BY year DESC,surname").fetchall()
    now=date.today(); balances=[]
    for r in residents:
        first=c.execute("SELECT MIN(year) FROM leave_years WHERE resident_id=?",(r["id"],)).fetchone()[0]
        first=first if first is not None else now.year
        for yr in range(first,now.year+1):
            c.execute("INSERT OR IGNORE INTO leave_years(resident_id,year,earned_days,created_by) VALUES(?,?,30,?)",(r["id"],yr,session.get("user","")))
        earned=c.execute("SELECT COALESCE(SUM(earned_days),0) FROM leave_years WHERE resident_id=? AND year<=?",(r["id"],now.year)).fetchone()[0]
        used=c.execute("SELECT COALESCE(SUM(days),0) FROM leave_entries WHERE resident_id=? AND end_date<=?",(r["id"],date(now.year,12,31).isoformat())).fetchone()[0]
        balances.append({"name":r["surname"]+" "+r["firstname"],"earned":earned,"used":used,"remaining":earned-used})
    entries=c.execute("SELECT l.*,r.surname,r.firstname FROM leave_entries l JOIN residents r ON r.id=l.resident_id ORDER BY l.start_date DESC").fetchall()
    c.commit(); c.close()
    return page("""<section><h3>الإجازات / Congés / Annual leave</h3><p>30 calendar days per year. Unused days carry forward.</p><form method="post"><label>Resident <select name="resident_id" required>{% for r in residents %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}}</option>{% endfor %}</select></label><div class="grid"><label>Start <input type="date" name="start_date" required></label><label>End <input type="date" name="end_date" required></label></div><label>Notes <input name="notes" maxlength="500"></label><button>Record leave</button></form></section><section><h3>Balances</h3><div class="wrap"><table><tr><th>Resident</th><th>Accrued</th><th>Used</th><th>Remaining</th></tr>{% for b in balances %}<tr><td>{{b.name}}</td><td>{{b.earned}}</td><td>{{b.used}}</td><td>{{b.remaining}}</td></tr>{% endfor %}</table></div></section><section><h3>History</h3><div class="wrap"><table><tr><th>Resident</th><th>Start</th><th>End</th><th>Days</th><th>Notes</th></tr>{% for e in entries %}<tr><td>{{e.surname}} {{e.firstname}}</td><td>{{e.start_date}}</td><td>{{e.end_date}}</td><td>{{e.days}}</td><td>{{e.notes}}</td></tr>{% endfor %}</table></div></section>""",residents=residents,balances=balances,entries=entries)

@app.route("/external-rotations", methods=["GET","POST"])
def external_rotations():
    if guard(): return redirect("/login")
    c=db()
    if request.method=="POST":
        if super_guard(): c.close(); return ("Forbidden — Super Admin only",403)
        try:
            rid=int(request.form["resident_id"])
            resident=c.execute("SELECT id FROM residents WHERE id=?",(rid,)).fetchone()
            if not resident: raise ValueError("Resident does not exist.")
            supervisor=request.form["supervisor"].strip(); destination=request.form["destination"].strip()
            hospital=request.form["hospital"].strip(); country=request.form["country"].strip(); kind=request.form["rotation_type"].strip()
            start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
            if not all([supervisor,destination,hospital,country,kind]): raise ValueError("All external rotation fields are required.")
            if end<start: raise ValueError("Return date must not precede start date.")
            overlap=c.execute("SELECT id FROM external_rotations WHERE resident_id=? AND status<>'cancelled' AND start_date<=? AND end_date>=?",(rid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError("External rotations for this resident cannot overlap.")
            reject_external_overlap_with_local_assignment(c,rid,start,end)
            status="returned" if end<date.today() else "active"
            c.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,notes,status,created_by) VALUES(?,?,?,?,?,?,?,?,?,?,?)",(rid,supervisor,destination,hospital,country,kind,start.isoformat(),end.isoformat(),request.form.get("notes","").strip(),status,session.get("user","")))
            audit_event(c,"external_rotation.created",f"resident_id={rid}; hospital={hospital}; country={country}; start={start}; end={end}")
            c.commit(); flash("External rotation saved; overlapping local assignments are blocked.")
        except Exception as exc: c.rollback(); flash("External rotation not saved: "+str(exc)[:220])
        c.close(); return redirect("/external-rotations")
    residents=c.execute("SELECT id,surname,firstname,year FROM residents ORDER BY year DESC,surname").fetchall()
    today=date.today().isoformat()
    c.execute("UPDATE external_rotations SET status='returned',updated_at=CURRENT_TIMESTAMP WHERE status='active' AND end_date<?",(today,))
    rows=c.execute("SELECT x.*,r.surname,r.firstname FROM external_rotations x JOIN residents r ON r.id=x.resident_id ORDER BY x.start_date DESC").fetchall()
    c.commit(); c.close()
    return page("""<section><h3>التربصات في الخارج / Stages à l’étranger / External rotations</h3><p>Residents are blocked from local assignments for any overlapping internal period and become eligible after the return date.</p><form method="post"><label>Resident <select name="resident_id" required>{% for r in residents %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}} ({{r.year}})</option>{% endfor %}</select></label><div class="grid"><label>Responsible employee / Supervisor <input name="supervisor" required maxlength="160"></label><label>Destination <input name="destination" required maxlength="160"></label><label>Hospital <input name="hospital" required maxlength="160"></label><label>Country <input name="country" required maxlength="100"></label><label>Rotation type <input name="rotation_type" required maxlength="100"></label><label>Start date <input type="date" name="start_date" required></label><label>Return date <input type="date" name="end_date" required></label></div><label>Notes <input name="notes" maxlength="1000"></label><button>Save external rotation</button></form></section><section><h3>External rotation history</h3><div class="wrap"><table><tr><th>Resident</th><th>Supervisor</th><th>Destination</th><th>Hospital</th><th>Country</th><th>Type</th><th>Dates</th><th>Status</th><th>Actions</th></tr>{% for x in rows %}<tr><td>{{x.surname}} {{x.firstname}}</td><td>{{x.supervisor}}</td><td>{{x.destination}}</td><td>{{x.hospital}}</td><td>{{x.country}}</td><td>{{x.rotation_type}}</td><td>{{x.start_date}} – {{x.end_date}}</td><td>{{x.status}}</td><td><a href="/external-rotations/{{x.id}}/edit">Edit</a>{% if role=='super_admin' and x.status!='cancelled' %}<form method="post" action="/external-rotations/{{x.id}}/cancel" style="display:inline"><button>Cancel</button></form>{% endif %}</td></tr>{% endfor %}</table></div></section>""",residents=residents,rows=rows)

@app.route("/external-rotations/<int:xid>/edit", methods=["GET","POST"])
def edit_external_rotation(xid):
    if guard(): return redirect("/login")
    denied=super_guard()
    if denied: return denied
    c=db()
    row=c.execute("SELECT * FROM external_rotations WHERE id=?",(xid,)).fetchone()
    if not row:
        c.close(); return ("External rotation not found",404)
    if row["status"]=="cancelled":
        c.close(); return ("Cancelled external rotations are immutable",409)
    if request.method=="POST":
        try:
            rid=int(request.form["resident_id"])
            if not c.execute("SELECT 1 FROM residents WHERE id=?",(rid,)).fetchone(): raise ValueError("Resident does not exist.")
            supervisor=request.form["supervisor"].strip(); destination=request.form["destination"].strip()
            hospital=request.form["hospital"].strip(); country=request.form["country"].strip(); kind=request.form["rotation_type"].strip()
            start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
            if not all([supervisor,destination,hospital,country,kind]): raise ValueError("All external rotation fields are required.")
            if end<start: raise ValueError("Return date must not precede start date.")
            overlap=c.execute("SELECT id FROM external_rotations WHERE resident_id=? AND id<>? AND status<>'cancelled' AND start_date<=? AND end_date>=?",(rid,xid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError("External rotations for this resident cannot overlap.")
            reject_external_overlap_with_local_assignment(c,rid,start,end)
            old=f"resident_id={row['resident_id']}; hospital={row['hospital']}; country={row['country']}; start={row['start_date']}; end={row['end_date']}"
            status="returned" if end<date.today() else "active"
            c.execute("UPDATE external_rotations SET resident_id=?,supervisor=?,destination=?,hospital=?,country=?,rotation_type=?,start_date=?,end_date=?,notes=?,status=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(rid,supervisor,destination,hospital,country,kind,start.isoformat(),end.isoformat(),request.form.get("notes","").strip(),status,xid))
            audit_event(c,"external_rotation.updated",f"id={xid}; old=({old}); new_resident_id={rid}; hospital={hospital}; country={country}; start={start}; end={end}")
            c.commit(); c.close(); flash("External rotation updated and audited."); return redirect("/external-rotations")
        except Exception as exc:
            c.rollback(); flash("External rotation not updated: "+str(exc)[:220])
    residents=c.execute("SELECT id,surname,firstname,year FROM residents ORDER BY year DESC,surname").fetchall()
    c.close()
    return page("""<section><h3>Edit external rotation / Modifier le stage externe / تعديل التربص الخارجي</h3><form method="post"><label>Resident <select name="resident_id" required>{% for r in residents %}<option value="{{r.id}}" {% if r.id==row.resident_id %}selected{% endif %}>{{r.surname}} {{r.firstname}} ({{r.year}})</option>{% endfor %}</select></label><div class="grid"><label>Responsible employee / Supervisor <input name="supervisor" required maxlength="160" value="{{row.supervisor}}"></label><label>Destination <input name="destination" required maxlength="160" value="{{row.destination}}"></label><label>Hospital <input name="hospital" required maxlength="160" value="{{row.hospital}}"></label><label>Country <input name="country" required maxlength="100" value="{{row.country}}"></label><label>Rotation type <input name="rotation_type" required maxlength="100" value="{{row.rotation_type}}"></label><label>Start date <input type="date" name="start_date" required value="{{row.start_date}}"></label><label>Return date <input type="date" name="end_date" required value="{{row.end_date}}"></label></div><label>Notes <input name="notes" maxlength="1000" value="{{row.notes}}"></label><button>Save changes</button></form><p><a href="/external-rotations">Back to external rotations</a></p></section>""",residents=residents,row=row)

@app.post("/external-rotations/<int:xid>/cancel")
def cancel_external_rotation(xid):
    if guard(): return redirect("/login")
    denied=super_guard()
    if denied: return denied
    c=db()
    row=c.execute("SELECT status FROM external_rotations WHERE id=?",(xid,)).fetchone()
    if not row:
        c.close(); return ("External rotation not found",404)
    if row["status"]=="cancelled":
        c.close(); flash("External rotation already cancelled."); return redirect("/external-rotations")
    c.execute("UPDATE external_rotations SET status='cancelled',updated_at=CURRENT_TIMESTAMP WHERE id=?",(xid,))
    audit_event(c,"external_rotation.cancelled",f"id={xid}; previous_status={row['status']}")
    c.commit(); c.close(); flash("External rotation cancelled; its audit history is retained.")
    return redirect("/external-rotations")

if __name__=="__main__":
 init(); app.run(host="127.0.0.1",port=int(os.environ.get("PORT","5000")),debug=False)
